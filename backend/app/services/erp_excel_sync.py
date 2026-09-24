"""ERP Excel eşitleme: "Sipariş ve Depo Miktarları.xlsx" (Power Query ile ERP'den çekilen) dosyasından
açık siparişleri ve bitmiş ürün depo bakiyelerini programa aktarır.

Sayfalar:
  * "Sipariş Ana Veri": FISNO, TARIH, TERMIN, CARI_ISIM, STOK_KODU, MIKTAR, TLTUTAR, DOVIZTUTAR, DTIP, GRUP_KODU, INCKEYNO ...
      -> sipariş no = FISNO, poz no = INCKEYNO (ERP satır anahtarı), termin = TERMIN, miktar = MIKTAR (açık miktar kabul edilir).
  * "Depo - Ana Veri": STOK_KODU, DEPOKOD, HUCRE_KODU, BAKIYE ...
      -> seçilen depolardaki BAKIYE toplamı ürünün hedef stoğudur; programdaki mevcut (girişler − sevkler) ile fark
         "ERP eşitleme" notlu depo girişi olarak yazılır (artı veya eksi). Dosyada olmayan ürün 0 kabul edilir (seçilebilir).
Sipariş satırları mevcut sipariş aktarımından geçer (rota kapısı dahil); listede olmayan açık siparişler isteğe bağlı kapatılır.
"""
from __future__ import annotations

import io
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date

from openpyxl import load_workbook
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.models import ImportLog, Item, Order, Reservation, Shipment, StockReceipt
from app.services.excel import _date, _float, _str, import_orders

ORDERS_SHEET = "Sipariş Ana Veri"
STOCK_SHEET = "Depo - Ana Veri"
ORDER_NOTE = "ERP Excel eşitleme"
DEFAULT_WAREHOUSES = ["69"]  # sevke hazır bitmiş ürün deposu (kullanıcı kararı)
RESERVATION_SOURCE = "erp"


def _norm(s: str) -> str:
    return "".join(ch for ch in str(s or "").upper() if ch.isalnum())


def _sheet(wb, wanted: str):
    key = _norm(wanted)
    for ws in wb.worksheets:
        if _norm(ws.title) == key:
            return ws
    return None


def _rows(ws) -> list[dict]:
    it = ws.iter_rows(values_only=True)
    header = next(it, None)
    if not header:
        return []
    cols = [_norm(h) for h in header]
    out = []
    for r in it:
        if r is None or all(v in (None, "") for v in r):
            continue
        out.append({cols[i]: r[i] for i in range(min(len(cols), len(r))) if cols[i]})
    return out


@dataclass
class ErpWorkbook:
    orders: list[dict] = field(default_factory=list)  # program sipariş şablonu anahtarlarıyla
    order_raw_count: int = 0
    skipped_orders: list[str] = field(default_factory=list)
    stock_rows: list[dict] = field(default_factory=list)  # {item_code, warehouse, balance}
    warehouses: dict[str, int] = field(default_factory=dict)  # depo kodu -> satır sayısı
    missing_sheets: list[str] = field(default_factory=list)


def parse_workbook(content: bytes) -> ErpWorkbook:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    out = ErpWorkbook()
    ws = _sheet(wb, ORDERS_SHEET)
    if ws is None:
        out.missing_sheets.append(ORDERS_SHEET)
    else:
        for r in _rows(ws):
            out.order_raw_count += 1
            fis = _str(r.get("FISNO"))
            code = _str(r.get("STOKKODU"))
            qty = _float(r.get("MIKTAR"), 0.0) or 0.0
            if not fis or not code:
                out.skipped_orders.append(f"{fis or '?'} / {code or '?'}: sipariş no veya stok kodu boş")
                continue
            if qty <= 0:
                out.skipped_orders.append(f"{fis} / {code}: miktar {qty:g} (atlandı)")
                continue
            due = r.get("TERMIN") or r.get("TARIH")
            if due in (None, ""):
                out.skipped_orders.append(f"{fis} / {code}: termin boş")
                continue
            tl = _float(r.get("TLTUTAR"), 0.0) or 0.0
            dv = _float(r.get("DOVIZTUTAR"), 0.0) or 0.0
            price = (tl / qty) if tl > 0 else ((dv / qty) if dv > 0 else 0.0)
            grp = _str(r.get("GRUPKODU")).upper()
            out.orders.append({
                "_row": out.order_raw_count + 1, "reserve": max(0.0, _float(r.get("REZERV"), 0.0) or 0.0),
                "order_no": fis, "position_no": _str(r.get("INCKEYNO")), "customer": _str(r.get("CARIISIM")),
                "order_date": r.get("TARIH"), "due_date": due, "revised_due_date": "",
                "market": "Yurtdışı" if grp and grp not in ("YURTICI", "YURTİÇİ", "YURTICI ") and "YURTD" in grp else "Yerli",
                "item_code": code, "quantity": qty, "unit_price": round(price, 4),
                "material_status": "", "material_ready_date": "", "material_note": "",
            })
    ws = _sheet(wb, STOCK_SHEET)
    if ws is None:
        out.missing_sheets.append(STOCK_SHEET)
    else:
        for r in _rows(ws):
            code = _str(r.get("STOKKODU"))
            if not code:
                continue
            wh = _str(r.get("DEPOKOD"))
            bal = _float(r.get("BAKIYE"), 0.0) or 0.0
            out.stock_rows.append({"item_code": code, "warehouse": wh, "balance": bal})
            out.warehouses[wh] = out.warehouses.get(wh, 0) + 1
    wb.close()
    return out


def _on_hand_map(db: Session) -> dict[int, float]:
    rec = {i: float(s or 0) for i, s in db.query(StockReceipt.item_id, func.sum(StockReceipt.quantity)).group_by(StockReceipt.item_id).all()}
    shp = {i: float(s or 0) for i, s in db.query(Shipment.item_id, func.sum(Shipment.quantity)).group_by(Shipment.item_id).all()}
    return {i: rec.get(i, 0.0) - shp.get(i, 0.0) for i in set(rec) | set(shp)}


def stock_targets(db: Session, wbk: ErpWorkbook, warehouses: list[str] | None) -> tuple[dict[int, float], list[str]]:
    """Seçilen depolardaki bakiye toplamı -> item_id. Programda olmayan stok kodları ayrı döner."""
    items = {i.code.upper(): i for i in db.query(Item).all()}
    sel = set(warehouses or DEFAULT_WAREHOUSES)
    target: dict[int, float] = defaultdict(float)
    unknown: dict[str, float] = defaultdict(float)
    for r in wbk.stock_rows:
        if r["warehouse"] not in sel:
            continue
        it = items.get(r["item_code"].upper())
        if it is None:
            unknown[r["item_code"]] += r["balance"]
        else:
            target[it.id] += r["balance"]
    return dict(target), [f"{k} ({v:g})" for k, v in sorted(unknown.items())]


def stock_adjustments(db: Session, target: dict[int, float], *, zero_missing: bool) -> list[tuple[int, float, float]]:
    """(item_id, mevcut, fark) listesi; fark 0 olanlar atlanır. zero_missing: dosyada olmayan (seçili depolarda) ürün 0'a çekilir."""
    on_hand = _on_hand_map(db)
    ids = set(target) | (set(i for i, v in on_hand.items() if abs(v) > 1e-9) if zero_missing else set())
    out = []
    for iid in ids:
        cur = on_hand.get(iid, 0.0)
        want = target.get(iid, 0.0)
        d = want - cur
        if abs(d) > 1e-6:
            out.append((iid, cur, d))
    return out


def preview(db: Session, content: bytes, *, warehouses: list[str] | None, close_missing: bool, zero_missing: bool) -> dict:
    wbk = parse_workbook(content)
    items = {i.code.upper(): i for i in db.query(Item).all()}
    open_orders = db.query(Order).filter(Order.status == "open").all()
    existing = {(o.order_no.upper(), (o.position_no or "").upper()) for o in open_orders}
    file_keys = {(r["order_no"].upper(), r["position_no"].upper()) for r in wbk.orders}
    unknown_items = sorted({r["item_code"] for r in wbk.orders if r["item_code"].upper() not in items})
    from app.services.order_routing_gate import routing_gaps_bulk
    file_items = [items[c] for c in {r["item_code"].upper() for r in wbk.orders} if c in items]
    gaps = routing_gaps_bulk(db, file_items)
    no_routing = sorted(f"{it.code}: {gaps[it.id]}" for it in file_items if gaps.get(it.id))
    to_close = [o for o in open_orders if (o.order_no.upper(), (o.position_no or "").upper()) not in file_keys]
    target, unknown_stock = stock_targets(db, wbk, warehouses)
    adjustments = stock_adjustments(db, target, zero_missing=zero_missing)
    code_by_id = {i.id: i.code for i in items.values()}
    return {
        "missing_sheets": wbk.missing_sheets,
        "orders": {
            "file_rows": wbk.order_raw_count, "valid_rows": len(wbk.orders), "skipped": wbk.skipped_orders[:50], "skipped_count": len(wbk.skipped_orders),
            "new": sum(1 for k in file_keys if k not in existing), "existing": sum(1 for k in file_keys if k in existing),
            "unknown_items": unknown_items[:50], "unknown_item_count": len(unknown_items),
            "no_routing": no_routing[:50], "no_routing_count": len(no_routing),
            "to_close": [f"{o.order_no}{'/' + o.position_no if o.position_no else ''} {o.item.code if o.item else ''} {o.quantity:g}" for o in to_close[:50]],
            "to_close_count": len(to_close) if close_missing else 0, "would_close_count": len(to_close),
        },
        "stock": {
            "file_rows": len(wbk.stock_rows), "warehouses": [{"code": k, "rows": v, "selected": k in set(warehouses or DEFAULT_WAREHOUSES)} for k, v in sorted(wbk.warehouses.items())],
            "items_in_file": len(target), "unknown_items": unknown_stock[:50], "unknown_item_count": len(unknown_stock),
            "adjustments": len(adjustments), "increase": round(sum(d for _, _, d in adjustments if d > 0), 2), "decrease": round(-sum(d for _, _, d in adjustments if d < 0), 2),
            "examples": [f"{code_by_id.get(i, i)}: {c:g} → {c + d:g} ({d:+g})" for i, c, d in sorted(adjustments, key=lambda x: -abs(x[2]))[:30]],
        },
    }


def apply(db: Session, content: bytes, *, username: str, filename: str, do_orders: bool, do_stock: bool,
          warehouses: list[str] | None, close_missing: bool, zero_missing: bool) -> dict:
    """Sıra: önce depo bakiyesi (stok), sonra siparişler ve ERP rezervasyonları (dış stok kredisi)."""
    wbk = parse_workbook(content)
    result: dict = {"orders": None, "stock": None, "missing_sheets": wbk.missing_sheets}
    today = date.today()
    if do_stock and STOCK_SHEET not in wbk.missing_sheets:
        target, unknown = stock_targets(db, wbk, warehouses)
        adjustments = stock_adjustments(db, target, zero_missing=zero_missing)
        sel = ", ".join(warehouses or DEFAULT_WAREHOUSES)
        for iid, _cur, d in adjustments:
            db.add(StockReceipt(item_id=iid, receipt_date=today, quantity=round(d, 3), lot="", note=f"{ORDER_NOTE} (depo {sel})", source="import", created_by=username))
        db.add(ImportLog(kind="stock_receipts", filename=f"{filename} ({ORDER_NOTE})", username=username, inserted=len(adjustments), updated=0,
                         errors="", warnings="\n".join(f"Programda olmayan stok: {u}" for u in unknown)[:10000]))
        db.flush()
        result["stock"] = {"adjustments": len(adjustments), "increase": round(sum(d for _, _, d in adjustments if d > 0), 2),
                           "decrease": round(-sum(d for _, _, d in adjustments if d < 0), 2), "unknown_item_count": len(unknown), "warehouses": sel}
    if do_orders and ORDERS_SHEET not in wbk.missing_sheets:
        ins, upd, _removed, errs = import_orders(db, wbk.orders, remove_missing=False)
        db.flush()
        closed = 0
        file_keys = {(r["order_no"].upper(), r["position_no"].upper()) for r in wbk.orders}
        open_by_key = {(o.order_no.upper(), (o.position_no or "").upper()): o for o in db.query(Order).filter(Order.status == "open").all()}
        if close_missing:
            for k, o in open_by_key.items():
                if k not in file_keys:
                    o.status = "closed"
                    closed += 1
                    db.query(Reservation).filter(Reservation.order_id == o.id, Reservation.source == RESERVATION_SOURCE).delete(synchronize_session=False)
        # REZERV: siparişe ayrılmış mamul stoğu -> dış stok rezervasyonu (üretim ihtiyacından düşülür); her eşitlemede yenilenir
        reserved = 0
        for r in wbk.orders:
            o = open_by_key.get((r["order_no"].upper(), r["position_no"].upper()))
            if o is None:
                continue
            db.query(Reservation).filter(Reservation.order_id == o.id, Reservation.source == RESERVATION_SOURCE).delete(synchronize_session=False)
            q = min(float(r.get("reserve") or 0.0), float(o.quantity or 0.0))
            if q > 1e-6:
                db.add(Reservation(item_id=o.item_id, order_id=o.id, quantity=round(q, 3), source=RESERVATION_SOURCE,
                                   stock_provenance="external_finished_stock", note=f"{ORDER_NOTE}: ERP REZERV", created_by=username))
                reserved += 1
        applied = (ins + upd) > 0 or not errs
        db.add(ImportLog(kind="orders", filename=f"{filename} ({ORDER_NOTE})", username=username, inserted=ins, updated=upd,
                         errors="" if applied else "\n".join(errs)[:10000], warnings="\n".join((errs if applied else []) + wbk.skipped_orders)[:10000]))
        result["orders"] = {"inserted": ins, "updated": upd, "closed": closed, "reserved": reserved, "errors": errs[:100], "error_count": len(errs), "skipped_count": len(wbk.skipped_orders)}
    db.commit()
    return result
