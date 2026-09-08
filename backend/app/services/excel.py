"""Excel import (sablon + yukleme) ve export (yedek, raporlar)."""

from __future__ import annotations

import io
import re
import unicodedata
from datetime import date, datetime, time, timedelta
from typing import Any, Callable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, joinedload

from app.models import (
    BomLine,
    Downtime,
    Employee,
    ImportLog,
    Item,
    Machine,
    OpTransitionRule,
    Order,
    PlanLine,
    ProductionActual,
    Reservation,
    RoutingOperation,
    Shipment,
    StockReceipt,
    User,
    WorkCenter,
    WorkCenterShift,
    WorkCenterWeek,
)
from app.schemas import ImportResult, OrderImportChangePreview, OrderImportPreview, OrderImportRowPreview

TR_MAP = str.maketrans("çğıöşüÇĞİÖŞÜ", "cgiosuCGIOSU")


def norm(s: Any) -> str:
    s = str(s or "").translate(TR_MAP)
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", s.lower())


# ---- Sablon tanimlari: kind -> (baslik listesi, ornek satir, zorunlu alanlar) ----
# key -> (Excel basligi, kabul edilen alternatif basliklar)
TEMPLATES: dict[str, dict] = {
    "workcenters": {
        "title": "İş Merkezleri",
        "columns": [
            ("code", "İş Merkezi Kodu", ["kod", "ismerkezi"]),
            ("name", "İş Merkezi Adı", ["ad", "adi"]),
            ("description", "Açıklama", []),
            ("is_active", "Aktif (E/H)", ["aktif"]),
            ("is_planned", "Planlanıyor (E/H)", ["planlaniyor", "pilot"]),
            ("capacity_unit_hours", "Birim Saat", ["kapasitebirimi", "birim"]),
            ("default_efficient_hours", "Kişi Başı Verimli Saat", ["verimlisaat", "verimlisure"]),
            ("area_code", "Alan Kodu", ["alan", "alankodu", "grupkodu"]),
            ("area_name", "Alan Adı", ["alanadi", "grup", "grupadi"]),
            ("capacity_source", "Kapasite Kaynağı (İM/Makine)", ["kapasitekaynagi", "kaynak"]),
        ],
        "example": ["PRESHANE 1", "PRESHANE 1", "", "E", "E", 10, 4, "PRS", "PRESHANELER", "İM"],
        "required": ["code", "name"],
    },
    "machines": {
        "title": "Makineler",
        "columns": [
            ("wc_code", "İş Merkezi Kodu", ["ismerkezi", "ismerkezikodu"]),
            ("code", "Makine Kodu", ["makine", "makinekodu", "tezgahkodu", "kod"]),
            ("name", "Makine Adı", ["makineadi", "ad", "adi"]),
            ("description", "Açıklama", []),
            ("is_active", "Aktif (E/H)", ["aktif"]),
        ],
        "example": ["PRESHANE 1", "PRS1-EKS-01", "Eksantrik Pres 60t", "", "E"],
        "required": ["wc_code", "code"],
    },
    "shifts": {
        "title": "Vardiyalar",
        "columns": [
            ("wc_code", "İş Merkezi Kodu", ["ismerkezi", "ismerkezikodu"]),
            ("name", "Vardiya", ["vardiyaadi", "ad"]),
            ("weekdays", "Günler (Pzt=0..Paz=6)", ["gunler"]),
            ("start_time", "Başlangıç", ["baslangicsaati", "baslama"]),
            ("end_time", "Bitiş", ["bitissaati"]),
            ("headcount", "Kişi Sayısı", ["kisi", "kisisayisi"]),
            ("efficient_hours_per_person", "Kişi Başı Verimli Saat", ["verimlisaat", "verimlisure"]),
        ],
        "example": ["TZG-A", "Gündüz", "0,1,2,3,4", "08:00", "18:00", 10, 4],
        "required": ["wc_code"],
    },
    "employees": {
        "title": "Personel",
        "columns": [
            ("code", "Sicil No", ["sicil", "kod", "personelkodu"]),
            ("name", "Ad Soyad", ["ad", "adsoyad", "personel"]),
            ("wc_code", "İş Merkezi Kodu", ["ismerkezi", "ismerkezikodu"]),
            ("machine_code", "Makine Kodu", ["makine", "makinekodu", "tezgah"]),
            ("is_active", "Aktif (E/H)", ["aktif"]),
        ],
        "example": ["1001", "Ahmet Yılmaz", "PRESHANE 1", "PRS1-EKS-01", "E"],
        "required": ["code", "name"],
    },
    "items": {
        "title": "Stok Kodları",
        "columns": [
            ("code", "Stok Kodu", ["stokkodu", "kod", "malzeme"]),
            ("name", "Stok Adı", ["stokadi", "ad", "aciklama"]),
            ("product_group", "Ürün Grubu", ["grup", "urungrubu"]),
            ("unit", "Birim", []),
        ],
        "example": ["MAM-0001", "Endüstriyel Ocak 4 Gözlü", "OCAK", "AD"],
        "required": ["code"],
    },
    "bom": {
        "title": "BOM (Hammadde)",
        "columns": [
            ("item_code", "Stok Kodu", ["stokkodu", "mamul", "anastok"]),
            ("component_code", "Bileşen Kodu", ["hammaddekodu", "bilesen", "hammadde"]),
            ("component_name", "Bileşen Adı", ["hammaddeadi", "bilesenadi"]),
            ("quantity", "Miktar", ["kullanimmiktari"]),
            ("unit", "Birim", []),
        ],
        "example": ["MAM-0001", "HM-SAC-2MM", "Paslanmaz Sac 2mm", 3.5, "KG"],
        "required": ["item_code", "component_code"],
    },
    "routing": {
        "title": "Rota / Çevrim Süreleri",
        "columns": [
            ("item_code", "Stok Kodu", ["stokkodu", "mamul"]),
            ("seq", "Sıra", ["operasyonsira", "sira", "asama"]),
            ("operation_name", "Operasyon", ["operasyonadi", "islem"]),
            ("wc_code", "İş Merkezi Kodu", ["ismerkezi", "tezgah", "tezgahkodu"]),
            ("cycle_time_sec", "Çevrim Süresi (sn)", ["cycletime", "cevrimsuresi", "cevrimsuresisn", "cevrim"]),
            ("setup_time_min", "Setup (dk)", ["setup", "hazirlik", "setupsuresi"]),
            ("semi_finished_code", "Yarımamül Kodu", ["yarimamul", "yarimamulkodu", "wip", "wipkodu", "semifinished"]),
        ],
        "example": ["MAM-0001", 10, "Kesim", "TZG-A", 50, 15, "MAM-0001-K10"],
        "required": ["item_code", "seq", "wc_code", "cycle_time_sec"],
    },
    "orders": {
        "title": "Siparişler",
        "columns": [
            ("order_no", "Sipariş No", ["siparis", "siparisno", "belgeno"]),
            ("position_no", "Poz No", ["poz", "pozno", "pozisyon", "pozisyonno"]),
            ("customer", "Müşteri", ["musteriadi", "cari"]),
            ("due_date", "Termin", ["termintarihi", "teslimtarihi", "tarih"]),
            ("item_code", "Stok Kodu", ["stokkodu", "malzeme"]),
            ("quantity", "Miktar", ["adet"]),
            ("unit_price", "Birim Fiyat", ["fiyat", "birimfiyat", "satisfiyati", "birimsatisfiyati"]),
        ],
        "example": ["SIP-2026-001", "10", "ABC Otel", "2026-10-15", "MAM-0001", 40, 1250],
        "required": ["order_no", "due_date", "item_code", "quantity"],
    },
    "production": {
        "title": "Günlük Üretim",
        "columns": [
            ("prod_date", "Tarih", ["uretimtarihi", "gun"]),
            ("semi_finished_code", "Yarımamül Kodu", ["yarimamul", "yarimamulkodu", "wip", "wipkodu"]),
            ("quantity", "Miktar", ["adet", "uretilen", "uretimmiktari"]),
            ("order_no", "Sipariş No", ["siparis", "siparisno"]),
            ("wc_code", "İş Merkezi Kodu (opsiyonel)", ["ismerkezi", "tezgah", "ismerkezikodu"]),
            ("item_code", "Stok Kodu (opsiyonel)", ["stokkodu", "malzeme"]),
            ("operation_seq", "Operasyon Sıra (opsiyonel)", ["sira", "operasyon", "operasyonsira"]),
            ("reported_hours", "Fiili Süre (saat)", ["fiilisure", "calismasuresi", "sure"]),
        ],
        "example": ["2026-09-06", "MAM-0001-K10", 120, "SIP-2026-001", "", "", "", ""],
        "required": ["prod_date", "quantity"],
    },
    "downtime": {
        "title": "Günlük Duruşlar",
        "columns": [
            ("dt_date", "Tarih", ["durustarihi", "gun"]),
            ("wc_code", "İş Merkezi Kodu", ["ismerkezi", "tezgah"]),
            ("reason_code", "Sebep Kodu", ["sebepkodu", "kod"]),
            ("reason_desc", "Sebep", ["sebepaciklama", "aciklama", "durussebebi"]),
            ("minutes", "Süre (dk)", ["sure", "dakika", "durussuresi"]),
        ],
        "example": ["2026-09-06", "TZG-A", "MLZ", "Malzeme bekleme", 45],
        "required": ["dt_date", "wc_code", "minutes"],
    },
    "wc_weeks": {
        "title": "Haftalık İş Gücü",
        "columns": [
            ("wc_code", "İş Merkezi Kodu", ["ismerkezi", "ismerkezikodu"]),
            ("week_start", "Hafta (Pzt tarihi veya 2026-W37)", ["hafta", "haftabaslangici", "haftano"]),
            ("headcount", "Kişi Sayısı", ["kisi", "kisisayisi"]),
            ("efficient_hours_per_person", "Kişi Başı Verimli Saat", ["verimlisaat", "verimlisure"]),
            ("working_days", "Çalışma Günü", ["gun", "gunsayisi", "calismagunu"]),
            ("note", "Not", ["aciklama"]),
        ],
        "example": ["PRESHANE 3", "2026-09-07", 8, 4.5, 5, "1 kişi izinli"],
        "required": ["wc_code", "week_start"],
    },
    "op_rules": {
        "title": "Senaryo Kuralları",
        "columns": [
            ("product_group", "Ürün Grubu", ["grup", "urungrubu"]),
            ("item_code", "Stok Kodu (boş = grup geneli)", ["stok", "stokkodu", "malzemekodu"]),
            ("from_op", "Önceki Operasyon", ["onceki", "oncekioperasyon", "kaynak"]),
            ("to_op", "Sonraki Operasyon", ["sonraki", "sonrakioperasyon", "hedef"]),
            ("from_wip_code", "Önceki Yarımamül Kodu", ["oncekiyarimamul", "kaynakwip"]),
            ("to_wip_code", "Sonraki Yarımamül Kodu", ["sonrakiyarimamul", "hedefwip"]),
            ("rule", "Kural (Bitiş/Çevrim)", ["kuraltipi", "tip"]),
            ("lag_cycles", "Çevrim Sayısı", ["cevrim", "cevrimsayisi", "adet"]),
            ("wait_minutes", "Bekleme (dk)", ["bekleme", "beklemesuresi"]),
            ("note", "Not", ["aciklama"]),
        ],
        "example": ["EVYE", "", "Sıvama", "Forma", "6005510-10", "6005510-11", "Çevrim", 5, 0, "5 parça sıvandıktan sonra forma başlar"],
        "required": ["from_op", "to_op"],
    },
    "stock_receipts": {
        "title": "Depo Girişi (Bitmiş Ürün)",
        "columns": [
            ("receipt_date", "Tarih", ["giristarihi", "gun"]),
            ("item_code", "Stok Kodu", ["stok", "stokkodu", "malzemekodu"]),
            ("quantity", "Miktar", ["adet", "miktar"]),
            ("lot", "Lot / Parti", ["lot", "parti", "partino"]),
            ("note", "Not", ["aciklama"]),
        ],
        "example": ["2026-09-06", "MAM-0001", 120, "L-2609", ""],
        "required": ["receipt_date", "item_code", "quantity"],
    },
}


def parse_week(v: Any) -> date:
    """'2026-09-07', '07.09.2026' veya '2026-W37' / '2026-H37' / '37' (yil = bugun) -> haftanin Pazartesisi."""
    s = _str(v)
    m = re.fullmatch(r"(?:(\d{4})[-/ ]?)?[WwHh]?(\d{1,2})", s)
    if m and not re.search(r"[./]", s) and len(s) <= 8:
        year = int(m.group(1)) if m.group(1) else date.today().year
        week = int(m.group(2))
        if 1 <= week <= 53:
            return date.fromisocalendar(year, week, 1)
    d = _date(v)
    return d - timedelta(days=d.weekday())


def sheet_title(title: str) -> str:
    """Excel sayfa adi kurallari: max 31 karakter, [] : * ? / \\ yasak."""
    cleaned = re.sub(r"[\[\]:*?/\\]", "-", title).strip()
    return (cleaned or "Sayfa")[:31]


def _autosize(ws) -> None:
    for i, col in enumerate(ws.columns, start=1):
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)


def _style_header(ws) -> None:
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")


def build_template(kind: str) -> bytes:
    t = TEMPLATES[kind]
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title(t["title"])
    ws.append([c[1] for c in t["columns"]])
    ws.append(t["example"])
    _style_header(ws)
    _autosize(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- Parse ----

def read_rows(content: bytes, kind: str) -> tuple[list[dict], list[str]]:
    """Excel -> [ {key: value} ], hatalar. Basliklar Turkce/alias uyumlu eslenir."""
    t = TEMPLATES[kind]
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return [], ["Dosya bos"]
    alias_map: dict[str, str] = {}
    for key, title, aliases in t["columns"]:
        alias_map[norm(title)] = key
        alias_map[norm(key)] = key
        for a in aliases:
            alias_map[norm(a)] = key
    col_keys: list[str | None] = [alias_map.get(norm(h)) for h in header]
    missing = [title for key, title, _ in t["columns"] if key in t["required"] and key not in col_keys]
    if missing:
        return [], [f"Zorunlu sutun(lar) bulunamadi: {', '.join(missing)}"]
    out = []
    for r_idx, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        rec = {"_row": r_idx}
        for key, val in zip(col_keys, row):
            if key:
                rec[key] = val
        out.append(rec)
    return out, []


def _bool(v: Any, default: bool = True) -> bool:
    if v is None or str(v).strip() == "":
        return default
    return norm(v) in {"e", "evet", "1", "true", "x", "yes", "y"}


def _float(v: Any, default: float | None = None) -> float | None:
    if v is None or str(v).strip() == "":
        return default
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", ".")
    return float(s)


def _int(v: Any, default: int | None = None) -> int | None:
    f = _float(v, None)
    return default if f is None else int(round(f))


def _str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _date(v: Any) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Tarih anlasilamadi: {v!r}")


def _time(v: Any, default: time) -> time:
    if v is None or str(v).strip() == "":
        return default
    if isinstance(v, datetime):
        return v.time()
    if isinstance(v, time):
        return v
    if isinstance(v, (int, float)):  # Excel gun kesri
        total_min = int(round(float(v) * 24 * 60)) % (24 * 60)
        return time(total_min // 60, total_min % 60)
    s = str(v).strip()
    for fmt in ("%H:%M", "%H:%M:%S", "%H.%M"):
        try:
            return datetime.strptime(s, fmt).time()
        except ValueError:
            pass
    raise ValueError(f"Saat anlasilamadi: {v!r}")


# ---- Importers ----

def _wc_lookup(db: Session) -> dict[str, WorkCenter]:
    return {w.code.upper(): w for w in db.query(WorkCenter).all()}


def _item_lookup(db: Session) -> dict[str, Item]:
    return {i.code.upper(): i for i in db.query(Item).all()}


def _get_or_create_item(db: Session, items: dict[str, Item], code: str) -> Item:
    key = code.upper()
    if key not in items:
        it = Item(code=code, name="")
        db.add(it)
        db.flush()
        items[key] = it
    return items[key]


def import_workcenters(db: Session, rows: list[dict]) -> tuple[int, int, list[str]]:
    ins = upd = 0
    errs = []
    existing = _wc_lookup(db)
    for r in rows:
        try:
            code = _str(r.get("code"))
            if not code:
                raise ValueError("Kod bos")
            wc = existing.get(code.upper())
            if not wc:
                wc = WorkCenter(code=code, name=_str(r.get("name")) or code)
                db.add(wc)
                existing[code.upper()] = wc
                ins += 1
            else:
                upd += 1
            if _str(r.get("name")):
                wc.name = _str(r.get("name"))
            wc.description = _str(r.get("description"))
            wc.is_active = _bool(r.get("is_active"), True)
            wc.is_planned = _bool(r.get("is_planned"), wc.is_planned or False)
            wc.capacity_unit_hours = _float(r.get("capacity_unit_hours"), wc.capacity_unit_hours or 10.0)
            wc.default_efficient_hours = _float(r.get("default_efficient_hours"), wc.default_efficient_hours or 4.0)
            if _str(r.get("area_code")) or _str(r.get("area_name")):
                wc.area_code = _str(r.get("area_code"))
                wc.area_name = _str(r.get("area_name"))
            src = norm(_str(r.get("capacity_source")))
            if src:
                wc.capacity_source = "machines" if src.startswith("makin") or src == "machines" else "work_center"
        except Exception as e:  # noqa: BLE001
            errs.append(f"Satir {r['_row']}: {e}")
    return ins, upd, errs


def import_machines(db: Session, rows: list[dict]) -> tuple[int, int, list[str]]:
    """Makine kodu benzersizdir; varsa gunceller (is merkezi degisebilir), yoksa ekler."""
    ins = upd = 0
    errs = []
    wcs = _wc_lookup(db)
    existing = {m.code.upper(): m for m in db.query(Machine).all()}
    for r in rows:
        try:
            wc = wcs.get(_str(r.get("wc_code")).upper())
            if not wc:
                raise ValueError(f"Is merkezi bulunamadi: {r.get('wc_code')}")
            code = _str(r.get("code"))
            if not code:
                raise ValueError("Makine kodu bos")
            m = existing.get(code.upper())
            if not m:
                m = Machine(code=code, work_center=wc)
                db.add(m)
                existing[code.upper()] = m
                ins += 1
            else:
                m.work_center_id = wc.id
                upd += 1
            m.name = _str(r.get("name")) or m.name or code
            m.description = _str(r.get("description"))
            m.is_active = _bool(r.get("is_active"), True)
        except Exception as e:  # noqa: BLE001
            errs.append(f"Satir {r['_row']}: {e}")
    db.flush()
    return ins, upd, errs


def import_shifts(db: Session, rows: list[dict]) -> tuple[int, int, list[str]]:
    """Ayni is merkezi + vardiya adi varsa gunceller, yoksa ekler."""
    ins = upd = 0
    errs = []
    wcs = _wc_lookup(db)
    for r in rows:
        try:
            wc = wcs.get(_str(r.get("wc_code")).upper())
            if not wc:
                raise ValueError(f"Is merkezi bulunamadi: {r.get('wc_code')}")
            name = _str(r.get("name")) or "Gunduz"
            sh = next((s for s in wc.shifts if s.name.lower() == name.lower()), None)
            if not sh:
                sh = WorkCenterShift(work_center=wc, name=name)
                db.add(sh)
                ins += 1
            else:
                upd += 1
            wd = _str(r.get("weekdays")).replace(";", ",").replace(" ", "")
            sh.weekdays = wd or "0,1,2,3,4"
            sh.start_time = _time(r.get("start_time"), time(8, 0))
            sh.end_time = _time(r.get("end_time"), time(18, 0))
            sh.headcount = _int(r.get("headcount"), 0) or 0
            sh.efficient_hours_per_person = _float(r.get("efficient_hours_per_person"), None)
        except Exception as e:  # noqa: BLE001
            errs.append(f"Satir {r['_row']}: {e}")
    return ins, upd, errs


def import_employees(db: Session, rows: list[dict]) -> tuple[int, int, list[str]]:
    ins = upd = 0
    errs = []
    wcs = _wc_lookup(db)
    machines = {m.code.upper(): m for m in db.query(Machine).all()}
    existing = {e.code.upper(): e for e in db.query(Employee).all()}
    for r in rows:
        try:
            code = _str(r.get("code"))
            if not code:
                raise ValueError("Sicil bos")
            emp = existing.get(code.upper())
            if not emp:
                emp = Employee(code=code, name=_str(r.get("name")))
                db.add(emp)
                existing[code.upper()] = emp
                ins += 1
            else:
                upd += 1
            emp.name = _str(r.get("name")) or emp.name
            wc_code = _str(r.get("wc_code"))
            if wc_code:
                wc = wcs.get(wc_code.upper())
                if not wc:
                    raise ValueError(f"Is merkezi bulunamadi: {wc_code}")
                emp.work_center_id = wc.id
            m_code = _str(r.get("machine_code"))
            if m_code:
                m = machines.get(m_code.upper())
                if not m:
                    raise ValueError(f"Makine bulunamadi: {m_code}")
                if emp.work_center_id and emp.work_center_id != m.work_center_id:
                    raise ValueError(f"Makine {m.code} personelin is merkezine ait degil")
                emp.machine_id = m.id
                emp.work_center_id = m.work_center_id
            elif "machine_code" in r:
                emp.machine_id = None  # kolon var ama bos => atama kaldirildi
            emp.is_active = _bool(r.get("is_active"), True)
        except Exception as e:  # noqa: BLE001
            errs.append(f"Satir {r['_row']}: {e}")
    return ins, upd, errs


def import_items(db: Session, rows: list[dict]) -> tuple[int, int, list[str]]:
    ins = upd = 0
    errs = []
    items = _item_lookup(db)
    for r in rows:
        try:
            code = _str(r.get("code"))
            if not code:
                raise ValueError("Stok kodu bos")
            it = items.get(code.upper())
            if not it:
                it = Item(code=code)
                db.add(it)
                items[code.upper()] = it
                ins += 1
            else:
                upd += 1
            it.name = _str(r.get("name")) or it.name
            it.product_group = _str(r.get("product_group")) or it.product_group
            it.unit = _str(r.get("unit")) or it.unit or "AD"
        except Exception as e:  # noqa: BLE001
            errs.append(f"Satir {r['_row']}: {e}")
    return ins, upd, errs


def import_bom(db: Session, rows: list[dict]) -> tuple[int, int, list[str]]:
    ins = upd = 0
    errs = []
    items = _item_lookup(db)
    for r in rows:
        try:
            item = _get_or_create_item(db, items, _str(r.get("item_code")))
            comp = _str(r.get("component_code"))
            if not comp:
                raise ValueError("Bilesen kodu bos")
            line = next((b for b in item.bom_lines if b.component_code.upper() == comp.upper()), None)
            if not line:
                line = BomLine(item=item, component_code=comp)
                db.add(line)
                ins += 1
            else:
                upd += 1
            line.component_name = _str(r.get("component_name")) or line.component_name
            line.quantity = _float(r.get("quantity"), 1.0)
            line.unit = _str(r.get("unit")) or line.unit or "AD"
        except Exception as e:  # noqa: BLE001
            errs.append(f"Satir {r['_row']}: {e}")
    return ins, upd, errs


def import_routing(db: Session, rows: list[dict]) -> tuple[int, int, list[str]]:
    ins = upd = 0
    errs = []
    items = _item_lookup(db)
    wcs = _wc_lookup(db)
    for r in rows:
        try:
            item = _get_or_create_item(db, items, _str(r.get("item_code")))
            seq = _int(r.get("seq"))
            if seq is None:
                raise ValueError("Sira bos")
            wc = wcs.get(_str(r.get("wc_code")).upper())
            if not wc:
                raise ValueError(f"Is merkezi bulunamadi: {r.get('wc_code')}")
            op = next((o for o in item.operations if o.seq == seq), None)
            if not op:
                op = RoutingOperation(item=item, seq=seq, work_center_id=wc.id)
                db.add(op)
                ins += 1
            else:
                upd += 1
            op.work_center_id = wc.id
            op.operation_name = _str(r.get("operation_name")) or op.operation_name
            op.cycle_time_sec = _float(r.get("cycle_time_sec"), 0.0) or 0.0
            op.setup_time_min = _float(r.get("setup_time_min"), 0.0) or 0.0
            wip = _str(r.get("semi_finished_code"))
            if wip:
                op.semi_finished_code = wip
        except Exception as e:  # noqa: BLE001
            errs.append(f"Satir {r['_row']}: {e}")
    return ins, upd, errs


def _order_preview_from_row(r: dict, item_code: str, order_id: int | None = None) -> OrderImportRowPreview:
    due = None
    try:
        due = _date(r.get("due_date"))
    except Exception:  # noqa: BLE001
        pass
    return OrderImportRowPreview(
        order_id=order_id,
        order_no=_str(r.get("order_no")),
        position_no=_str(r.get("position_no")),
        item_code=item_code,
        customer=_str(r.get("customer")),
        due_date=due,
        quantity=_float(r.get("quantity")),
        unit_price=_float(r.get("unit_price"), None),
        excel_row=r.get("_row"),
    )


def _order_preview_from_model(o: Order) -> OrderImportRowPreview:
    return OrderImportRowPreview(
        order_id=o.id,
        order_no=o.order_no,
        position_no=o.position_no or "",
        item_code=o.item.code if o.item else "",
        customer=o.customer,
        due_date=o.due_date,
        quantity=o.quantity,
        unit_price=o.unit_price,
    )


def preview_orders_import(db: Session, rows: list[dict]) -> OrderImportPreview:
    """Dosyadaki acik siparis listesi ile sistemdeki acik siparisleri karsilastirir."""
    from app.services.orders import _index_orders, _order_lookup_key

    items = _item_lookup(db)
    parse_errors: list[str] = []
    file_keys: dict[tuple, OrderImportRowPreview] = {}

    for r in rows:
        try:
            order_no = _str(r.get("order_no"))
            if not order_no:
                raise ValueError("Siparis no bos")
            pos = _str(r.get("position_no"))
            code = _str(r.get("item_code"))
            item = items.get(code.upper())
            if not item:
                raise ValueError(f"Stok kodu bulunamadi: {code}")
            if _float(r.get("quantity")) is None:
                raise ValueError("Miktar bos")
            _date(r.get("due_date"))
            key = _order_lookup_key(order_no, pos, item.id)
            file_keys[key] = _order_preview_from_row(r, item.code)
        except Exception as e:  # noqa: BLE001
            parse_errors.append(f"Satir {r.get('_row', '?')}: {e}")

    open_orders = (
        db.query(Order)
        .options(joinedload(Order.item))
        .filter(Order.status == "open")
        .all()
    )
    system_idx = _index_orders(open_orders)
    file_key_set = set(file_keys.keys())
    system_key_set = set(system_idx.keys())

    only_in_system = [_order_preview_from_model(system_idx[k]) for k in sorted(system_key_set - file_key_set, key=str)]
    only_in_file = [file_keys[k] for k in sorted(file_key_set - system_key_set, key=str)]

    updated: list[OrderImportChangePreview] = []
    unchanged = 0
    for key in file_key_set & system_key_set:
        o = system_idx[key]
        fp = file_keys[key]
        changes: list[str] = []
        file_due = fp.due_date
        if file_due and file_due != o.due_date:
            changes.append(f"termin: {o.due_date} → {file_due}")
        if fp.quantity is not None and abs(fp.quantity - o.quantity) > 1e-9:
            changes.append(f"miktar: {o.quantity:g} → {fp.quantity:g}")
        if fp.customer and fp.customer != o.customer:
            changes.append(f"musteri: {o.customer or '—'} → {fp.customer}")
        if fp.unit_price is not None and abs(fp.unit_price - (o.unit_price or 0)) > 1e-9:
            changes.append(f"birim fiyat: {o.unit_price or 0:g} → {fp.unit_price:g}")
        if changes:
            updated.append(
                OrderImportChangePreview(
                    order_id=o.id,
                    order_no=fp.order_no,
                    position_no=fp.position_no,
                    item_code=fp.item_code,
                    customer=fp.customer,
                    due_date=fp.due_date,
                    quantity=fp.quantity,
                    unit_price=fp.unit_price,
                    excel_row=fp.excel_row,
                    changes=changes,
                )
            )
        else:
            unchanged += 1

    return OrderImportPreview(
        parse_errors=parse_errors,
        only_in_system=only_in_system,
        only_in_file=only_in_file,
        updated=updated,
        unchanged_count=unchanged,
        file_row_count=len(rows),
        system_open_count=len(open_orders),
    )


def import_orders(db: Session, rows: list[dict], remove_missing: bool = False) -> tuple[int, int, int, list[str]]:
    """Ayni siparis+poz (veya poz bos ise siparis+stok) varsa gunceller; istege bagli listede olmayan acik siparisleri siler."""
    from app.services.orders import _index_orders, _order_lookup_key, bulk_delete_orders

    ins = upd = removed = 0
    errs = []
    items = _item_lookup(db)
    existing = _index_orders(db.query(Order).filter(Order.status != "merged").all())
    file_keys: set[tuple] = set()
    for r in rows:
        try:
            order_no = _str(r.get("order_no"))
            if not order_no:
                raise ValueError("Siparis no bos")
            pos = _str(r.get("position_no"))
            item = _get_or_create_item(db, items, _str(r.get("item_code")))
            qty = _float(r.get("quantity"))
            if qty is None:
                raise ValueError("Miktar bos")
            key = _order_lookup_key(order_no, pos, item.id)
            file_keys.add(key)
            o = existing.get(key)
            if not o:
                o = Order(order_no=order_no, position_no=pos, item_id=item.id, quantity=qty, due_date=_date(r.get("due_date")))
                db.add(o)
                existing[key] = o
                ins += 1
            else:
                upd += 1
            o.customer = _str(r.get("customer")) or o.customer
            o.due_date = _date(r.get("due_date"))
            o.quantity = qty
            o.position_no = pos
            price = _float(r.get("unit_price"), None)
            if price is not None:
                o.unit_price = price
            o.status = "open"
        except Exception as e:  # noqa: BLE001
            errs.append(f"Satir {r['_row']}: {e}")

    if remove_missing and not errs:
        to_remove: list[int] = []
        open_orders = db.query(Order).filter(Order.status == "open").all()
        for o in open_orders:
            if o.position_no:
                key = ("pos", o.order_no.upper(), o.position_no.upper())
            else:
                key = ("item", o.order_no.upper(), o.item_id)
            if key not in file_keys:
                to_remove.append(o.id)
        if to_remove:
            try:
                removed = bulk_delete_orders(db, to_remove)
            except ValueError as e:
                errs.append(str(e))
    return ins, upd, removed, errs


def import_production(db: Session, rows: list[dict]) -> tuple[int, int, list[str]]:
    """Ayni gun/yarimamul/siparis satiri varsa uzerine yazar (idempotent).

    Birincil anahtar: tarih + yarimamul kodu + siparis no.
    Yarimamul kodu rota operasyonuna cozulur; stok kodu coklu eslesmede ayirt eder.
    Geriye uyumluluk: yarimamul yoksa is merkezi + stok + operasyon sira ile eslesir.
    """
    from app.services.wip import resolve_wip, wip_index

    ins = upd = 0
    errs = []
    items = _item_lookup(db)
    wcs = _wc_lookup(db)
    wip_idx = wip_index(db)
    for r in rows:
        try:
            d = _date(r.get("prod_date"))
            order_no = _str(r.get("order_no"))
            qty = _float(r.get("quantity"))
            if qty is None:
                raise ValueError("Miktar bos")
            wip_raw = _str(r.get("semi_finished_code"))
            op = None
            item = None
            wc = None
            seq = None
            if wip_raw:
                op = resolve_wip(db, wip_raw, _str(r.get("item_code")) or None, wip_idx)
                item = op.item
                wc = op.work_center
                seq = op.seq
            else:
                wc = wcs.get(_str(r.get("wc_code")).upper())
                if not wc:
                    raise ValueError("Yarimamul kodu veya is merkezi + stok kodu gerekli")
                item = items.get(_str(r.get("item_code")).upper())
                if not item:
                    raise ValueError(f"Stok kodu bulunamadi: {r.get('item_code')}")
                seq = _int(r.get("operation_seq"), None)
                if seq is not None:
                    op = next((o for o in item.operations if o.seq == seq), None)
                if op is None:
                    op = next((o for o in item.operations if o.work_center_id == wc.id), None)
            if item is None or wc is None:
                raise ValueError("Uretim satiri cozulemedi")
            earned = op.hours_for(qty) - (op.setup_time_min / 60.0) if op else 0.0
            wip_stored = wip_raw or (op.semi_finished_code if op else "")
            if wip_stored:
                pa = (
                    db.query(ProductionActual)
                    .filter(
                        ProductionActual.prod_date == d,
                        ProductionActual.semi_finished_code == wip_stored,
                        ProductionActual.order_no == order_no,
                    )
                    .first()
                )
            else:
                pa = (
                    db.query(ProductionActual)
                    .filter(
                        ProductionActual.prod_date == d,
                        ProductionActual.work_center_id == wc.id,
                        ProductionActual.item_id == item.id,
                        ProductionActual.operation_seq == seq,
                        ProductionActual.order_no == order_no,
                    )
                    .first()
                )
            if not pa:
                pa = ProductionActual(
                    prod_date=d,
                    work_center_id=wc.id,
                    item_id=item.id,
                    operation_seq=seq,
                    order_no=order_no,
                    semi_finished_code=wip_stored,
                    quantity=qty,
                )
                db.add(pa)
                ins += 1
            else:
                pa.quantity = qty
                pa.work_center_id = wc.id
                pa.item_id = item.id
                pa.operation_seq = seq
                pa.semi_finished_code = wip_stored
                upd += 1
            pa.earned_hours = round(max(earned, 0.0), 4)
            pa.reported_hours = _float(r.get("reported_hours"), None)
        except Exception as e:  # noqa: BLE001
            errs.append(f"Satir {r['_row']}: {e}")
    return ins, upd, errs


def import_downtime(db: Session, rows: list[dict]) -> tuple[int, int, list[str]]:
    """Ayni gun icin is merkezinin duruslari yeniden yuklenirse eski satirlar silinir."""
    ins = 0
    errs = []
    wcs = _wc_lookup(db)
    cleared: set[tuple[int, date]] = set()
    for r in rows:
        try:
            d = _date(r.get("dt_date"))
            wc = wcs.get(_str(r.get("wc_code")).upper())
            if not wc:
                raise ValueError(f"Is merkezi bulunamadi: {r.get('wc_code')}")
            minutes = _float(r.get("minutes"))
            if minutes is None:
                raise ValueError("Sure bos")
            if (wc.id, d) not in cleared:
                db.query(Downtime).filter(Downtime.work_center_id == wc.id, Downtime.dt_date == d).delete(synchronize_session=False)
                cleared.add((wc.id, d))
            db.add(Downtime(dt_date=d, work_center_id=wc.id, reason_code=_str(r.get("reason_code")), reason_desc=_str(r.get("reason_desc")), minutes=minutes))
            ins += 1
        except Exception as e:  # noqa: BLE001
            errs.append(f"Satir {r['_row']}: {e}")
    return ins, 0, errs


def import_wc_weeks(db: Session, rows: list[dict]) -> tuple[int, int, list[str]]:
    """Is merkezi x hafta istisnalari; ayni hafta varsa guncellenir. Tum degerler bos ise kayit silinir."""
    ins = upd = 0
    errs = []
    wcs = _wc_lookup(db)
    for r in rows:
        try:
            wc = wcs.get(_str(r.get("wc_code")).upper())
            if not wc:
                raise ValueError(f"Is merkezi bulunamadi: {r.get('wc_code')}")
            wk = parse_week(r.get("week_start"))
            hc = _int(r.get("headcount"))
            eff = _float(r.get("efficient_hours_per_person"))
            days = _int(r.get("working_days"))
            if days is not None and not 0 <= days <= 7:
                raise ValueError("Calisma gunu 0-7 arasinda olmali")
            note = _str(r.get("note"))
            row = db.query(WorkCenterWeek).filter(WorkCenterWeek.work_center_id == wc.id, WorkCenterWeek.week_start == wk).first()
            if hc is None and eff is None and days is None and not note:
                if row:
                    db.delete(row)
                    upd += 1
                continue
            if row:
                upd += 1
            else:
                row = WorkCenterWeek(work_center_id=wc.id, week_start=wk)
                db.add(row)
                ins += 1
            row.headcount, row.efficient_hours_per_person, row.working_days, row.note = hc, eff, days, note
        except Exception as e:  # noqa: BLE001
            errs.append(f"Satir {r['_row']}: {e}")
    return ins, upd, errs


def import_op_rules(db: Session, rows: list[dict]) -> tuple[int, int, list[str]]:
    from app.services import scenarios as scen

    ins = upd = 0
    errs = []
    for r in rows:
        try:
            item_code = _str(r.get("item_code"))
            scope = "item" if item_code else "group"
            group = _str(r.get("product_group"))
            if scope == "group" and not group:
                raise ValueError("Urun grubu veya stok kodu gerekli")
            kind_raw = norm(r.get("rule"))
            rule = "cycles" if kind_raw in ("cevrim", "cycles", "cycle", "adet", "birlikte", "overlap") else "finish"
            existed = db.query(OpTransitionRule).count()
            scen.upsert_rule(
                db, scope, group, item_code or None, _str(r.get("from_op")), _str(r.get("to_op")),
                rule, _float(r.get("lag_cycles"), 0.0) or 0.0, _float(r.get("wait_minutes"), 0.0) or 0.0, _str(r.get("note")),
                _str(r.get("from_wip_code")), _str(r.get("to_wip_code")),
            )
            if db.query(OpTransitionRule).count() > existed:
                ins += 1
            else:
                upd += 1
        except Exception as e:  # noqa: BLE001
            errs.append(f"Satir {r['_row']}: {e}")
    return ins, upd, errs


def import_stock_receipts(db: Session, rows: list[dict]) -> tuple[int, int, list[str]]:
    ins = 0
    errs = []
    items = _item_lookup(db)
    for r in rows:
        try:
            d = _date(r.get("receipt_date"))
            code = _str(r.get("item_code")).upper()
            item = items.get(code)
            if not item:
                raise ValueError(f"Stok kodu bulunamadi: {code}")
            qty = _float(r.get("quantity"))
            if qty is None or qty <= 0:
                raise ValueError("Miktar pozitif olmali")
            db.add(StockReceipt(item_id=item.id, receipt_date=d, quantity=qty, lot=_str(r.get("lot")), note=_str(r.get("note")), source="import"))
            ins += 1
        except Exception as e:  # noqa: BLE001
            errs.append(f"Satir {r['_row']}: {e}")
    return ins, 0, errs


IMPORTERS: dict[str, Callable[[Session, list[dict]], tuple[int, int, list[str]]]] = {
    "workcenters": import_workcenters,
    "machines": import_machines,
    "shifts": import_shifts,
    "wc_weeks": import_wc_weeks,
    "employees": import_employees,
    "items": import_items,
    "bom": import_bom,
    "routing": import_routing,
    "op_rules": import_op_rules,
    "orders": import_orders,
    "production": import_production,
    "downtime": import_downtime,
    "stock_receipts": import_stock_receipts,
}


def run_import(db: Session, kind: str, content: bytes, filename: str, username: str, remove_missing: bool = False) -> ImportResult:
    if kind not in IMPORTERS:
        raise ValueError(f"Bilinmeyen import turu: {kind}")
    rows, errs = read_rows(content, kind)
    ins = upd = removed = 0
    if not errs:
        if kind == "orders":
            ins, upd, removed, errs = import_orders(db, rows, remove_missing=remove_missing)
        else:
            ins, upd, imp_errs = IMPORTERS[kind](db, rows)
            errs = imp_errs
    if kind == "production" and not errs:
        from app.services.stock import sync_progress_receipts

        sync_progress_receipts(db, username=username)
    db.add(ImportLog(kind=kind, filename=filename, username=username, inserted=ins, updated=upd, errors="\n".join(errs)[:10000]))
    db.commit()
    return ImportResult(kind=kind, inserted=ins, updated=upd, removed=removed, errors=errs)


# ---- Export ----

def _ws_from_rows(wb: Workbook, title: str, header: list[str], rows: list[list[Any]]) -> None:
    ws = wb.create_sheet(sheet_title(title))
    ws.append(header)
    for r in rows:
        ws.append([v.isoformat() if isinstance(v, (date, datetime)) else (v.strftime("%H:%M") if isinstance(v, time) else v) for v in r])
    _style_header(ws)
    _autosize(ws)


def workbook_bytes(wb: Workbook) -> bytes:
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
        del wb["Sheet"]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_backup(db: Session) -> bytes:
    """Tum tablolar tek Excel dosyasinda (her tablo bir sayfa). Sablon formatiyla uyumlu => geri yuklenebilir."""
    wb = Workbook()
    wc_code = {w.id: w.code for w in db.query(WorkCenter).all()}
    item_code = {i.id: i.code for i in db.query(Item).all()}

    _ws_from_rows(wb, "İş Merkezleri", [c[1] for c in TEMPLATES["workcenters"]["columns"]],
                  [[w.code, w.name, w.description, "E" if w.is_active else "H", "E" if w.is_planned else "H", w.capacity_unit_hours, w.default_efficient_hours, w.area_code, w.area_name, "Makine" if w.capacity_source == "machines" else "İM"] for w in db.query(WorkCenter).order_by(WorkCenter.code)])
    machine_code = {m.id: m.code for m in db.query(Machine).all()}
    _ws_from_rows(wb, "Makineler", [c[1] for c in TEMPLATES["machines"]["columns"]],
                  [[wc_code.get(m.work_center_id), m.code, m.name, m.description, "E" if m.is_active else "H"] for m in db.query(Machine).order_by(Machine.work_center_id, Machine.code)])
    _ws_from_rows(wb, "Vardiyalar", [c[1] for c in TEMPLATES["shifts"]["columns"]],
                  [[wc_code.get(s.work_center_id), s.name, s.weekdays, s.start_time, s.end_time, s.headcount, s.efficient_hours_per_person] for s in db.query(WorkCenterShift).order_by(WorkCenterShift.work_center_id, WorkCenterShift.id)])
    _ws_from_rows(wb, "Personel", [c[1] for c in TEMPLATES["employees"]["columns"]],
                  [[e.code, e.name, wc_code.get(e.work_center_id), machine_code.get(e.machine_id), "E" if e.is_active else "H"] for e in db.query(Employee).order_by(Employee.code)])
    _ws_from_rows(wb, "Stok Kodları", [c[1] for c in TEMPLATES["items"]["columns"]],
                  [[i.code, i.name, i.product_group, i.unit] for i in db.query(Item).order_by(Item.code)])
    _ws_from_rows(wb, "BOM", [c[1] for c in TEMPLATES["bom"]["columns"]],
                  [[item_code.get(b.item_id), b.component_code, b.component_name, b.quantity, b.unit] for b in db.query(BomLine).order_by(BomLine.item_id, BomLine.id)])
    _ws_from_rows(wb, "Rota", [c[1] for c in TEMPLATES["routing"]["columns"]],
                  [[item_code.get(o.item_id), o.seq, o.operation_name, wc_code.get(o.work_center_id), o.cycle_time_sec, o.setup_time_min, o.semi_finished_code] for o in db.query(RoutingOperation).order_by(RoutingOperation.item_id, RoutingOperation.seq)])
    _ws_from_rows(wb, "Siparişler", [c[1] for c in TEMPLATES["orders"]["columns"]] + ["Durum"],
                  [[o.order_no, o.position_no or "", o.customer, o.due_date, item_code.get(o.item_id), o.quantity, o.unit_price, o.status] for o in db.query(Order).order_by(Order.due_date, Order.order_no, Order.position_no)])
    _ws_from_rows(wb, "Plan", ["Hafta", "İş Merkezi Kodu", "Sipariş No", "Stok Kodu", "Operasyon Id", "Planlanan Saat", "Planlanan Miktar", "Mod", "Oluşturan"],
                  [[p.week_start, wc_code.get(p.work_center_id), p.order.order_no, item_code.get(p.order.item_id), p.operation_id, p.planned_hours, p.planned_qty, p.mode, p.created_by] for p in db.query(PlanLine).order_by(PlanLine.week_start, PlanLine.work_center_id)])
    _ws_from_rows(wb, "Günlük Üretim", [c[1] for c in TEMPLATES["production"]["columns"]] + ["Kazanılan Saat"],
                  [[p.prod_date, p.semi_finished_code or "", p.quantity, p.order_no, wc_code.get(p.work_center_id), item_code.get(p.item_id), p.operation_seq, p.reported_hours, p.earned_hours] for p in db.query(ProductionActual).order_by(ProductionActual.prod_date)])
    _ws_from_rows(wb, "Günlük Duruşlar", [c[1] for c in TEMPLATES["downtime"]["columns"]],
                  [[d.dt_date, wc_code.get(d.work_center_id), d.reason_code, d.reason_desc, d.minutes] for d in db.query(Downtime).order_by(Downtime.dt_date)])
    _ws_from_rows(wb, "Haftalık İş Gücü", [c[1] for c in TEMPLATES["wc_weeks"]["columns"]],
                  [[wc_code.get(w.work_center_id), w.week_start, w.headcount, w.efficient_hours_per_person, w.working_days, w.note] for w in db.query(WorkCenterWeek).order_by(WorkCenterWeek.work_center_id, WorkCenterWeek.week_start)])
    _ws_from_rows(wb, "Senaryo Kuralları", [c[1] for c in TEMPLATES["op_rules"]["columns"]],
                  [[r.product_group, item_code.get(r.item_id) if r.item_id else "", r.from_op, r.to_op, r.from_wip_code, r.to_wip_code, "Çevrim" if r.rule == "cycles" else "Bitiş", r.lag_cycles, r.wait_minutes, r.note] for r in db.query(OpTransitionRule).order_by(OpTransitionRule.product_group, OpTransitionRule.scope, OpTransitionRule.id)])
    _ws_from_rows(wb, "Depo Girişi", [c[1] for c in TEMPLATES["stock_receipts"]["columns"]],
                  [[s.receipt_date, item_code.get(s.item_id), s.quantity, s.lot, s.note] for s in db.query(StockReceipt).order_by(StockReceipt.receipt_date, StockReceipt.id)])
    order_no = {o.id: o.order_no for o in db.query(Order).all()}
    order_pos = {o.id: o.position_no or "" for o in db.query(Order).all()}
    _ws_from_rows(wb, "Rezervasyonlar", ["Stok Kodu", "Sipariş No", "Poz No", "Miktar", "Kaynak", "Not", "Oluşturan", "Tarih"],
                  [[item_code.get(r.item_id), order_no.get(r.order_id), order_pos.get(r.order_id), r.quantity, "manuel" if r.source == "manual" else "otomatik", r.note, r.created_by, r.created_at] for r in db.query(Reservation).order_by(Reservation.id)])
    _ws_from_rows(wb, "Sevkler", ["Tarih", "Stok Kodu", "Sipariş No", "Poz No", "Miktar", "Not", "Oluşturan"],
                  [[s.ship_date, item_code.get(s.item_id), order_no.get(s.order_id), order_pos.get(s.order_id), s.quantity, s.note, s.created_by] for s in db.query(Shipment).order_by(Shipment.ship_date, Shipment.id)])
    _ws_from_rows(wb, "Kullanıcılar", ["Kullanıcı", "Ad Soyad", "Rol", "Aktif"],
                  [[u.username, u.full_name, u.role, "E" if u.is_active else "H"] for u in db.query(User).order_by(User.username)])
    _ws_from_rows(wb, "Import Logu", ["Tarih", "Tür", "Dosya", "Kullanıcı", "Eklenen", "Güncellenen", "Hatalar"],
                  [[l.created_at, l.kind, l.filename, l.username, l.inserted, l.updated, l.errors] for l in db.query(ImportLog).order_by(ImportLog.created_at.desc()).limit(500)])
    return workbook_bytes(wb)


def build_report(sheets: dict[str, tuple[list[str], list[list[Any]]]]) -> bytes:
    """Genel rapor: {sayfa adi: (basliklar, satirlar)}."""
    wb = Workbook()
    for title, (header, rows) in sheets.items():
        _ws_from_rows(wb, title, header, rows)
    return workbook_bytes(wb)
