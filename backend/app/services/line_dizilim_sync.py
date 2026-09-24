"""Hat dizilimi eşitleme (saha tablosu): "Tavlama_Yikama_Onceliklendirme.xlsx".

  * "YIKAMA ÖNCELİK": stok kodu + yk2/yk3/yk4 sütunları = YK-02 / YK-03 / YK-04 makinelerinin konveyörüne tek seferde
    konan parça adedi (dizilim); boş hücre = ürün o makinede yıkanamaz. YK-05, YK-03 ile aynı kabul edilir (kullanıcı kararı).
  * "TAVLAMA ÖNCELİK": stok kodu + "Tavlama dizilim (adet)" = her iki fırın (TAV-06, TAV-07) için aynı dizilim.

Uygulama (saha tablosu geçerli): mamulün tüm yıkama/tavlama operasyonlarında (kendi rotası + montaj parçası rotaları)
istasyon bağları dosyadaki uygunlukla değiştirilir, her istasyon bağına kendi dizilimi yazılır, birincil istasyon en yüksek
dizilimli makine (eşitlikte YK-04 / TAV-06), operasyonun dizilimi (units_per_cycle) birincil makineninkidir.
Tekrarlı yıkamalar aynı dizilimi alır. Paylaşılan yarımamül rotasında farklı mamuller farklı değer verirse makine bazında en düşük dizilim alınır
(temkinli) ve çakışma raporlanır.
"""
from __future__ import annotations

import io
from collections import Counter, defaultdict
from dataclasses import dataclass, field

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import ImportLog, Item, Machine, Order, RoutingOperation, RoutingOperationStation
from app.services.bom_tree import flatten_fg_operations
from app.services.routing_resource import conveyor_kind, line_run_hours

WASH_SHEET = "YIKAMA ÖNCELİK"
ANNEAL_SHEET = "TAVLAMA ÖNCELİK"
WASH_COLS = {"yk2": "YK-02", "yk3": "YK-03", "yk4": "YK-04"}
WASH_ALIAS = {"YK-05": "YK-03"}  # YK-05 = YK-03 gibi
ANNEAL_MACHINES = ["TAV-06", "TAV-07"]
WASH_PRIMARY_ORDER = ["YK-04", "YK-02", "YK-03", "YK-05"]
NOTE = "Hat dizilimi (saha tablosu)"


def _norm(s) -> str:
    return "".join(ch for ch in str(s or "").upper().replace("İ", "I") if ch.isalnum())


def _sheet_rows(wb, wanted: str):
    key = _norm(wanted)
    ws = next((w for w in wb.worksheets if _norm(w.title) == key), None)
    if ws is None:
        return None, []
    rows = list(ws.iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(rows) if r and sum(1 for v in r if v not in (None, "")) >= 5), None)
    if hi is None:
        return [], []
    hdr = [str(v) if v is not None else "" for v in rows[hi]]
    return hdr, [r for r in rows[hi + 1:] if r and r[0] not in (None, "")]


def _int(v) -> int | None:
    if v in (None, ""):
        return None
    try:
        n = int(float(v))
    except (TypeError, ValueError):
        return None
    return n if n > 0 else None


@dataclass
class DizilimFile:
    washing: dict[str, dict[str, int]] = field(default_factory=dict)  # stok kodu -> {makine kodu: dizilim}
    annealing: dict[str, int] = field(default_factory=dict)  # stok kodu -> dizilim
    missing_sheets: list[str] = field(default_factory=list)
    bad_rows: list[str] = field(default_factory=list)


def parse_file(content: bytes) -> DizilimFile:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    out = DizilimFile()
    hdr, rows = _sheet_rows(wb, WASH_SHEET)
    if hdr is None:
        out.missing_sheets.append(WASH_SHEET)
    else:
        idx = {c: next((i for i, h in enumerate(hdr) if _norm(h) == _norm(c)), None) for c in WASH_COLS}
        if any(i is None for i in idx.values()):
            out.bad_rows.append(f"{WASH_SHEET}: yk2/yk3/yk4 sütunları bulunamadı")
        else:
            for r in rows:
                code = str(r[0]).strip()
                d = {}
                for col, mcode in WASH_COLS.items():
                    n = _int(r[idx[col]])
                    if n:
                        d[mcode] = n
                if not d:
                    out.bad_rows.append(f"{code}: hiçbir yıkama makinesi için dizilim yok")
                    continue
                if "YK-03" in d:
                    d["YK-05"] = d["YK-03"]
                out.washing[code] = d
    hdr, rows = _sheet_rows(wb, ANNEAL_SHEET)
    if hdr is None:
        out.missing_sheets.append(ANNEAL_SHEET)
    else:
        di = next((i for i, h in enumerate(hdr) if "DIZILIM" in _norm(h)), None)
        if di is None:
            out.bad_rows.append(f"{ANNEAL_SHEET}: dizilim sütunu bulunamadı")
        else:
            for r in rows:
                n = _int(r[di])
                if n:
                    out.annealing[str(r[0]).strip()] = n
    wb.close()
    return out


def _machines(db: Session) -> dict[str, Machine]:
    return {m.code.upper(): m for m in db.query(Machine).all()}


def _plan(db: Session, f: DizilimFile) -> dict:
    """Uygulanacak değişiklikleri hesaplar: op_id -> {"kind", "stations": {machine_id: units}, "primary": machine_id, "units": int}."""
    items = {i.code.upper(): i for i in db.query(Item).all()}
    machines = _machines(db)
    missing_machines = sorted({m for d in f.washing.values() for m in d if m not in machines} | {m for m in ANNEAL_MACHINES if m not in machines})
    changes: dict[int, dict] = {}
    conflicts: list[str] = []
    not_found: list[str] = []
    no_ops: list[str] = []
    seen_src: dict[int, str] = {}
    for code, d in list(f.washing.items()) + [(c, {"__anneal__": n}) for c, n in f.annealing.items()]:
        it = items.get(code.upper())
        if it is None:
            not_found.append(code)
            continue
        kind = "annealing" if "__anneal__" in d else "washing"
        if kind == "annealing":
            st = {m: d["__anneal__"] for m in ANNEAL_MACHINES if m in machines}
            order = ANNEAL_MACHINES
        else:
            st = {m: n for m, n in d.items() if m in machines}
            order = WASH_PRIMARY_ORDER
        if not st:
            continue
        ops = [fl.operation for fl in flatten_fg_operations(db, it) if conveyor_kind(fl.operation) == kind]
        if not ops:
            no_ops.append(f"{code} ({'tavlama' if kind == 'annealing' else 'yıkama'})")
            continue
        best = max(st.values())
        primary = next(m for m in order if m in st and st[m] == best)
        for op in ops:
            spec = {"kind": kind, "stations": {machines[m].id: n for m, n in st.items()}, "primary": machines[primary].id, "units": best, "src": code}
            prev = changes.get(op.id)
            if prev and prev["stations"] != spec["stations"]:
                # Paylaşılan yarımamül rotası: farklı mamuller farklı değer verdiyse makine bazında EN DÜŞÜK dizilim
                # (temkinli kapasite); uygunluk kümesi kesişim değil birleşim (bir mamulde yıkanabiliyorsa makine uygundur).
                merged = dict(prev["stations"])
                for mid, n in spec["stations"].items():
                    merged[mid] = min(merged[mid], n) if mid in merged else n
                m_best = max(merged.values())
                order_ids = [machines[m].id for m in order if m in machines]
                m_primary = next(mid for mid in order_ids if mid in merged and merged[mid] == m_best)
                conflicts.append(f"op #{op.id} ({op.operation_name}): {seen_src[op.id]} ile {code} farklı dizilim verdi; en düşük değer alındı")
                spec = {**spec, "stations": merged, "primary": m_primary, "units": m_best}
            changes[op.id] = spec
            seen_src[op.id] = code
    return {"changes": changes, "conflicts": conflicts, "not_found": not_found, "no_ops": no_ops, "missing_machines": missing_machines}


def _current(db: Session, op_ids: list[int]) -> tuple[dict[int, RoutingOperation], dict[int, dict[int, int | None]], dict[int, bool]]:
    ops = {o.id: o for o in db.query(RoutingOperation).filter(RoutingOperation.id.in_(op_ids)).all()} if op_ids else {}
    links: dict[int, dict[int, int | None]] = defaultdict(dict)
    prim: dict[int, bool] = {}
    for st in db.query(RoutingOperationStation).filter(RoutingOperationStation.operation_id.in_(op_ids)).all() if op_ids else []:
        links[st.operation_id][st.machine_id] = getattr(st, "units_per_cycle", None)
    return ops, links, prim


def _open_order_hours(db: Session, ops: dict[int, RoutingOperation], units_of) -> dict[str, float]:
    """Açık siparişlerin yıkama/tavlama hat-saati (birincil makine dizilimiyle)."""
    from app.services.bom_tree import flatten_fg_operations as _flat
    hours: dict[str, float] = defaultdict(float)
    for o in db.query(Order).filter(Order.status == "open").all():
        if not o.item:
            continue
        for fl in _flat(db, o.item):
            op = fl.operation
            k = conveyor_kind(op)
            if not k:
                continue
            hours[k] += line_run_hours(op, float(o.quantity or 0), units=units_of(op))
    return dict(hours)


def preview(db: Session, content: bytes) -> dict:
    f = parse_file(content)
    p = _plan(db, f)
    changes = p["changes"]
    ops, links, _ = _current(db, list(changes))
    link_changed = units_changed = primary_changed = 0
    for op_id, spec in changes.items():
        op = ops[op_id]
        cur = links.get(op_id, {})
        if set(cur) != set(spec["stations"]) or any(cur.get(m) != n for m, n in spec["stations"].items()):
            link_changed += 1
        if int(op.units_per_cycle or 1) != spec["units"]:
            units_changed += 1
        if op.primary_machine_id != spec["primary"]:
            primary_changed += 1
    before = _open_order_hours(db, ops, lambda op: None)
    after = _open_order_hours(db, ops, lambda op: changes[op.id]["units"] if op.id in changes else None)
    by_kind = Counter(spec["kind"] for spec in changes.values())
    return {
        "missing_sheets": f.missing_sheets, "bad_rows": f.bad_rows[:30], "bad_row_count": len(f.bad_rows),
        "file": {"washing_items": len(f.washing), "annealing_items": len(f.annealing),
                 "washing_combos": [{"machines": " + ".join(sorted(k)), "items": v} for k, v in sorted(Counter(tuple(sorted(m for m in d if m != "YK-05")) for d in f.washing.values()).items(), key=lambda x: -x[1])]},
        "operations": {"washing": by_kind.get("washing", 0), "annealing": by_kind.get("annealing", 0), "station_links_changed": link_changed,
                       "units_changed": units_changed, "primary_changed": primary_changed},
        "not_found": p["not_found"][:50], "not_found_count": len(p["not_found"]), "no_ops": p["no_ops"][:50], "no_ops_count": len(p["no_ops"]),
        "conflicts": p["conflicts"][:50], "conflict_count": len(p["conflicts"]), "missing_machines": p["missing_machines"],
        "open_order_hours": {"before": {k: round(v, 1) for k, v in before.items()}, "after": {k: round(v, 1) for k, v in after.items()}},
    }


def apply(db: Session, content: bytes, *, username: str, filename: str) -> dict:
    f = parse_file(content)
    p = _plan(db, f)
    changes = p["changes"]
    if p["missing_machines"]:
        raise ValueError("Programda tanımlı olmayan makine: " + ", ".join(p["missing_machines"]))
    op_ids = list(changes)
    if op_ids:
        db.query(RoutingOperationStation).filter(RoutingOperationStation.operation_id.in_(op_ids)).delete(synchronize_session=False)
    ops = {o.id: o for o in db.query(RoutingOperation).filter(RoutingOperation.id.in_(op_ids)).all()} if op_ids else {}
    for op_id, spec in changes.items():
        op = ops[op_id]
        for mid, n in spec["stations"].items():
            db.add(RoutingOperationStation(operation_id=op_id, machine_id=mid, is_primary=(mid == spec["primary"]), units_per_cycle=n))
        op.primary_machine_id = spec["primary"]
        op.units_per_cycle = spec["units"]
    by_kind = Counter(spec["kind"] for spec in changes.values())
    db.add(ImportLog(kind="line_dizilim", filename=f"{filename} ({NOTE})", username=username, inserted=len(changes), updated=0,
                     errors="", warnings="\n".join(p["conflicts"] + [f"Stok kartı yok: {c}" for c in p["not_found"]] + [f"Operasyon yok: {c}" for c in p["no_ops"]])[:10000]))
    db.commit()
    return {"operations": len(changes), "washing": by_kind.get("washing", 0), "annealing": by_kind.get("annealing", 0),
            "conflict_count": len(p["conflicts"]), "not_found_count": len(p["not_found"]), "no_ops_count": len(p["no_ops"])}
