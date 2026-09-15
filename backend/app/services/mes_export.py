from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from app.services.excel import _export_cell


def export_report(r):
    wb = Workbook()
    wb.remove(wb.active)
    sheets = [
        ("Bilgi", ["Alan", "Değer"], [["Hafta", str(r["week"])], ["Rapor tarihi", str(r["as_of"])],
          ["Referans", r["baseline"]], *[["Hesaplama", n] for n in r["notes"]]]),
        ("Günlük", ["Tarih", "Üretim standart saat", "Planla eşleşen saat", "Plan dışı saat", "Kümülatif saat"],
         [[d[k] for k in ("day", "hours", "matched_hours", "off_plan_hours", "cumulative_hours")] for d in r["daily"] if d["reported"]]),
        ("Operasyonlar", ["Malzeme", "Operasyon", "Bitmiş ürün adayları", "Plan adet", "Bu hafta adet", "Önceden adet", "Kalan adet", "Kalan standart saat"],
         [[o["material_code"], o["operation_name"], ", ".join(o["products"]), o["quantity"], o["actual_qty"], o["early_qty"], o["remaining_qty"], o["remaining_hours"]] for o in r["operations"]]),
        ("Plan dışı", ["Detay ID", "Tarih", "Malzeme", "Net adet", "Standart saat", "Eşleşen hafta"],
         [[d["detail_id"], d["prod_date"], d["material_code"], d["quantity"], d["standard_hours"], d["match_week"]] for d in r["off_plan"]]),
        ("Havuz dağıtımı", ["Bitiş haftası", "Bitmiş ürün", "Yarımamül", "Ayrılan adet"],
         [[a[k] for k in ("finish_week", "finished_item_code", "material_code", "quantity")] for a in r["allocations"]]),
        ("Kontroller", ["Detay ID", "Tarih", "Malzeme", "Net adet", "Sorun"],
         [[d["detail_id"], d["prod_date"], d["material_code"], d["quantity"], d["mapping"]["reason"]] for d in r["unresolved"]]),
        ("Havuz", ["Yarımamül", "Net bakiye"], [[p["material_code"], p["quantity"]] for p in r["pool"]]),
    ]
    for name, headers, rows in sheets:
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append([_export_cell(v) for v in row])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="167C80")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(18, len(headers[i-1]) + 4))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
