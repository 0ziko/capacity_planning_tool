"""ERP Production BOM (RECETELER) import: ters rota, istasyon, yari mamul agaci."""

from __future__ import annotations

import io
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.models import BomLine, Item, Machine, RoutingOperation, RoutingOperationStation
from app.services.stations import ensure_extra_work_centers, get_or_create_wc_by_name, wc_lookup_by_name

GOLDEN_FGS = ("6000006", "6012181", "6000087", "6004761", "6000705")


def S(x) -> str:
    if x is None:
        return ""
    return str(x).strip()


def code_str(x) -> str:
    if x is None:
        return ""
    if isinstance(x, int):
        return str(x)
    if isinstance(x, float) and x == int(x):
        return str(int(x))
    s = str(x).strip()
    if s.endswith(".0"):
        try:
            return str(int(float(s)))
        except ValueError:
            pass
    return s


def base_wip(stok: str) -> str:
    return stok.split("-")[0] if "-" in stok else stok


def sure_dk(v) -> float | None:
    try:
        if v is None or S(v) in ("", "None"):
            return None
        return round(float(v), 3)
    except (TypeError, ValueError):
        return None


def cycle_time_sec(sure_dk_val: float | None) -> float:
    if not sure_dk_val or sure_dk_val <= 0:
        return 0.0
    return round(sure_dk_val * 60.0 * get_settings().bom_cycle_factor, 2)


def is_wip_code(code: str) -> bool:
    return bool(re.match(r"^5\d{5,}$", code))


def is_fg_code(code: str) -> bool:
    return bool(re.match(r"^6\d{5,}$", code))


@dataclass
class MaterialRow:
    level: str
    code: str
    name: str
    qty: float


@dataclass
class ParsedOp:
    listing_sira: int
    wip: str
    wip_op_code: str
    wip_name: str
    op_code: str
    op_name: str
    sure_dk: float | None
    station: str
    alt_station: str
    wc_name: str
    materials: list[MaterialRow] = field(default_factory=list)


@dataclass
class Branch:
    wip: str
    instance: int
    ops: list[ParsedOp]


@dataclass
class ParsedFG:
    fg: str
    name: str
    branches: list[Branch]
    finish_wip: str


@dataclass
class ImportProductionBomResult:
    fg_count: int = 0
    items_created: int = 0
    items_updated: int = 0
    routes_written: int = 0
    bom_lines: int = 0
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    ambiguous_parents: list[str] = field(default_factory=list)


def _parse_fg_records(fg: str, recs: list[dict], header: list[str]) -> ParsedFG:
    k_level, _, k_sira, k_stok, k_ad, k_qty, k_op, k_opis, k_sure = header[:9]
    k_ist, k_alt, _, _, k_wc = header[9:14]

    name = ""
    steps: list[ParsedOp] = []
    pending_mats: list[MaterialRow] = []
    current: ParsedOp | None = None

    for r in sorted(recs, key=lambda x: int(x.get(k_sira) or 0)):
        st = code_str(r.get(k_stok))
        on = S(r.get(k_opis))
        lv = code_str(r.get(k_level))
        if lv == "6" or st == fg:
            name = S(r.get(k_ad)) or name
            continue
        if on:
            if current:
                current.materials = pending_mats
                steps.append(current)
                pending_mats = []
            ist = S(r.get(k_ist))
            alt = S(r.get(k_alt))
            wc = S(r.get(k_wc))
            current = ParsedOp(
                listing_sira=int(r.get(k_sira) or 0),
                wip=base_wip(st),
                wip_op_code=st,
                wip_name=S(r.get(k_ad)),
                op_code=code_str(r.get(k_op)),
                op_name=on,
                sure_dk=sure_dk(r.get(k_sure)),
                station="" if ist in ("0", "None") else ist,
                alt_station="" if alt in ("0", "None") else alt,
                wc_name="" if wc in ("0", "None") else wc,
            )
        else:
            qty_raw = r.get(k_qty)
            try:
                qty = float(qty_raw) if qty_raw not in (None, "") else 1.0
            except (TypeError, ValueError):
                qty = 1.0
            mat = MaterialRow(level=lv, code=st, name=S(r.get(k_ad)), qty=qty)
            if current:
                pending_mats.append(mat)
            else:
                pending_mats.append(mat)
    if current:
        current.materials = pending_mats
        steps.append(current)

    runs: list[list[ParsedOp]] = []
    for s in steps:
        if not runs or runs[-1][-1].wip != s.wip:
            runs.append([s])
        else:
            runs[-1].append(s)

    branches: list[Branch] = []
    seen: dict[str, int] = defaultdict(int)
    for ops in runs:
        if not ops:
            continue
        mfg = list(reversed(ops))
        # Montaja giden kod Excel'deki son operasyon ciktisi; base_wip() sadece dal gruplama icin.
        output = mfg[-1].wip_op_code or ops[0].wip
        seen[output] += 1
        branches.append(Branch(wip=output, instance=seen[output], ops=mfg))

    finish_wip = ""
    if branches:
        finish_wip = min(branches, key=lambda b: b.ops[0].listing_sira if b.ops else 999999).wip

    return ParsedFG(fg=fg, name=name, branches=branches, finish_wip=finish_wip)


def load_bom_grouped(content: bytes | None = None, path: Path | None = None) -> tuple[dict[str, list[dict]], list[str]]:
    if content is not None:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    elif path is not None:
        wb = load_workbook(path, data_only=True, read_only=True)
    else:
        raise ValueError("content veya path gerekli")
    ws = wb[wb.sheetnames[0]]
    header: list[str] | None = None
    by_fg: dict[str, list[dict]] = defaultdict(list)
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        vals = list(row[:14])
        if i == 1:
            header = [S(x) or f"COL{j}" for j, x in enumerate(vals)]
            continue
        if not header:
            continue
        rec = {header[j]: vals[j] if j < len(vals) else None for j in range(len(header))}
        fg = code_str(rec.get(header[1]))
        if fg and is_fg_code(fg):
            by_fg[fg].append(rec)
    wb.close()
    return by_fg, header or []


def _get_or_create_item(
    db: Session,
    cache: dict[str, Item],
    code: str,
    name: str = "",
    *,
    counters: dict[str, int],
) -> Item:
    key = code.upper()
    it = cache.get(key)
    if it:
        if name and not it.name:
            it.name = name
        return it
    it = db.query(Item).filter(Item.code.ilike(code)).first()
    if it:
        cache[key] = it
        counters["updated"] += 1
        if name and (not it.name or it.name == it.code):
            it.name = name
        return it
    it = Item(code=code, name=name or code, product_group="ERP", unit="AD")
    db.add(it)
    db.flush()
    cache[key] = it
    counters["created"] += 1
    return it


def _clear_item_master(db: Session, item_id: int) -> None:
    db.query(RoutingOperation).filter(RoutingOperation.item_id == item_id).delete(synchronize_session=False)
    db.query(BomLine).filter(BomLine.item_id == item_id).delete(synchronize_session=False)


def _attach_stations(
    db: Session,
    ro: RoutingOperation,
    primary: Machine | None,
    alt_code: str,
    machines: dict[str, Machine],
    warnings: list[str],
    ctx: str,
) -> None:
    if primary:
        db.add(RoutingOperationStation(operation_id=ro.id, machine_id=primary.id, is_primary=True))
    if alt_code:
        alt = machines.get(alt_code.upper())
        if alt and (not primary or alt.id != primary.id):
            db.add(RoutingOperationStation(operation_id=ro.id, machine_id=alt.id, is_primary=False))
        elif alt_code and not alt:
            warnings.append(f"{ctx}: alternatif istasyon yok {alt_code}")


def _resolve_wc(
    db: Session,
    wc_idx: dict,
    wc_name: str,
    machine: Machine | None,
) -> int:
    if wc_name:
        return get_or_create_wc_by_name(db, wc_name, wc_idx).id
    if machine:
        return machine.work_center_id
    return get_or_create_wc_by_name(db, "MONTAJ", wc_idx).id


def _write_route(
    db: Session,
    item: Item,
    branches: list[Branch],
    finish_wip: str,
    *,
    wc_idx: dict,
    machines: dict[str, Machine],
    warnings: list[str],
) -> int:
    _clear_item_master(db, item.id)
    seq = 10
    written = 0
    finish_branch = next((b for b in branches if b.wip == finish_wip), None)
    if not finish_branch:
        return 0
    for op in finish_branch.ops:
        wc_name = op.wc_name
        primary = machines.get(op.station.upper()) if op.station else None
        if not wc_name and primary:
            wc_name = primary.work_center.name if primary.work_center else ""
        wc_id = _resolve_wc(db, wc_idx, wc_name, primary)
        ro = RoutingOperation(
            item_id=item.id,
            seq=seq,
            operation_name=op.op_name,
            work_center_id=wc_id,
            cycle_time_sec=cycle_time_sec(op.sure_dk),
            semi_finished_code=op.wip_op_code or op.wip,
            primary_machine_id=primary.id if primary else None,
        )
        if op.sure_dk is None or op.sure_dk <= 0:
            warnings.append(f"{item.code} {op.op_name}: SURE bos/0")
        db.add(ro)
        db.flush()
        _attach_stations(db, ro, primary, op.alt_station, machines, warnings, f"{item.code} {op.op_name}")
        seq += 1
        written += 1
    return written


def _branch_listing_sira(branch: Branch) -> int:
    return max((o.listing_sira for o in branch.ops), default=0)


def _write_wip_item(
    db: Session,
    branch: Branch,
    cache: dict[str, Item],
    counters: dict[str, int],
    wc_idx: dict,
    machines: dict[str, Machine],
    warnings: list[str],
) -> Item:
    wip = branch.wip
    name = branch.ops[-1].wip_name.split("-")[0].strip() if branch.ops else wip
    item = _get_or_create_item(db, cache, wip, name, counters=counters)
    _clear_item_master(db, item.id)
    branch_mats: dict[tuple[str, str], tuple[str, float, int]] = {}
    branch_sira = _branch_listing_sira(branch)
    seq = 10
    for mfg_idx, op in enumerate(branch.ops):
        primary = machines.get(op.station.upper()) if op.station else None
        wc_name = op.wc_name
        if not wc_name and primary:
            wc_name = primary.work_center.name if primary.work_center else ""
        wc_id = _resolve_wc(db, wc_idx, wc_name, primary)
        ro = RoutingOperation(
            item_id=item.id,
            seq=seq,
            operation_name=op.op_name,
            work_center_id=wc_id,
            cycle_time_sec=cycle_time_sec(op.sure_dk),
            semi_finished_code=op.wip_op_code or wip,
            primary_machine_id=primary.id if primary else None,
        )
        if op.sure_dk is None or op.sure_dk <= 0:
            warnings.append(f"{wip} {op.op_name}: SURE bos/0")
        db.add(ro)
        db.flush()
        _attach_stations(db, ro, primary, op.alt_station, machines, warnings, f"{wip} {op.op_name}")
        seq += 10
        mat_seq = (mfg_idx + 1) * 10 + 1
        for m in op.materials:
            if not m.code:
                continue
            k = (m.code.upper(), branch.wip)
            prev = branch_mats.get(k)
            prev_seq = prev[2] if prev else mat_seq
            branch_mats[k] = (
                m.name or (prev[0] if prev else m.name),
                (prev[1] if prev else 0) + m.qty,
                min(prev_seq, mat_seq),
            )
    for (comp, src), (cname, qty, rseq) in branch_mats.items():
        _get_or_create_item(db, cache, comp, cname, counters=counters)
        db.add(
            BomLine(
                item_id=item.id,
                component_code=comp,
                component_name=cname,
                quantity=qty,
                unit="AD",
                source_wip=src,
                branch_listing_sira=branch_sira,
                recipe_seq=rseq,
            )
        )
        counters["bom"] += 1
    return item


def _write_fg_bom(
    db: Session,
    fg_item: Item,
    branches: list[Branch],
    finish_wip: str,
    cache: dict[str, Item],
    counters: dict[str, int],
) -> None:
    wip_info: dict[str, dict[str, float | int]] = {}
    for b in branches:
        if b.wip == finish_wip:
            continue
        max_sira = max((o.listing_sira for o in b.ops), default=0)
        info = wip_info.setdefault(b.wip, {"qty": 0, "listing_sira": 0})
        info["qty"] = int(info["qty"]) + 1
        info["listing_sira"] = max(int(info["listing_sira"]), max_sira)
    for wip, info in sorted(wip_info.items(), key=lambda x: int(x[1]["listing_sira"]), reverse=True):
        _get_or_create_item(db, cache, wip, "", counters=counters)
        db.add(
            BomLine(
                item_id=fg_item.id,
                component_code=wip,
                component_name=cache[wip.upper()].name,
                quantity=float(info["qty"]),
                unit="AD",
                source_wip=wip,
                branch_listing_sira=int(info["listing_sira"]),
                recipe_seq=0,
            )
        )
        counters["bom"] += 1


def _write_fg_recipe_detail(
    db: Session,
    fg_item: Item,
    branches: list[Branch],
    cache: dict[str, Item],
    counters: dict[str, int],
) -> None:
    """FG BOM'a tum dallarin hammadde, alt yari mamul ve operasyon adim satirlarini yazar."""
    recipe_agg: dict[tuple[str, str], tuple[str, float, int, int]] = {}

    def _add(comp: str, src: str, name: str, qty: float, branch_sira: int, recipe_seq: int) -> None:
        k = (comp.upper(), src)
        prev = recipe_agg.get(k)
        if prev:
            recipe_agg[k] = (
                name or prev[0],
                prev[1] + qty,
                branch_sira,
                min(prev[3], recipe_seq),
            )
        else:
            recipe_agg[k] = (name, qty, branch_sira, recipe_seq)

    for branch in branches:
        src = branch.wip
        branch_sira = _branch_listing_sira(branch)
        for mfg_idx, op in enumerate(branch.ops):
            step_seq = (mfg_idx + 1) * 10
            mat_seq = step_seq + 1
            # Son op kodu montaj baglantisi ile ayni (item, code, source_wip); adim olarak tekrar yazma.
            if op.wip_op_code and op.wip_op_code != branch.wip:
                _add(op.wip_op_code, src, op.wip_name or op.op_name, 1.0, branch_sira, step_seq)
            for m in op.materials:
                if m.code:
                    _add(m.code, src, m.name, m.qty, branch_sira, mat_seq)

    for (comp, src), (cname, qty, branch_sira, recipe_seq) in recipe_agg.items():
        _get_or_create_item(db, cache, comp, cname, counters=counters)
        db.add(
            BomLine(
                item_id=fg_item.id,
                component_code=comp,
                component_name=cname,
                quantity=qty,
                unit="AD",
                source_wip=src,
                branch_listing_sira=branch_sira,
                recipe_seq=recipe_seq,
            )
        )
        counters["bom"] += 1


def import_parsed_fg(
    db: Session,
    parsed: ParsedFG,
    *,
    cache: dict[str, Item],
    wc_idx: dict,
    machines: dict[str, Machine],
    counters: dict[str, int],
    warnings: list[str],
) -> None:
    fg_item = _get_or_create_item(db, cache, parsed.fg, parsed.name, counters=counters)
    _clear_item_master(db, fg_item.id)

    for branch in parsed.branches:
        if branch.wip != parsed.finish_wip:
            _write_wip_item(db, branch, cache, counters, wc_idx, machines, warnings)

    counters["routes"] += _write_route(db, fg_item, parsed.branches, parsed.finish_wip, wc_idx=wc_idx, machines=machines, warnings=warnings)
    _write_fg_bom(db, fg_item, parsed.branches, parsed.finish_wip, cache, counters)
    _write_fg_recipe_detail(db, fg_item, parsed.branches, cache, counters)


def run_production_bom_import(
    db: Session,
    *,
    content: bytes | None = None,
    path: Path | None = None,
    fg_filter: set[str] | None = None,
) -> ImportProductionBomResult:
    result = ImportProductionBomResult()
    by_fg, header = load_bom_grouped(content=content, path=path)
    if not header:
        result.errors.append("BOM baslik satiri okunamadi")
        return result

    wc_idx = wc_lookup_by_name(db)
    ensure_extra_work_centers(db, wc_idx)
    machines = {m.code.upper(): m for m in db.query(Machine).filter(Machine.is_active.is_(True)).all()}
    cache: dict[str, Item] = {i.code.upper(): i for i in db.query(Item).all()}
    counters = {"created": 0, "updated": 0, "routes": 0, "bom": 0}

    targets = sorted(by_fg.keys())
    if fg_filter:
        targets = [f for f in targets if f in fg_filter]

    for fg in targets:
        try:
            parsed = _parse_fg_records(fg, by_fg[fg], header)
            if not parsed.branches:
                result.warnings.append(f"{fg}: dal/operasyon yok")
                continue
            import_parsed_fg(db, parsed, cache=cache, wc_idx=wc_idx, machines=machines, counters=counters, warnings=result.warnings)
            result.fg_count += 1
        except Exception as e:  # noqa: BLE001
            result.errors.append(f"{fg}: {e}")

    result.items_created = counters["created"]
    result.items_updated = counters["updated"]
    result.routes_written = counters["routes"]
    result.bom_lines = counters["bom"]
    db.flush()
    return result


def import_production_bom_excel(db: Session, content: bytes) -> tuple[int, int, list[str]]:
    """Excel import handler: inserted=FG sayisi, updated=diger sayaclar toplami."""
    r = run_production_bom_import(db, content=content)
    errs = r.errors + [f"UYARI: {w}" for w in r.warnings[:200]]
    if len(r.warnings) > 200:
        errs.append(f"UYARI: ... ve {len(r.warnings) - 200} uyari daha")
    upd = r.items_created + r.items_updated + r.routes_written + r.bom_lines
    return r.fg_count, upd, errs
