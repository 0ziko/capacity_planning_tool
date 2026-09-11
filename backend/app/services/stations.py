"""Istasyon (makine) import: Is Merkezi adi ile eslestirme, eksik IM olusturma."""

from __future__ import annotations

import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import Employee, Machine, WorkCenter
from app.services.excel import norm

TR_MAP = str.maketrans("çğıöşüÇĞİÖŞÜ", "cgiosuCGIOSU")

# BOM'da gorunen ama istasyon listesinde olmayan is merkezleri
EXTRA_WORK_CENTER_NAMES = ("MONTAJ", "PAKETLEME")


def norm_wc_key(name: str) -> str:
    s = str(name or "").strip().translate(TR_MAP)
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", s.lower())


def wc_code_from_name(name: str) -> str:
    s = str(name or "").strip().translate(TR_MAP)
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    s = re.sub(r"[^A-Za-z0-9]+", " ", s).strip().upper()
    return s[:32] if s else "WC"


def wc_lookup_by_name(db: Session) -> dict[str, WorkCenter]:
    idx: dict[str, WorkCenter] = {}
    for wc in db.query(WorkCenter).all():
        for key in (norm_wc_key(wc.name), norm(wc.name), norm(wc_key(wc.code))):
            if key:
                idx[key] = wc
    return idx


def wc_key(code: str) -> str:
    return norm(code)


def get_or_create_wc_by_name(db: Session, name: str, idx: dict[str, WorkCenter]) -> WorkCenter:
    raw = str(name or "").strip()
    if not raw:
        raise ValueError("Is merkezi adi bos")
    key = norm_wc_key(raw)
    wc = idx.get(key)
    if wc:
        return wc
    code = wc_code_from_name(raw)
    existing = db.query(WorkCenter).filter(WorkCenter.code.ilike(code)).first()
    if existing:
        idx[key] = existing
        return existing
    wc = WorkCenter(
        code=code,
        name=raw,
        is_active=True,
        is_planned=True,
    )
    db.add(wc)
    db.flush()
    idx[key] = wc
    idx[norm(wc.code)] = wc
    return wc


def ensure_extra_work_centers(db: Session, idx: dict[str, WorkCenter]) -> None:
    for name in EXTRA_WORK_CENTER_NAMES:
        get_or_create_wc_by_name(db, name, idx)


def import_stations(
    db: Session,
    rows: list[dict],
    *,
    replace_missing: bool = True,
) -> tuple[int, int, int, list[str]]:
    """Istasyon satirlarini Machine olarak yukler. replace_missing=True ise listede olmayan makineler pasif yapilir."""
    ins = upd = deactivated = 0
    errs: list[str] = []
    idx = wc_lookup_by_name(db)
    ensure_extra_work_centers(db, idx)
    existing = {m.code.upper(): m for m in db.query(Machine).all()}
    file_codes: set[str] = set()

    for r in rows:
        try:
            code = str(r.get("station_code") or r.get("code") or "").strip()
            if not code:
                raise ValueError("Istasyon kodu bos")
            wc_name = str(r.get("wc_name") or r.get("work_center_name") or "").strip()
            if not wc_name:
                raise ValueError("Is merkezi adi bos")
            wc = get_or_create_wc_by_name(db, wc_name, idx)
            name = str(r.get("station_name") or r.get("name") or code).strip()
            file_codes.add(code.upper())
            m = existing.get(code.upper())
            if not m:
                m = Machine(code=code, work_center_id=wc.id, name=name, is_active=True)
                db.add(m)
                existing[code.upper()] = m
                ins += 1
            else:
                m.work_center_id = wc.id
                m.name = name or m.name
                m.is_active = True
                upd += 1
        except Exception as e:  # noqa: BLE001
            errs.append(f"Satir {r.get('_row', '?')}: {e}")

    if replace_missing and file_codes:
        for m in db.query(Machine).all():
            if m.code.upper() not in file_codes and m.is_active:
                m.is_active = False
                for emp in db.query(Employee).filter(Employee.machine_id == m.id).all():
                    emp.machine_id = None
                deactivated += 1

    db.flush()
    return ins, upd, deactivated, errs


def load_stations_xlsx(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    header = None
    rows: list[dict] = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        vals = list(row[:5])
        if i == 1:
            header = [str(x).strip() if x is not None else f"COL{j}" for j, x in enumerate(vals)]
            continue
        if not any(v is not None and str(v).strip() for v in vals):
            continue
        rec = {header[j]: vals[j] if j < len(vals) else None for j in range(len(header))}
        rec["_row"] = i
        # map Turkish headers
        mapped: dict = {"_row": i}
        for h, v in rec.items():
            hn = norm(h)
            if "istasyon" in hn and "kod" in hn:
                mapped["station_code"] = v
            elif "istasyon" in hn and ("tanim" in hn or "ad" in hn):
                mapped["station_name"] = v
            elif "merkez" in hn:
                mapped["wc_name"] = v
            elif hn == "code" or hn == "kod":
                mapped["station_code"] = v
        if mapped.get("station_code"):
            rows.append(mapped)
    wb.close()
    return rows
