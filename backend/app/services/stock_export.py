"""Excel representation of the order fulfilment screen."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.schemas import OrderStockRow, StockRow


def build(rows: list[OrderStockRow], summary: list[StockRow]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sipariş Karşılama"
    ws.append(["Sipariş", "Poz", "Müşteri", "Termin", "Stok kodu", "Stok adı",
               "Sipariş miktarı", "Sevk edilen", "Açık miktar", "Rezerve", "Karşılanmamış kalan",
               "Serbest stok (ürün toplamı)", "Plan bitişi", "Karşılama durumu"])
    free_by = {r.item_id: r.free for r in summary}
    for r in rows:
        free = free_by.get(r.item_id, 0)
        status = ("Karşılandı" if r.remaining <= 0 else "Kısmi" if r.reserved + r.shipped > 0
                  else "Stoktan verilebilir" if free > 0 else "Üretim bekliyor")
        ws.append([r.order_no, r.position_no, r.customer, r.due_date, r.item_code, r.item_name,
                   r.quantity, r.shipped, round(r.quantity - r.shipped, 3), r.reserved, r.remaining,
                   free, r.planned_end, status])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
        for col in (4, 13):
            ws.cell(ws.max_row, col).number_format = "dd.mm.yyyy"
        for col in range(7, 13):
            ws.cell(ws.max_row, col).number_format = '#,##0.###'
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
    for i, width in enumerate([20, 14, 30, 15, 22, 40, 20, 18, 18, 18, 25, 29, 15, 23], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    help_ws = wb.create_sheet("Açıklamalar")
    for line in ["Açık miktar = sipariş miktarı − sevk edilen.",
                 "Karşılanmamış kalan = açık miktar − rezerve.",
                 "Karşılandı: siparişin tamamı rezerve veya sevk edilmiştir; tamamının sevk edildiği anlamına gelmez.",
                 "Serbest stok ürün toplamıdır; aynı ürünün sipariş satırlarında tekrar eder. Bu sütunu toplamayın.",
                 "Rapor yalnızca açık siparişleri ve ekrandaki poz / kalanı olan siparişler filtrelerini içerir."]:
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 120
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
