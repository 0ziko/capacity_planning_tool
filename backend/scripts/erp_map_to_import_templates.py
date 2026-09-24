"""ERP 'Sipariş ve Depo Miktarları.xlsx' -> program import sablonlari (doldurulmus).

Cikti: Desktop'ta iki xlsx (siparis + depo girisi). Bos birakilan hucreler kasitli:
ERP'de karsiligi yok veya bu export'ta veri yok.
"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.services.excel import TEMPLATES, _date_cell, _style_header, _autosize, sheet_title

ERP_PATH = Path(r"c:\Users\ozan.deniz\Desktop\Sipariş ve Depo Miktarları.xlsx")
OUT_ORDERS = Path(r"c:\Users\ozan.deniz\Desktop\KapasitePlanlama_Sablon_Siparisler_ERP_dolu.xlsx")
OUT_STOCK = Path(r"c:\Users\ozan.deniz\Desktop\KapasitePlanlama_Sablon_DepoGirisi_ERP_dolu.xlsx")


def _code(v) -> str:
    if v is None or str(v).strip() == "":
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _market_label(grup_kodu: str) -> str:
    s = (grup_kodu or "").strip().upper()
    if s in ("YURTDISI", "YURT DIŞI", "IHRACAT"):
        return "Yurtdışı"
    if s == "YURTICI":
        return "Yerli"
    return ""


def _pos(v) -> str:
    if v is None or str(v).strip() == "":
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _note(*parts: str, max_len: int = 256) -> str:
    s = " | ".join(p for p in parts if p and str(p).strip())
    return s[:max_len]


def build_orders_rows() -> list[list]:
    wb = load_workbook(ERP_PATH, read_only=True, data_only=True)
    ws = wb["Sipariş Ana Veri"]
    all_rows = list(ws.iter_rows(values_only=True))
    header = {h: i for i, h in enumerate(all_rows[0])}
    out: list[list] = []
    for r in all_rows[1:]:
        g = lambda name: r[header[name]]
        tlt = g("TLTUTAR")
        mik = g("MIKTAR") or 0
        unit = ""
        if mik and tlt not in (None, "", 0):
            try:
                unit = round(float(tlt) / float(mik), 4)
            except (TypeError, ValueError, ZeroDivisionError):
                unit = ""
        out.append(
            [
                _code(g("FISNO")),  # Siparis No
                _pos(g("STRA_SIPKONT")),  # Poz No
                (g("CARI_ISIM") or "").strip(),  # Musteri
                _date_cell(g("TARIH")),  # Siparis Tarihi
                _date_cell(g("TERMIN")),  # Termin
                "",  # Revize Termin
                _market_label(g("GRUP_KODU")),  # Pazar
                _code(g("STOK_KODU")),  # Stok Kodu
                g("MIKTAR"),  # Miktar
                unit if unit != 0 else "",  # Birim Fiyat (export'ta tutar 0)
                "",  # Malzeme Durumu
                "",  # Malzeme Hazir Tarihi
                "",  # Malzeme Notu
            ]
        )
    wb.close()
    return out


def build_stock_rows() -> list[list]:
    """Depo satirlari -> sablon. Tarih bilerek bos (ERP anlik bakiye verir, giris tarihi vermez)."""
    wb = load_workbook(ERP_PATH, read_only=True, data_only=True)
    ws = wb["Depo - Ana Veri"]
    all_rows = list(ws.iter_rows(values_only=True))
    header = {h: i for i, h in enumerate(all_rows[0])}
    out: list[list] = []
    for r in all_rows[1:]:
        g = lambda name: r[header[name]]
        lot = _code(g("HUCRE_KODU")) or _code(g("SERI_NO"))
        note = _note(
            f"DEPO={g('DEPOKOD')}" if g("DEPOKOD") is not None else "",
            f"SIP={g('SIPARIS_NO')}" if g("SIPARIS_NO") else "",
            (g("SIPARIS_CARI_ISIM") or "")[:80] if g("SIPARIS_NO") else "",
        )
        out.append(
            [
                "",  # Tarih — ERP'de yok
                _code(g("STOK_KODU")),
                g("BAKIYE"),
                lot,
                note,
            ]
        )
    wb.close()
    return out


def _save(kind: str, data_rows: list[list], path: Path) -> None:
    t = TEMPLATES[kind]
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title(t["title"])
    headers = [c[1] for c in t["columns"]]
    ws.append(headers)
    for row in data_rows:
        ws.append(row)
    # ikinci sayfa: esleme ozeti
    ws2 = wb.create_sheet("ERP esleme notu")
    ws2.append(["Sablon kolonu", "ERP kaynagi / not"])
    if kind == "orders":
        notes = [
            ("Sipariş No", "FISNO"),
            ("Poz No", "STRA_SIPKONT"),
            ("Müşteri", "CARI_ISIM"),
            ("Sipariş Tarihi", "TARIH"),
            ("Termin", "TERMIN"),
            ("Revize Termin", "— (ERP yok, bos)"),
            ("Pazar", "GRUP_KODU → Yerli / Yurtdışı"),
            ("Stok Kodu", "STOK_KODU"),
            ("Miktar", "MIKTAR"),
            ("Birim Fiyat", "TLTUTAR/MIKTAR — bu exportta tutar 0, bos"),
            ("Malzeme *", "— (ERP yok, bos)"),
        ]
    else:
        notes = [
            ("Tarih", "— (ERP anlik bakiye; giris tarihi yok, bos)"),
            ("Stok Kodu", "STOK_KODU"),
            ("Miktar", "BAKIYE (satir bazli; import icin mutabakat gerekir)"),
            ("Lot / Parti", "HUCRE_KODU veya SERI_NO"),
            ("Not", "DEPOKOD + varsa SIPARIS_NO / cari"),
        ]
    for a, b in notes:
        ws2.append([a, b])
    _style_header(ws)
    _autosize(ws)
    _autosize(ws2)
    wb.save(path)


def main() -> None:
    if not ERP_PATH.is_file():
        raise SystemExit(f"ERP dosyasi bulunamadi: {ERP_PATH}")
    orders = build_orders_rows()
    stock = build_stock_rows()
    _save("orders", orders, OUT_ORDERS)
    _save("stock_receipts", stock, OUT_STOCK)
    print(f"Siparis: {len(orders)} satir -> {OUT_ORDERS}")
    print(f"Depo:    {len(stock)} satir -> {OUT_STOCK}")


if __name__ == "__main__":
    main()
