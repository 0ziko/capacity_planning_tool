"""Desktop'taki ERP-dolu sablonlari programa import et."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from app.db.session import SessionLocal
from app.services import excel

ORDERS_PATH = Path(r"c:\Users\ozan.deniz\Desktop\KapasitePlanlama_Sablon_Siparisler_ERP_dolu.xlsx")
STOCK_PATH = Path(r"c:\Users\ozan.deniz\Desktop\KapasitePlanlama_Sablon_DepoGirisi_ERP_dolu.xlsx")
SYNC_DATE = date.today().isoformat()
USERNAME = "admin"


def stock_bytes_with_sync_date() -> bytes:
    """Gorsel sablonda bos birakilan Tarih -> import icin bugun (ERP'de yoktu)."""
    wb = load_workbook(STOCK_PATH)
    ws = wb.worksheets[0]
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 1).value in (None, ""):
            ws.cell(row, 1).value = SYNC_DATE
    import io

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main() -> None:
    db = SessionLocal()
    try:
        ob = ORDERS_PATH.read_bytes()
        rows, parse_errs = excel.read_rows(ob, "orders")
        print("=== SIPARIS ONIZLEME ===")
        print(f"Parse hatalari: {len(parse_errs)}")
        if parse_errs[:5]:
            print("  ornek:", parse_errs[:5])
        if not parse_errs:
            prev = excel.preview_orders_import(db, rows)
            print(f"Dosya satir: {prev.file_row_count}, sistem acik: {prev.system_open_count}")
            print(f"Hata satiri: {len(prev.error_rows)}, eksik stok: {len(prev.missing_item_codes)}")
            print(f"Yeni (dosyada): {len(prev.only_in_file)}, guncellenecek: {len(prev.updated)}, ayni: {prev.unchanged_count}")
        print("\n=== SIPARIS IMPORT (remove_missing=False) ===")
        res_o = excel.run_import(db, "orders", ob, ORDERS_PATH.name, USERNAME, remove_missing=False)
        print(f"inserted={res_o.inserted} updated={res_o.updated} removed={res_o.removed} errors={len(res_o.errors)}")
        if res_o.errors[:8]:
            for e in res_o.errors[:8]:
                print(" ", e)
            if len(res_o.errors) > 8:
                print(f"  ... +{len(res_o.errors) - 8} hata")

        sb = stock_bytes_with_sync_date()
        rows_s, parse_s = excel.read_rows(sb, "stock_receipts")
        print("\n=== DEPO GIRISI ONIZLEME ===")
        print(f"Parse hatalari: {len(parse_s)}, satir: {len(rows_s)}")
        if parse_s[:3]:
            print("  ornek:", parse_s[:3])
        print(f"Tarih alani import icin {SYNC_DATE} ile dolduruldu (ERP'de yoktu).")
        print("\n=== DEPO GIRISI IMPORT ===")
        print("UYARI: Her satir yeni depo girisi EKLER; mevcut stokla TOPLANIR (snapshot degil).")
        res_s = excel.run_import(db, "stock_receipts", sb, STOCK_PATH.name, USERNAME)
        print(f"inserted={res_s.inserted} updated={res_s.updated} errors={len(res_s.errors)}")
        if res_s.errors[:8]:
            for e in res_s.errors[:8]:
                print(" ", e)
    finally:
        db.close()


if __name__ == "__main__":
    main()
