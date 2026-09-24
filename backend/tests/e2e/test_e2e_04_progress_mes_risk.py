"""E2E-04: Günlük üretim/duruş ilerlemesi (legacy), MES önizleme→import→ilerleme/stok defteri, teslimat riski, veri tazeliği."""
import io
from datetime import date, timedelta

from openpyxl import Workbook

from tests.e2e.conftest import WEEK, is_xlsx, plan_auto, upload

ORDER_HDR = ["Sipariş No", "Poz No", "Müşteri", "Termin", "Stok Kodu", "Miktar", "Birim Fiyat"]


def mes_workbook(rows: list[tuple]) -> bytes:
    """Gerçek MES rapor düzeni: sadece 'Üretim Detay Id, Tarih, Malzeme Kodu, Net Üretilen Miktar, İş Merkezi Kodu(makine)' okunur."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Üretim Detay Id", "Tarih", "Malzeme Kodu", "Net Üretilen Miktar", "İş Merkezi Kodu",
               "Sağlam Adet", "Üretilen Miktar", "Başlangıç", "Müşteri", "Süre"])
    for row in rows:
        ws.append([*row, -999, 9999, "IGNORE", "IGNORE", 9999])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _mes_files(content: bytes):
    return {"file": ("mes.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


def test_legacy_production_progress_and_downtime(client, auth, master, sc):
    """Pazartesi 2880 adet Kesim (50 sn) = 40 saat gerçekleşme; duruş 3700 dk ölçülen vs 3600 beklenen → 100 dk fazla; çevrim süresi analizi."""
    upload(client, auth, "orders", ORDER_HDR, [["E2E-P1", "10", "ABC", (WEEK + timedelta(weeks=4)).isoformat(), "E2E-MAM", 9000, 100]])
    plan_auto(client, auth, start_week=WEEK.isoformat(), weeks=8, work_center_ids=master["wc_ids"], replace_existing=True)

    res = upload(client, auth, "production",
                 ["Tarih", "İş Merkezi Kodu", "Stok Kodu", "Operasyon Sıra", "Yarımamül Kodu", "Sipariş No", "Miktar"],
                 [[WEEK.isoformat(), "E2E-PRS", "E2E-MAM", 10, "E2E-MAM-10", "E2E-P1", 2880]])
    sc.step("Günlük üretim import", res.get("errors") in ([], None), created=res.get("created"))
    prog = client.get("/api/progress", headers=auth, params={"week": WEEK.isoformat(), "as_of": (WEEK + timedelta(days=1)).isoformat(), "work_center_ids": [master["prs_id"]]}).json()
    row = prog[0]
    sc.step("Gerçekleşen 40 saat", round(row["actual_hours_to_date"], 1) == 40.0, actual=row["actual_hours_to_date"], expected=row["expected_hours_to_date"], planned=row["planned_hours"], status=row["status"])
    sc.step("Beklenen = planlanan/5 gün", abs(row["expected_hours_to_date"] - row["planned_hours"] / 5) < 0.01)
    sc.step("Kalan saat = plan − gerçekleşen", abs(row["remaining_hours"] - (row["planned_hours"] - row["actual_hours_to_date"])) < 0.01, remaining=row["remaining_hours"], days=row.get("remaining_days"))
    po = client.get("/api/progress/orders", headers=auth, params={"as_of": (WEEK + timedelta(days=1)).isoformat()})
    sc.step("Sipariş ilerlemesi", po.status_code == 200 and any(x.get("order_no") == "E2E-P1" for x in po.json()), count=len(po.json()))
    sc.step("Sipariş ilerlemesi Excel", is_xlsx(client.get("/api/progress/orders.xlsx", headers=auth, params={"as_of": (WEEK + timedelta(days=1)).isoformat()})))
    g = client.get("/api/plan/gantt", headers=auth, params={"work_center_id": master["prs_id"], "start": WEEK.isoformat(), "end": (WEEK + timedelta(days=27)).isoformat(), "as_of": (WEEK + timedelta(days=1)).isoformat()}).json()
    bar = next((b for b in g["bars"] if b.get("semi_finished_code") == "E2E-MAM-10"), None)
    sc.step("Gantt üretileni gösterir (2880)", bar is not None and bar.get("produced_qty") == 2880, produced=bar.get("produced_qty") if bar else None, status=bar.get("status") if bar else None)

    res = upload(client, auth, "downtime",
                 ["Tarih", "İş Merkezi Kodu", "Sebep Kodu", "Sebep", "Süre (dk)", "Süre Değeri (dk)", "Süre Temeli"],
                 [[WEEK.isoformat(), "E2E-PRS", "MLZ", "Malzeme bekleme", 3000, 3000, "labor_minutes"],
                  [WEEK.isoformat(), "E2E-PRS", "SET", "Setup", 700, 700, "labor_minutes"]])
    sc.step("Duruş import", res.get("errors") in ([], None))
    dt = client.get("/api/analysis/downtime", headers=auth, params={"start": WEEK.isoformat(), "end": WEEK.isoformat(), "work_center_ids": [master["prs_id"]]}).json()
    sc.step("Fazla duruş 100 dk (3700 − (10−4)×10×60)", dt["totals"][0]["excess_minutes"] == 100, excess=dt["totals"][0]["excess_minutes"])
    sc.step("En büyük sebep MLZ", dt["reasons"][0]["reason_code"] == "MLZ", reasons=[(r["reason_code"], r.get("minutes")) for r in dt["reasons"]])
    sc.step("Duruş Excel", is_xlsx(client.get("/api/analysis/downtime.xlsx", headers=auth, params={"start": WEEK.isoformat(), "end": WEEK.isoformat()})))
    ct = client.get("/api/analysis/cycletime", headers=auth, params={"start": WEEK.isoformat(), "end": (WEEK + timedelta(days=6)).isoformat()})
    sc.step("Çevrim süresi analizi", ct.status_code == 200, rows=len(ct.json()) if isinstance(ct.json(), list) else list(ct.json().keys())[:6])
    sc.step("Çevrim süresi Excel", is_xlsx(client.get("/api/analysis/cycletime.xlsx", headers=auth)))


def test_mes_preview_import_progress_inventory_risk(client, auth, master, sc):
    """MES raporu: önizleme yazmaz, token'sız import 422, token ile import, tekrar import mükerrer yaratmaz; tanımsız kod serbest stok; ilerleme/defter/risk."""
    upload(client, auth, "orders", ORDER_HDR, [["E2E-MS", "10", "ABC", (WEEK + timedelta(weeks=4)).isoformat(), "E2E-MAM", 3600, 100]])
    plan_auto(client, auth, start_week=WEEK.isoformat(), weeks=8, work_center_ids=master["wc_ids"], replace_existing=True)
    status0 = client.get("/api/mes/source-status", headers=auth).json()
    sc.step("Kaynak durumu (legacy, MES kaydı yok)", status0["production_source"] == "legacy" and status0["has_mes_records"] is False, status=status0)

    content = mes_workbook([("E2E-M1", WEEK, "E2E-MAM-10", 100, "E2E-PRS-M1"),
                            ("E2E-M2", WEEK + timedelta(days=1), "E2E-MAM-10", 50, "E2E-PRS-M1"),
                            ("E2E-M3", WEEK, "E2E-TANIMSIZ-99", 7, "E2E-PRS-M1")])
    pv = client.post("/api/mes/preview", headers=auth, files=_mes_files(content))
    sc.step("MES önizleme", pv.status_code == 200 and pv.json().get("token"), keys=list(pv.json().keys())[:14] if pv.status_code == 200 else pv.text[:200])
    body = pv.json()
    summary = body.get("summary") or body.get("counts") or {k: v for k, v in body.items() if isinstance(v, (int, float))}
    sc.info("Önizleme sayaçları", **{str(k): v for k, v in list(summary.items())[:12]})
    sc.step("Tanımsız kod serbest stok olarak sınıflandı", "free_stock" in pv.text, mentions=pv.text.count("free_stock"))
    sc.step("Önizleme yazmadı", client.get("/api/mes/source-status", headers=auth).json()["has_mes_records"] is False)
    sc.step("Token'sız import reddedilir (422)", client.post("/api/mes/import", headers=auth, files=_mes_files(content)).status_code == 422)
    imp = client.post("/api/mes/import", headers=auth, files=_mes_files(content), params={"token": body["token"]})
    sc.step("Token ile import", imp.status_code == 200, result={k: v for k, v in imp.json().items() if isinstance(v, (int, str, bool))} if imp.status_code == 200 else imp.text[:200])
    sc.step("MES kaydı var", client.get("/api/mes/source-status", headers=auth).json()["has_mes_records"] is True)

    pv2 = client.post("/api/mes/preview", headers=auth, files=_mes_files(content)).json()
    imp2 = client.post("/api/mes/import", headers=auth, files=_mes_files(content), params={"token": pv2["token"]})
    sc.step("Aynı dosya tekrar import (idempotent)", imp2.status_code == 200, result={k: v for k, v in imp2.json().items() if isinstance(v, (int, str, bool))})
    inv = client.get("/api/mes/inventory", headers=auth, params={"as_of": (WEEK + timedelta(days=2)).isoformat()}).json()
    inv_text = str(inv)
    sc.step("Yarımamul stok defteri E2E-MAM-10 içeriyor", "E2E-MAM-10" in inv_text, keys=list(inv.keys())[:8] if isinstance(inv, dict) else len(inv))
    sc.step("Defterde 150 adet tek sefer (mükerrer yok)", "150" in inv_text and "300" not in inv_text.replace("3000", ""))
    sc.step("Stok defteri Excel", is_xlsx(client.get("/api/mes/inventory.xlsx", headers=auth, params={"as_of": (WEEK + timedelta(days=2)).isoformat()})))
    fs = client.get("/api/mes/free-stock", headers=auth).json()
    sc.step("Serbest stok listesinde tanımsız kod", "E2E-TANIMSIZ-99" in str(fs), count=len(fs) if isinstance(fs, list) else None)

    mp = client.get("/api/mes/progress", headers=auth, params={"as_of": (WEEK + timedelta(days=2)).isoformat(), "work_center_ids": [master["prs_id"]]})
    sc.step("MES ilerleme raporu", mp.status_code == 200, keys=list(mp.json().keys())[:10] if isinstance(mp.json(), dict) else len(mp.json()))
    sc.step("MES ilerleme Excel", is_xlsx(client.get("/api/mes/progress.xlsx", headers=auth, params={"as_of": (WEEK + timedelta(days=2)).isoformat()})))
    dr = client.get("/api/mes/delivery-risk", headers=auth, params={"as_of": (WEEK + timedelta(days=2)).isoformat(), "horizon": 8})
    sc.step("Teslimat riski analizi", dr.status_code == 200, keys=list(dr.json().keys())[:10] if isinstance(dr.json(), dict) else len(dr.json()))
    sc.step("Teslimat riski E2E-MS siparişini kapsıyor", "E2E-MS" in dr.text)
    sc.step("Teslimat riski salt okunur (sipariş/plan değişmedi)", any(o["order_no"] == "E2E-MS" for o in client.get("/api/orders", headers=auth).json()))


def test_data_freshness_checkpoints(client, auth, master, sc):
    """Veri tazeliği: kontrol noktaları listelenir; 'değişiklik yok' onayı bugünkü durumu günceller."""
    r = client.get("/api/data-freshness", headers=auth)
    sc.step("Veri tazeliği raporu", r.status_code == 200, keys=list(r.json().keys())[:8])
    cps = r.json().get("checkpoints", [])
    keys = {c.get("key") for c in cps}
    sc.step("Dört kontrol noktası", {"finished_stock", "wip_stock", "production_output", "open_orders"} <= keys, keys=sorted(keys))
    c = client.post("/api/data-freshness/checkpoints/open_orders/confirm-no-change", headers=auth)
    sc.step("'Değişiklik yok' onayı", c.status_code == 200, status=c.status_code, body=c.text[:160] if c.status_code >= 400 else None)
    after = next(x for x in client.get("/api/data-freshness", headers=auth).json()["checkpoints"] if x["key"] == "open_orders")
    sc.step("Onay sonrası durum güncel", after.get("status") in ("fresh", "ok", "confirmed", "current", "acked") or after.get("acknowledged_today") or after.get("ack_today"), checkpoint={k: v for k, v in after.items() if isinstance(v, (str, int, bool))})
    bad = client.post("/api/data-freshness/checkpoints/yok/confirm-no-change", headers=auth)
    sc.step("Bilinmeyen kontrol noktası reddedilir", bad.status_code in (400, 404, 422), status=bad.status_code)
