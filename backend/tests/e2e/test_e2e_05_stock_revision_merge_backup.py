"""E2E-05: Stok & rezervasyon & sevk, plan revizyonu (iş taşıma), üretim partisi birleştirme, birlikte sevk,
yedek/şablon/import günlüğü ve owner temizlik akışı."""
import io
import json
from datetime import date, timedelta

from openpyxl import load_workbook

from tests.e2e.conftest import WEEK, is_xlsx, plan_auto, upload

ORDER_HDR = ["Sipariş No", "Poz No", "Müşteri", "Termin", "Stok Kodu", "Miktar", "Birim Fiyat"]


def _summary(client, auth, code):
    return next(x for x in client.get("/api/stock/summary", headers=auth).json() if x["item_code"] == code)


def _stock_orders(client, auth, **params):
    return {x["order_no"]: x for x in client.get("/api/stock/orders", headers=auth, params=params).json()}


def _res_ids(client, auth, order_id):
    return [r["id"] for r in client.get("/api/stock/reservations", headers=auth, params={"order_id": order_id}).json()]


def test_stock_receipt_reserve_ship_close_reopen(client, auth, master, sc):
    """Depo girişi 60 → manuel 30 + otomatik (termin sırası) 30 → kısmi/tam sevk → sipariş kapanır → sevk iptali siparişi açar."""
    upload(client, auth, "items", ["Stok Kodu", "Stok Adı", "Ürün Grubu"], [["E2E-STK", "E2E Stoklu Ürün", "E2EGRP"]])
    upload(client, auth, "routing", ["Stok Kodu", "Sıra", "Operasyon", "İş Merkezi Kodu", "Çevrim Süresi (sn)"], [["E2E-STK", 10, "Montaj", "E2E-MNT", 30]])  # giriş kapısı: rota şart
    upload(client, auth, "orders", ORDER_HDR, [["E2E-ST-A", "10", "A", (WEEK + timedelta(weeks=1)).isoformat(), "E2E-STK", 50, 10],
                                              ["E2E-ST-B", "10", "B", (WEEK + timedelta(weeks=2)).isoformat(), "E2E-STK", 40, 10],
                                              ["E2E-ST-C", "10", "C", (WEEK + timedelta(weeks=3)).isoformat(), "E2E-STK", 40, 10]])
    today = date.today().isoformat()
    r = client.post("/api/stock/receipts", headers=auth, json={"item_code": "E2E-STK", "receipt_date": today, "quantity": 60, "lot": "E2E-L1"})
    sc.step("Depo girişi 60", r.status_code == 201, status=r.status_code)
    row = _summary(client, auth, "E2E-STK")
    sc.step("Stok özeti: eldeki 60, serbest 60", row["on_hand"] == 60 and row["free"] == 60, row={k: row[k] for k in ("on_hand", "reserved", "free", "shipped")})
    orders = _stock_orders(client, auth)
    r = client.post("/api/stock/reservations", headers=auth, json={"order_id": orders["E2E-ST-C"]["order_id"], "quantity": 30, "note": "e2e kritik"})
    sc.step("Manuel rezervasyon C=30", r.status_code == 201, status=r.status_code)
    sc.step("Serbest stoğu aşan rezervasyon reddedilir", client.post("/api/stock/reservations", headers=auth, json={"order_id": orders["E2E-ST-B"]["order_id"], "quantity": 31}).status_code == 400)
    pv = client.post("/api/stock/reservations/auto/preview", headers=auth, json={"item_ids": [row["item_id"]]})
    sc.step("Otomatik rezervasyon önizleme", pv.status_code == 200 and pv.json().get("preview_token"), keys=list(pv.json().keys())[:8])
    ar = client.post("/api/stock/reservations/auto", headers=auth, json={"preview_token": pv.json()["preview_token"]})
    sc.step("Otomatik rezervasyon uygulandı", ar.status_code == 200, status=ar.status_code, body=ar.text[:160] if ar.status_code >= 400 else None)
    orders = _stock_orders(client, auth)
    sc.step("Termin sırası: A 30 aldı, B 0", orders["E2E-ST-A"]["reserved"] == 30 and orders["E2E-ST-B"]["reserved"] == 0, a=orders["E2E-ST-A"]["reserved"], b=orders["E2E-ST-B"]["reserved"], c=orders["E2E-ST-C"]["reserved"])
    sc.step("Serbest stok 0", _summary(client, auth, "E2E-STK")["free"] == 0)

    c_id = _res_ids(client, auth, orders["E2E-ST-C"]["order_id"])[0]
    sc.step("Kısmi sevk C 10", client.post(f"/api/stock/reservations/{c_id}/ship", headers=auth, json={"quantity": 10}).status_code == 201)
    sc.step("Kalan sevk C 20", client.post(f"/api/stock/reservations/{c_id}/ship", headers=auth, json={}).status_code == 201)
    row = _summary(client, auth, "E2E-STK")
    sc.step("Özet: eldeki 30, sevk 30, rezerve 30", row["on_hand"] == 30 and row["shipped"] == 30 and row["reserved"] == 30, row={k: row[k] for k in ("on_hand", "reserved", "free", "shipped")})
    oc = _stock_orders(client, auth, include_closed=True)["E2E-ST-C"]
    sc.step("C: 30 sevk, 10 kalan, açık", oc["shipped"] == 30 and oc["remaining"] == 10 and oc["status"] == "open")

    upload(client, auth, "stock_receipts", ["Tarih", "Stok Kodu", "Miktar", "Lot / Parti"], [[today, "E2E-STK", 20, "E2E-L2"]])
    sc.step("Excel ile depo girişi 20", _summary(client, auth, "E2E-STK")["free"] == 20)
    sc.step("A'ya 20 daha rezerve (toplam 50)", client.post("/api/stock/reservations", headers=auth, json={"order_id": orders["E2E-ST-A"]["order_id"], "quantity": 20}).status_code == 201)
    sc.step("İhtiyacı aşan rezervasyon reddedilir", client.post("/api/stock/reservations", headers=auth, json={"order_id": orders["E2E-ST-A"]["order_id"], "quantity": 1}).status_code == 400)
    for rid in _res_ids(client, auth, orders["E2E-ST-A"]["order_id"]):
        assert client.post(f"/api/stock/reservations/{rid}/ship", headers=auth, json={}).status_code == 201
    oa = _stock_orders(client, auth, include_closed=True)["E2E-ST-A"]
    sc.step("A tamamen sevk edildi → kapandı", oa["status"] == "closed" and oa["remaining"] == 0, shipped=oa["shipped"])
    ships = client.get("/api/stock/shipments", headers=auth, params={"order_id": orders["E2E-ST-A"]["order_id"]}).json()
    sc.step("Sevk iptali", client.delete(f"/api/stock/shipments/{ships[0]['id']}", headers=auth).status_code == 204, shipments=len(ships))
    sc.step("İptal sonrası A yeniden açık", _stock_orders(client, auth)["E2E-ST-A"]["status"] == "open")
    sc.step("Karşılama Excel", is_xlsx(client.get("/api/stock/orders/export.xlsx", headers=auth)))
    rc = client.get("/api/stock/receipts", headers=auth, params={"item_id": row["item_id"]}).json()
    sc.step("Depo giriş listesi (2 giriş)", len(rc) == 2, lots=[x.get("lot") for x in rc])


def test_plan_revision_job_move_approve_reject_cancel(client, auth, master, sc):
    """İş taşıma revizyonu: taşıma önizleme → revizyon → değişiklik → hesapla → onayla canlı plana yazar; ayrıca reddet ve iptal akışları."""
    upload(client, auth, "items", ["Stok Kodu", "Stok Adı", "Ürün Grubu"], [["E2E-RVI", "E2E Revizyon Ürünü", "E2ERV"]])
    upload(client, auth, "routing", ["Stok Kodu", "Sıra", "Operasyon", "İş Merkezi Kodu", "Çevrim Süresi (sn)", "Yarımamül Kodu"],
           [["E2E-RVI", 10, "Pres", "E2E-PRS", 100, "E2E-RVI-10"]])
    a = client.post("/api/orders", headers=auth, json={"order_no": "E2E-RV-A", "due_date": (WEEK + timedelta(weeks=2)).isoformat(), "item_code": "E2E-RVI", "quantity": 3600, "unit_price": 10}).json()
    b = client.post("/api/orders", headers=auth, json={"order_no": "E2E-RV-B", "due_date": (WEEK + timedelta(weeks=6)).isoformat(), "item_code": "E2E-RVI", "quantity": 1800, "unit_price": 10}).json()
    plan_auto(client, auth, start_week=WEEK.isoformat(), weeks=8, work_center_ids=[master["prs_id"]], replace_existing=True)

    def hours_by_week(order_no):
        out = {}
        for l in client.get("/api/plan/lines", headers=auth, params={"start": WEEK.isoformat(), "work_center_ids": [master["prs_id"]]}).json():
            if l.get("order_no") == order_no:
                out[l["week_start"]] = round(out.get(l["week_start"], 0) + l["planned_hours"], 1)
        return out

    before = hours_by_week("E2E-RV-B")
    sc.step("Başlangıç: B (50 saat) 1. haftada", before.get(WEEK.isoformat()) == 50.0, b_weeks=before, a_weeks=hours_by_week("E2E-RV-A"))
    pv = client.get("/api/plan/revisions/move-preview", headers=auth, params={"order_id": b["id"]}).json()
    sc.step("Taşıma önizleme: 1800 adet taşınabilir, kilit yok", pv["movable_qty"] == 1800 and pv["ops"][0]["locked"] is False, consumed=pv.get("consumed_work_centers"))

    body = {"reason_codes": ["vip_pull_in"], "note": "e2e taşıma", "start_week": WEEK.isoformat(), "weeks": 8, "work_center_ids": [master["prs_id"]], "mode": "due_date"}
    rev = client.post("/api/plan/revisions", headers=auth, json=body)
    sc.step("Revizyon taslağı", rev.status_code in (200, 201), status=rev.status_code, body=rev.text[:200] if rev.status_code >= 400 else None)
    rev = rev.json()
    target = (WEEK + timedelta(weeks=2)).isoformat()
    ch = client.post(f"/api/plan/revisions/{rev['id']}/changes", headers=auth, json={
        "entity_type": "order", "entity_id": b["id"], "extra_key": "E2E-RVI", "field": "job_move",
        "new_value": json.dumps({"item_code": "E2E-RVI", "start_date": target, "qty_mode": "remaining", "quantity": None})})
    sc.step("İş taşıma değişikliği eklendi", ch.status_code == 200, status=ch.status_code, body=ch.text[:200] if ch.status_code >= 400 else None)
    calc = client.post(f"/api/plan/revisions/{rev['id']}/calculate", headers=auth)
    sc.step("Revizyon hesaplandı", calc.status_code == 200, compare_keys=list(calc.json().get("compare", {}).keys())[:10] if calc.status_code == 200 else calc.text[:200])
    sc.step("Boş kapasiteye taşıma kimseyi ötelemez", calc.json()["compare"].get("bumped_orders") == [], bumped=calc.json()["compare"].get("bumped_orders"))
    sc.step("Canlı plan hesaplamadan etkilenmedi", hours_by_week("E2E-RV-B") == before)
    ap = client.post(f"/api/plan/revisions/{rev['id']}/approve", headers=auth)
    sc.step("Revizyon onaylandı", ap.status_code == 200, status=ap.status_code, body=ap.text[:200] if ap.status_code >= 400 else None)
    after = hours_by_week("E2E-RV-B")
    sc.step("B artık 3. haftada (50 saat)", after.get(target) == 50.0 and after.get(WEEK.isoformat()) is None, b_weeks=after)
    sc.step("A yerinde kaldı", hours_by_week("E2E-RV-A").get(WEEK.isoformat()) == 100.0, a_weeks=hours_by_week("E2E-RV-A"))
    detail = client.get(f"/api/plan/revisions/{rev['id']}", headers=auth).json()
    sc.step("Revizyon detayı 'approved'", detail.get("status") in ("approved", "applied"), status=detail.get("status"), events=len(detail.get("events", [])))

    rev2 = client.post("/api/plan/revisions", headers=auth, json=body).json()
    client.post(f"/api/plan/revisions/{rev2['id']}/calculate", headers=auth)
    rj = client.post(f"/api/plan/revisions/{rev2['id']}/reject", headers=auth, json={"note": "e2e red"})
    sc.step("Revizyon reddedildi", rj.status_code == 200 and rj.json().get("status") == "rejected", status=rj.json().get("status") if rj.status_code == 200 else rj.text[:160])
    rev3 = client.post("/api/plan/revisions", headers=auth, json=body).json()
    cn = client.post(f"/api/plan/revisions/{rev3['id']}/cancel", headers=auth)
    sc.step("Revizyon iptal edildi", cn.status_code == 200 and cn.json().get("status") in ("cancelled", "canceled"), status=cn.json().get("status") if cn.status_code == 200 else cn.text[:160])
    lst = client.get("/api/plan/revisions", headers=auth).json()
    sc.step("Revizyon listesi 3 kayıt", len(lst) >= 3, statuses=[r.get("status") for r in lst][:6])
    sc.step("Canlı plan red/iptalden etkilenmedi", hours_by_week("E2E-RV-B") == after)


def test_merge_batches_and_co_shipment(client, auth, master, sc):
    """Aynı ürün/termin siparişleri birleştirme önerisi → üretim partisi → plan → parti çözme; birlikte sevk seçenekli plan."""
    upload(client, auth, "items", ["Stok Kodu", "Stok Adı", "Ürün Grubu"], [["E2E-MRG", "E2E Parti Ürünü", "E2EMRG"]])
    upload(client, auth, "routing", ["Stok Kodu", "Sıra", "Operasyon", "İş Merkezi Kodu", "Çevrim Süresi (sn)", "Yarımamül Kodu"],
           [["E2E-MRG", 10, "Pres", "E2E-PRS", 100, "E2E-MRG-10"]])
    due = (WEEK + timedelta(weeks=3)).isoformat()
    m1 = client.post("/api/orders", headers=auth, json={"order_no": "E2E-MG-1", "customer": "K1", "due_date": due, "item_code": "E2E-MRG", "quantity": 360, "unit_price": 5}).json()
    m2 = client.post("/api/orders", headers=auth, json={"order_no": "E2E-MG-2", "customer": "K2", "due_date": due, "item_code": "E2E-MRG", "quantity": 720, "unit_price": 5}).json()
    sug = client.get("/api/plan/merge-suggestions", headers=auth, params={"tolerance_days": 5})
    sc.step("Birleştirme önerisi ikisini gruplar", sug.status_code == 200 and "E2E-MG-1" in sug.text and "E2E-MG-2" in sug.text, groups=len(sug.json()))
    mg = client.post("/api/plan/merge", headers=auth, json={"order_ids": [m1["id"], m2["id"]], "note": "e2e parti"})
    sc.step("Üretim partisi oluşturuldu", mg.status_code in (200, 201), status=mg.status_code, body=mg.text[:200] if mg.status_code >= 400 else None)
    batch = mg.json()
    batches = client.get("/api/plan/production-batches", headers=auth).json()
    sc.step("Parti listede", any(b["id"] == batch["id"] for b in batches), count=len(batches))
    out = plan_auto(client, auth, start_week=WEEK.isoformat(), weeks=8, work_center_ids=[master["prs_id"]], replace_existing=True)
    sc.step("Parti planlandı", out.get("created", 0) > 0 and out.get("unplanned") == [], created=out.get("created"))
    lines = client.get("/api/plan/lines", headers=auth, params={"start": WEEK.isoformat(), "work_center_ids": [master["prs_id"]]}).json()
    total = round(sum(l["planned_hours"] for l in lines if l.get("production_batch_id") == batch["id"] or l.get("order_no") in ("E2E-MG-1", "E2E-MG-2")), 1)
    sc.step("Parti toplam 30 saat (1080 adet × 100 sn)", total == 30.0, hours=total)
    sc.step("Parti çözüldü", client.delete(f"/api/plan/merge/{batch['id']}", headers=auth).status_code in (200, 204))
    sc.step("Siparişler tekrar bağımsız açık", {"E2E-MG-1", "E2E-MG-2"} <= {o["order_no"] for o in client.get("/api/orders", headers=auth).json()})

    upload(client, auth, "orders", ORDER_HDR, [["E2E-CS", "10", "Müşteri X", due, "E2E-MAM", 360, 10],
                                              ["E2E-CS", "20", "Müşteri X", due, "E2E-MRG", 360, 10]])
    out = plan_auto(client, auth, start_week=WEEK.isoformat(), weeks=8, work_center_ids=master["wc_ids"], replace_existing=True,
                    co_shipment={"enabled": True, "ready_before_delivery_days": 3, "selections": [{"order_no": "E2E-CS"}]})
    sc.step("Birlikte sevk planı", out.get("created", 0) > 0, created=out.get("created"), unplanned=out.get("unplanned"))
    sched = [s for s in client.get("/api/plan/orders", headers=auth).json() if s.get("order_no") == "E2E-CS"]
    sc.step("Her iki poz sipariş takviminde", len(sched) == 2, positions=[s.get("position_no") for s in sched])


def test_backup_templates_import_log(client, auth, master, sc):
    """Tek tık yedek Excel'i tüm ana sayfaları içerir; tüm şablonlar iner; import günlüğü kayıtları tutar."""
    b = client.get("/api/backup.xlsx", headers=auth)
    sc.step("Yedek Excel", is_xlsx(b), size_kb=round(len(b.content) / 1024, 1))
    names = load_workbook(io.BytesIO(b.content), read_only=True).sheetnames
    expected = {"Siparişler", "Haftalık İş Gücü"}
    sc.step("Yedekte beklenen sayfalar", expected <= set(names), sheets=names)
    kinds = [k["kind"] for k in client.get("/api/imports/kinds", headers=auth).json()]
    sc.step("Import türleri", {"workcenters", "items", "routing", "orders", "production", "downtime", "wc_weeks", "stock_receipts"} <= set(kinds), kinds=kinds)
    bad = [k for k in kinds if not is_xlsx(client.get(f"/api/imports/template/{k}", headers=auth))]
    sc.step("Tüm şablonlar iniyor", bad == [], failed=bad, count=len(kinds))
    log = client.get("/api/imports/log", headers=auth).json()
    rows = log if isinstance(log, list) else log.get("items", log.get("rows", []))
    sc.step("Import günlüğü", len(rows) > 0 and any(r.get("kind") in kinds for r in rows), count=len(rows), last=({k: rows[0].get(k) for k in ("kind", "created", "updated", "errors")} if rows else None))


def test_owner_purge_flow(client, auth, master, sc):
    """Owner rolü: istatistik, kayıt listesi, temizlik önizleme, 'SIL' onayıyla seçili siparişleri kalıcı silme; admin purge yapamaz."""
    r = client.post("/api/auth/login", data={"username": "owner", "password": "owner123"})
    sc.step("Owner girişi", r.status_code == 200, status=r.status_code)
    oh = {"Authorization": f"Bearer {r.json()['access_token']}"}
    sc.step("Owner istatistik", client.get("/api/owner/stats", headers=oh).status_code == 200)
    sc.step("Admin owner paneline giremez", client.get("/api/owner/stats", headers=auth).status_code == 403)
    for i in (1, 2):
        client.post("/api/orders", headers=auth, json={"order_no": f"E2E-PG-{i}", "due_date": (WEEK + timedelta(weeks=1)).isoformat(), "item_code": "E2E-MAM", "quantity": 10})
    recs = client.get("/api/owner/records", headers=oh, params={"target": "orders", "limit": 50}).json()
    rows = recs if isinstance(recs, list) else recs.get("records", [])
    ids = [x["id"] for x in rows if str(x.get("label", x.get("order_no", ""))).startswith("E2E-PG-")]
    sc.step("Owner kayıt listesi E2E siparişlerini görüyor", len(ids) == 2, ids=ids)
    pv = client.post("/api/owner/purge/preview", headers=oh, json={"target": "orders", "confirm": "SIL", "ids": ids})
    sc.step("Temizlik önizleme", pv.status_code == 200, body={k: v for k, v in pv.json().items() if isinstance(v, (int, str, bool))} if pv.status_code == 200 else pv.text[:160])
    sc.step("Onay kelimesi olmadan silme reddedilir", client.post("/api/owner/purge", headers=oh, json={"target": "orders", "confirm": "HAYIR", "ids": ids}).status_code in (400, 422))
    sc.step("Admin purge yapamaz (403)", client.post("/api/owner/purge", headers=auth, json={"target": "orders", "confirm": "SIL", "ids": ids}).status_code == 403)
    p = client.post("/api/owner/purge", headers=oh, json={"target": "orders", "confirm": "SIL", "ids": ids})
    sc.step("Seçili siparişler kalıcı silindi", p.status_code == 200, result=p.json() if p.status_code == 200 else p.text[:160])
    sc.step("Siparişler listede yok", all(not o["order_no"].startswith("E2E-PG-") for o in client.get("/api/orders", headers=auth).json()))
