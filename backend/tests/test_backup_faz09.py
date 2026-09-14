"""FAZ 09 — Excel disa aktarim, migrate idempotent, restore hedef guvenligi."""

import io
import json
from datetime import date

import pytest
from openpyxl import load_workbook

from app.db.migrate import ensure_columns
from app.db.session import SessionLocal, engine
from app.models import BomLine, Item, Machine, RoutingOperation, RoutingOperationStation, WorkCenter
from app.services.backup_guard import validate_restore_target
from app.services.data_integrity import audit_data
from app.services.excel import build_backup, import_bom, import_routing

TEST_PG = __import__("os").environ.get("TEST_PG_URL", "").strip()


def test_restore_target_rejects_live_db():
    with pytest.raises(ValueError, match="Canli"):
        validate_restore_target("kapasite", source_db="kapasite")
    validate_restore_target("kapasite_audit_test", source_db="kapasite")


def test_excel_export_bom_routing_fields(db):
    wc = WorkCenter(code="EXP-WC", name="X", is_planned=True, planning_reserve_pct=12.5)
    db.add(wc)
    db.flush()
    m = Machine(work_center_id=wc.id, code="EXP-M1", name="M")
    db.add(m)
    db.flush()
    it = Item(code="EXP-FG", name="FG", product_group="G")
    db.add(it)
    db.flush()
    db.add(
        BomLine(
            item_id=it.id,
            component_code="C1",
            quantity=2,
            source_wip="WIP-1",
            branch_listing_sira=5,
            recipe_seq=10,
        )
    )
    op = RoutingOperation(item_id=it.id, seq=10, work_center_id=wc.id, cycle_time_sec=60, primary_machine_id=m.id)
    db.add(op)
    db.flush()
    db.add(RoutingOperationStation(operation_id=op.id, machine_id=m.id, is_primary=True))
    db.commit()

    wb = load_workbook(io.BytesIO(build_backup(db)))
    bom = wb["BOM"]
    assert bom.max_row >= 2
    hdr = [bom.cell(1, c).value for c in range(1, bom.max_column + 1)]
    row = [bom.cell(2, c).value for c in range(1, bom.max_column + 1)]
    assert "Kaynak Yarımamül" in hdr
    assert row[hdr.index("Kaynak Yarımamül")] == "WIP-1"
    assert row[hdr.index("Dal Sıra")] == 5
    assert row[hdr.index("Reçete Sıra")] == 10

    rota = wb["Rota"]
    rh = [rota.cell(1, c).value for c in range(1, rota.max_column + 1)]
    rr = [rota.cell(2, c).value for c in range(1, rota.max_column + 1)]
    assert rr[rh.index("Birincil Makine Kodu")] == "EXP-M1"
    assert "Rota İstasyonları" in wb.sheetnames


def test_bom_routing_import_roundtrip_fields(db):
    wc = WorkCenter(code="IMP-WC", name="X", is_planned=True)
    db.add(wc)
    db.flush()
    m = Machine(work_center_id=wc.id, code="IMP-M1", name="M")
    db.add(m)
    db.flush()
    it = Item(code="IMP-FG", name="FG")
    db.add(it)
    db.commit()

    _, _, errs = import_bom(
        db,
        [
            {
                "_row": 2,
                "item_code": "IMP-FG",
                "component_code": "C1",
                "quantity": 1,
                "source_wip": "DAL",
                "branch_listing_sira": 3,
                "recipe_seq": 7,
            }
        ],
    )
    assert errs == []
    _, _, errs = import_routing(
        db,
        [
            {
                "_row": 2,
                "item_code": "IMP-FG",
                "seq": 10,
                "wc_code": "IMP-WC",
                "cycle_time_sec": 30,
                "primary_machine_code": "IMP-M1",
            }
        ],
    )
    db.commit()
    assert errs == []
    bl = db.query(BomLine).join(Item).filter(Item.code == "IMP-FG", BomLine.component_code == "C1").one()
    assert bl.source_wip == "DAL" and bl.branch_listing_sira == 3 and bl.recipe_seq == 7
    op = db.query(RoutingOperation).join(Item).filter(Item.code == "IMP-FG").one()
    assert op.primary_machine_id == m.id


def test_ensure_columns_idempotent(db):
    db.add(Item(code="MIG-1", name="x"))
    db.commit()
    n1 = db.query(Item).count()
    added1 = ensure_columns(engine)
    added2 = ensure_columns(engine)
    n2 = db.query(Item).count()
    assert n2 == n1
    assert isinstance(added1, list) and isinstance(added2, list)


def test_audit_reports_negative_cycle(db):
    wc = WorkCenter(code="AUD-WC", name="X")
    db.add(wc)
    db.flush()
    it = Item(code="AUD-1", name="x")
    db.add(it)
    db.flush()
    db.add(RoutingOperation(item_id=it.id, seq=10, work_center_id=wc.id, cycle_time_sec=-1))
    db.commit()
    rep = audit_data(db)
    assert rep["error_count"] >= 1
    assert any(i["code"] == "negative_cycle_time" for i in rep["issues"])


@pytest.mark.skipif(not TEST_PG, reason="TEST_PG_URL yok — pg_dump/restore dongusu icin")
def test_pg_dump_restore_critical_counts_match():
    """Sentetik kaynak -> dump -> bos audit DB; kritik tablo sayilari esit."""
    import os
    import subprocess
    import tempfile
    from pathlib import Path
    from urllib.parse import urlparse

    from sqlalchemy import create_engine, text
    from sqlalchemy.orm import Session

    from app.services.data_integrity import critical_table_counts

    src_url = TEST_PG
    parsed = urlparse(src_url.replace("postgresql+psycopg://", "postgresql://"))
    backend = Path(__file__).resolve().parents[1]
    pg_bin = backend / ".postgres" / "pgsql" / "bin"
    dump_exe = pg_bin / "pg_dump.exe"
    if not dump_exe.exists():
        pytest.skip("Portable pg_dump yok")

    target_db = "kapasite_audit_test"
    validate_restore_target(target_db, source_db=parsed.path.lstrip("/"))

    with Session(create_engine(src_url)) as db:
        expected = critical_table_counts(db)

    with tempfile.TemporaryDirectory() as td:
        dump_path = Path(td) / "test.dump"
        env = os.environ.copy()
        if parsed.password:
            env["PGPASSWORD"] = parsed.password
        subprocess.run(
            [str(dump_exe), "-Fc", "-h", parsed.hostname or "localhost", "-p", str(parsed.port or 5432), "-U", parsed.username, "-d", parsed.path.lstrip("/"), "-f", str(dump_path), "--no-owner", "--no-acl"],
            check=True,
            env=env,
            capture_output=True,
        )
        ps1 = backend / "scripts" / "restore_postgres_test.ps1"
        counts_file = backend / "backup" / "restore_expected_counts.json"
        counts_file.parent.mkdir(exist_ok=True)
        counts_file.write_text(json.dumps(expected), encoding="utf-8")
        try:
            r = subprocess.run(
                ["powershell", "-NoProfile", "-File", str(ps1), "-DumpFile", str(dump_path), "-TargetDatabase", target_db],
                cwd=str(backend),
                capture_output=True,
                text=True,
            )
            assert r.returncode == 0, r.stdout + r.stderr
        finally:
            counts_file.unlink(missing_ok=True)
