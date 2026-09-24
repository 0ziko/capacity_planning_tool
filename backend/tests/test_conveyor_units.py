from datetime import datetime, timedelta
from types import SimpleNamespace
import pytest
from app.models import Item, RoutingOperation, WorkCenter
from app.services.routing_resource import compute_operation_need, standard_unit_hours
from app.services.mes import standard_unit_hours as mes_unit_hours
from app.services.daily_scheduler import _process_outputs
from app.services.excel import import_routing


@pytest.mark.parametrize('name', ['TAVLAMA', '2.TAVLAMA', '5.TAVLAMA', 'AĞIZ TAVLAMA', 'YIKAMA', '3.YIKAMA', 'ARA YIKAMA', '4.ARA YIKAMA'])
@pytest.mark.parametrize('basis', ['legacy_unspecified', 'labor_seconds_per_unit', 'machine_seconds_per_cycle'])
def test_proportional_conveyor_and_split_invariance(name, basis):
    op=RoutingOperation(operation_name=name, cycle_time_sec=600, setup_time_min=2, setup_machine_minutes=2, units_per_cycle=5, time_basis=basis, crew_size=1, machine_cycle_time_sec=600)
    result=compute_operation_need(op,6)
    assert result.labor_hours == pytest.approx(14/60)
    assert result.machine_hours == pytest.approx(14/60)
    assert mes_unit_hours(op)*6 == pytest.approx(12/60)
    assert sum(compute_operation_need(op,q,setup_required=False).labor_hours for q in [1,2,3]) == pytest.approx(12/60)


def test_default_and_unrelated_operations():
    op=RoutingOperation(operation_name='TAVLAMA',cycle_time_sec=600,setup_time_min=0,time_basis='legacy_unspecified')
    assert standard_unit_hours(op)==pytest.approx(10/60)
    op.operation_name='PRES'; op.units_per_cycle=5
    assert standard_unit_hours(op)==pytest.approx(10/60)


def test_daily_release_uses_proportional_piece_time():
    op=RoutingOperation(operation_name='ARA YIKAMA',cycle_time_sec=600,units_per_cycle=5,time_basis='machine_seconds_per_cycle')
    start=datetime(2026,9,21,8)
    seg=SimpleNamespace(segment_kind='process',start_at=start,end_at=start+timedelta(minutes=12),good_qty=6,is_locked=False)
    events=_process_outputs(op,[seg])
    assert sum(e[1] for e in events)==6
    assert events[-1][0]==seg.end_at


def test_weekly_capacity_places_six_units_in_twelve_minutes():
    from datetime import date
    from app.services.planning import _place_quantity
    from app.services.scenarios import Rule
    week=date(2026,9,21)
    op=RoutingOperation(id=1,seq=10,work_center_id=1,operation_name='2.TAVLAMA',cycle_time_sec=600,setup_time_min=0,units_per_cycle=5,time_basis='legacy_unspecified')
    item=SimpleNamespace(code='CONVEYOR',operations=[op])
    order=SimpleNamespace(id=1,order_no='TEST',item=item)
    lines,unplanned,_=_place_quantity(order,6,'TEST',None,{1:SimpleNamespace(code='TAVLAMA')},[week],{(1,week):12/60},SimpleNamespace(get=lambda *args:Rule(rule='finish')))
    assert not unplanned
    assert sum(l.planned_qty for l in lines)==pytest.approx(6)
    assert sum(l.planned_hours for l in lines)==pytest.approx(12/60)


def test_recipe_group_separate_values_and_validation(client,auth,db):
    wc=WorkCenter(code='CONVEYOR-TEST',name='Conveyor')
    items=[Item(code=f'CONVEYOR-{i}',name='Test',product_group='CONVEYOR-G' if i<2 else '') for i in range(3)]
    db.add_all([wc,*items]);db.flush()
    ops=[]
    for item in items:
        for seq,name in enumerate(['2.TAVLAMA','3.ARA YIKAMA','PRES'],1):
            op=RoutingOperation(item_id=item.id,work_center_id=wc.id,seq=seq,operation_name=name,cycle_time_sec=600,units_per_cycle=1)
            db.add(op);ops.append(op)
    db.commit()
    url=f'/api/items/{items[0].id}/conveyor-units'
    response=client.patch(url,headers=auth,json={'annealing_units':5,'washing_units':3,'apply_to_group':True})
    assert response.status_code==200,response.text
    assert response.json()['operation_count']==4
    db.expire_all()
    assert [o.units_per_cycle for o in ops]==[5,3,1,5,3,1,1,1,1]
    from io import BytesIO
    from openpyxl import load_workbook
    from app.services.excel import build_backup
    workbook=load_workbook(BytesIO(build_backup(db)),data_only=True)
    sheet=next(s for s in workbook if 'Çevrim Başına Adet' in [c.value for c in s[1]])
    headers=[c.value for c in sheet[1]]
    exported=[r for r in sheet.iter_rows(min_row=2,values_only=True) if r[0]==items[0].code]
    assert [r[headers.index('Çevrim Başına Adet')] for r in exported]==[5,3,1]
    for invalid in [0,-1,1.5,None]:
        assert client.patch(url,headers=auth,json={'annealing_units':invalid}).status_code==422
    assert client.patch(f'/api/items/{items[2].id}/conveyor-units',headers=auth,json={'apply_to_group':True}).status_code==400
    # Invalid import must not partially overwrite the existing cycle time.
    row={'_row':2,'item_code':items[0].code,'seq':1,'wc_code':wc.code,'cycle_time_sec':999,'units_per_cycle':2.5}
    _,_,errors=import_routing(db,[row]);assert errors
    assert ops[0].cycle_time_sec==600
    row['units_per_cycle']=7
    _,_,errors=import_routing(db,[row]);assert not errors
    assert ops[0].units_per_cycle==7
    db.rollback()


def test_production_bom_reimport_preserves_reviewed_conveyor_units(db):
    from collections import defaultdict
    from app.services.production_bom import ParsedFG, Branch, ParsedOp, import_parsed_fg
    op=ParsedOp(1,'5998765-02','5998765-02','Test','OP','TAVLAMA',10,'','','TAVLAMA')
    parsed=ParsedFG('6998765','Test',[Branch('5998765-02',1,[op])],'5998765-02')
    kwargs=dict(cache={},wc_idx={},machines={},counters=defaultdict(int),warnings=[])
    import_parsed_fg(db,parsed,**kwargs);db.flush()
    item=db.query(Item).filter_by(code='6998765').one()
    route=db.query(RoutingOperation).filter_by(item_id=item.id).one()
    route.units_per_cycle=5;db.commit();db.expunge_all()
    import_parsed_fg(db,parsed,cache={},wc_idx={},machines={},counters=defaultdict(int),warnings=[])
    db.flush()
    new_route=db.query(RoutingOperation).filter_by(item_id=item.id).one()
    assert new_route.units_per_cycle==5
    db.rollback()
