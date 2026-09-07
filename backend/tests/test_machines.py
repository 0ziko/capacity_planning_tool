"""Makineler ve kapasite kaynagi: is merkezi personeli vs makine atamalari."""

from datetime import date, timedelta

from tests.test_capacity_flow import _upload


def _monday() -> str:
    d = date.today()
    return (d - timedelta(days=d.weekday())).isoformat()


def _wc(client, auth, code):
    return next(w for w in client.get("/api/workcenters", headers=auth).json() if w["code"] == code)


def _cap(client, auth, wc_id) -> float:
    rows = client.get("/api/capacity", headers=auth, params={"start": _monday()}).json()
    return next(c for c in rows if c["work_center_id"] == wc_id)["capacity_hours"]


def test_machines_and_capacity_source(client, auth):
    # Alan bilgisiyle is merkezi; vardiya tanimi yok => Pzt-Cum, kisi x 4 saat
    _upload(
        client, auth, "workcenters",
        ["İş Merkezi Kodu", "İş Merkezi Adı", "Planlanıyor (E/H)", "Birim Saat", "Kişi Başı Verimli Saat", "Alan Kodu", "Alan Adı"],
        [["MAK-1", "Makineli İM", "E", 10, 4, "PRS", "PRESHANELER"]],
    )
    wc = _wc(client, auth, "MAK-1")
    assert wc["area_code"] == "PRS" and wc["area_name"] == "PRESHANELER"
    assert wc["capacity_source"] == "work_center"

    # 3 personel is merkezine bagli
    _upload(client, auth, "employees", ["Sicil No", "Ad Soyad", "İş Merkezi Kodu"], [["M-1", "A", "MAK-1"], ["M-2", "B", "MAK-1"], ["M-3", "C", "MAK-1"]])
    assert _cap(client, auth, wc["id"]) == 3 * 4 * 5

    # Makineler: API ile ekle + Excel ile ekle
    r = client.post(f"/api/workcenters/{wc['id']}/machines", headers=auth, json={"code": "MK-01", "name": "Pres 60t"})
    assert r.status_code == 201, r.text
    m1 = r.json()
    r = _upload(client, auth, "machines", ["İş Merkezi Kodu", "Makine Kodu", "Makine Adı"], [["MAK-1", "MK-02", "Pres 100t"]])
    assert r["inserted"] == 1 and not r["errors"]
    # ayni kod tekrar eklenemez
    assert client.post(f"/api/workcenters/{wc['id']}/machines", headers=auth, json={"code": "mk-01"}).status_code == 400
    wc = _wc(client, auth, "MAK-1")
    assert [m["code"] for m in wc["machines"]] == ["MK-01", "MK-02"]

    # Personelden ikisini makinelere ata (biri form, biri Excel)
    emps = {e["code"]: e for e in client.get("/api/employees", headers=auth, params={"work_center_id": wc["id"]}).json()}
    r = client.put(f"/api/employees/{emps['M-1']['id']}", headers=auth, json={**emps["M-1"], "machine_id": m1["id"]})
    assert r.status_code == 200 and r.json()["machine_code"] == "MK-01"
    r = _upload(client, auth, "employees", ["Sicil No", "Ad Soyad", "Makine Kodu"], [["M-2", "B", "MK-02"]])
    assert not r["errors"]
    wc = _wc(client, auth, "MAK-1")
    assert wc["employee_count"] == 3 and wc["machine_employee_count"] == 2
    assert next(m for m in wc["machines"] if m["code"] == "MK-01")["employee_count"] == 1

    # Kaynak is merkezi iken kapasite degismez (3 kisi)
    assert _cap(client, auth, wc["id"]) == 60

    # Makine detayini aktif et => yalnizca makinelere atanan 2 kisi
    body = {k: v for k, v in wc.items() if k not in ("id", "shifts", "machines", "employee_count", "machine_employee_count", "capacity_headcount")}
    r = client.put(f"/api/workcenters/{wc['id']}", headers=auth, json={**body, "capacity_source": "machines"})
    assert r.status_code == 200 and r.json()["capacity_headcount"] == 2
    assert _cap(client, auth, wc["id"]) == 2 * 4 * 5

    # Makine modunda vardiya kisi sayisi (10) yok sayilir; makine atamasi gecerlidir
    client.post(f"/api/workcenters/{wc['id']}/shifts", headers=auth, json={"name": "Gündüz", "weekdays": "0,1,2,3,4", "start_time": "08:00", "end_time": "18:00", "headcount": 10})
    assert _cap(client, auth, wc["id"]) == 40
    # Pasif makine sayilmaz
    m2 = next(m for m in wc["machines"] if m["code"] == "MK-02")
    client.put(f"/api/machines/{m2['id']}", headers=auth, json={"code": "MK-02", "name": "Pres 100t", "is_active": False})
    assert _cap(client, auth, wc["id"]) == 20

    # Kaynak geri "is merkezi" => vardiya kisi sayisi 10 gecerli
    r = client.put(f"/api/workcenters/{wc['id']}", headers=auth, json={**body, "capacity_source": "work_center"})
    assert r.status_code == 200
    assert _cap(client, auth, wc["id"]) == 10 * 4 * 5

    # Baska is merkezine ait makine atanamaz
    _upload(client, auth, "workcenters", ["İş Merkezi Kodu", "İş Merkezi Adı"], [["MAK-2", "Diğer"]])
    wc2 = _wc(client, auth, "MAK-2")
    r = client.put(f"/api/employees/{emps['M-3']['id']}", headers=auth, json={**emps["M-3"], "work_center_id": wc2["id"], "machine_id": m1["id"]})
    assert r.status_code == 400

    # Makine silinince personel atamasi dusur
    assert client.delete(f"/api/machines/{m1['id']}", headers=auth).status_code == 204
    e1 = next(e for e in client.get("/api/employees", headers=auth).json() if e["code"] == "M-1")
    assert e1["machine_id"] is None and e1["work_center_id"] == wc["id"]

    # Yedek dosyasinda Makineler sayfasi var
    from io import BytesIO

    from openpyxl import load_workbook

    r = client.get("/api/backup.xlsx", headers=auth)
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert "Makineler" in wb.sheetnames
