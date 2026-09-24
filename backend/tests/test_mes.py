from datetime import date, timedelta
from io import BytesIO
from uuid import uuid4
from zipfile import ZipFile, ZIP_DEFLATED
import pytest
from openpyxl import Workbook
from app.models import Item, BomLine, WorkCenter, WorkCenterWeek, Machine, RoutingOperation, Order, PlanLine, StockReceipt
from app.models.mes import MesDetail, MesPlanBaseline
from app.services import mes, mes_progress
from app.services.remaining_work import produced_qty_map

DAY = date(2026, 9, 14)


def workbook(rows):
    wb = Workbook()
    wb.active.append(["Üretim Detay Id", "Tarih", "Malzeme Kodu", "Net Üretilen Miktar", "İş Merkezi Kodu",
                      "Sağlam Adet", "Üretilen Miktar", "Başlangıç", "Müşteri", "Süre"])
    for row in rows:
        wb.active.append([*row, -999, 9999, "IGNORE", "IGNORE", 9999])
        for cell in wb.active[wb.active.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def case(db):
    token = uuid4().hex[:7]
    wc = WorkCenter(code="MES-" + token, name="MES", is_planned=True)
    db.add(wc); db.flush()
    machine = Machine(code="M-" + token, work_center_id=wc.id)
    db.add(machine)
    db.add(WorkCenterWeek(work_center_id=wc.id, week_start=DAY, headcount=10,
                          efficient_hours_per_person=4, working_days=5))
    items, ops, orders = [], [], []
    # Numeric FG/WIP codes, unique across this persistent test database.
    suffix = str(int(token, 16)).zfill(9)
    shared = "5" + suffix + "-01"
    for i in range(2):
        item = Item(code="6" + suffix + str(i), name="Test FG")
        db.add(item); db.flush()
        finish = "5" + suffix + str(i) + "-50"
        db.add(BomLine(item_id=item.id, component_code=shared, source_wip=finish, recipe_seq=10, quantity=1))
        first = RoutingOperation(item_id=item.id, seq=10, operation_name="Kesim", semi_finished_code=shared,
                                 work_center_id=wc.id, cycle_time_sec=3600)
        final = RoutingOperation(item_id=item.id, seq=20, operation_name="Montaj", semi_finished_code=finish,
                                 work_center_id=wc.id, cycle_time_sec=1800)
        db.add_all([first, final]); db.flush()
        order = Order(order_no=f"MES-{token}-{i}", customer="irrelevant", item_id=item.id,
                      due_date=DAY + timedelta(days=7 * (1-i)), quantity=100)
        db.add(order); db.flush()
        items.append(item); ops.append((first, final)); orders.append(order)
    db.commit()
    return dict(wc=wc, machine=machine, items=items, ops=ops, orders=orders, shared=shared)


def import_rows(db, case, rows):
    content = workbook([[id, day, code, qty, case["machine"].code] for id, day, code, qty in rows])
    p = mes.preview(db, content)
    result = mes.apply_import(db, content, p["token"], "test")
    db.commit()
    return result, content


def plan(db, c, i, wk, qty, final=False):
    op = c["ops"][i][1 if final else 0]
    p = PlanLine(order_id=c["orders"][i].id, operation_id=op.id, work_center_id=c["wc"].id,
                 week_start=wk, planned_qty=qty, planned_hours=qty * (0.5 if final else 1), mode="auto")
    db.add(p); db.commit()
    return p


def test_reader_uses_only_net_and_date_and_ignores_broken_styles():
    source = workbook([["X1", DAY, "500001-01", 12, "M1"]])
    out = BytesIO()
    with ZipFile(BytesIO(source)) as z, ZipFile(out, "w", ZIP_DEFLATED) as target:
        for name in z.namelist():
            target.writestr(name, b"<invalid/>" if name == "xl/styles.xml" else z.read(name))
    r = mes.read_mes(out.getvalue())[0]
    assert r["quantity"] == 12 and r["prod_date"] == DAY
    assert "customer" not in r


@pytest.mark.parametrize("qty", [-1, "NaN", "inf", ""])
def test_invalid_quantity_rejected(qty):
    with pytest.raises(ValueError):
        mes.read_mes(workbook([["x", DAY, "500001-01", qty, "M"]]))


def test_duplicate_id_rejects_whole_file():
    with pytest.raises(ValueError, match="tekrarlı"):
        mes.read_mes(workbook([["x", DAY, "500001-01", 1, "M"], ["x", DAY, "500001-01", 2, "M"]]))


def test_idempotent_details_and_amendment_no_daily_collapse(db, case):
    result, content = import_rows(db, case, [("a", DAY, case["shared"], 3), ("b", DAY, case["shared"], 4)])
    assert result["counts"]["new"] == 2 and result["counts"]["unresolved"] == 0
    assert mes.preview(db, content)["counts"]["unchanged"] == 2
    assert db.query(MesDetail).count() == 2
    import_rows(db, case, [("a", DAY, case["shared"], 6)])
    assert mes.balances([mes.detail_dict(d) for d in db.query(MesDetail).all()])[case["shared"]] == 10


def test_finished_receipt_and_consumption_update_once(db, case):
    fg = case["items"][0].code
    import_rows(db, case, [("w", DAY, case["shared"], 10), ("f", DAY, fg, 4)])
    receipt = db.query(StockReceipt).filter_by(source="mes").one()
    assert receipt.quantity == 4
    result, _ = import_rows(db, case, [("f", DAY, fg, 6)])
    assert db.query(StockReceipt).filter_by(source="mes").count() == 1
    assert receipt.quantity == 6
    assert mes.balances([mes.detail_dict(d) for d in db.query(MesDetail).all()])[case["shared"]] == 4
    assert result["counts"]["updated"] == 1


def test_stale_preview_cannot_overwrite_changed_record(db, case):
    import_rows(db, case, [("a", DAY, case["shared"], 1)])
    target = workbook([["a", DAY, case["shared"], 9, case["machine"].code]])
    p = mes.preview(db, target)
    import_rows(db, case, [("a", DAY, case["shared"], 2)])
    with pytest.raises(ValueError, match="Önizlemeyi"):
        mes.apply_import(db, target, p["token"], "test")


def test_ambiguous_times_preserve_stock_without_earned_hours(db, case):
    case["ops"][1][0].cycle_time_sec = 7200
    db.commit()
    p, _ = import_rows(db, case, [("a", DAY, case["shared"], 10)])
    assert p["counts"]["free_stock"] == 1
    assert p["standard_hours"] == 0
    assert mes.balances([mes.detail_dict(d) for d in db.query(MesDetail).all()])[case["shared"]] == 10


def test_weekly_plan_early_output_and_frozen_baseline(db, case):
    current = plan(db, case, 0, DAY, 5)
    plan(db, case, 1, DAY + timedelta(weeks=1), 10)
    import_rows(db, case, [("a", DAY, case["shared"], 12), ("b", DAY + timedelta(days=1), case["shared"], 1)])
    r = mes_progress.progress(db, DAY)
    assert r["summary"]["actual_hours"] == 12
    assert r["summary"]["matched_hours"] == 5
    assert r["summary"]["early_hours"] == 7
    assert r["summary"]["capacity_usage_pct"] == 30
    assert r["daily"][1]["hours"] == 0
    current.planned_qty = 999
    db.commit()
    assert mes_progress.progress(db, DAY)["operations"][0]["quantity"] == 5
    future = mes_progress.progress(db, DAY + timedelta(weeks=1))
    assert future["operations"][0]["early_qty"] == 8
    assert future["operations"][0]["remaining_qty"] == 2
    assert future["summary"]["actual_hours"] == 0


def test_common_pool_prioritizes_finished_week_not_customer_due_date(db, case):
    # Earlier due order is item 1, but item 0 reaches finished goods this week.
    plan(db, case, 0, DAY, 30, final=True)
    plan(db, case, 1, DAY + timedelta(weeks=1), 40, final=True)
    import_rows(db, case, [("a", DAY, case["shared"], 50)])
    allocations = []
    credit, _ = produced_qty_map(db, include_mes=True, as_of=DAY, mes_allocations=allocations)
    assert credit[(case["orders"][0].id, case["ops"][0][0].id)] == 30
    assert credit[(case["orders"][1].id, case["ops"][1][0].id)] == 20
    assert sum(a["quantity"] for a in allocations) == 50
    assert allocations[0]["finish_week"] == DAY.isoformat()


def test_split_finish_plan_cannot_take_all_stock_before_other_fg(db, case):
    plan(db, case, 0, DAY, 10, final=True)
    plan(db, case, 0, DAY + timedelta(weeks=2), 90, final=True)
    plan(db, case, 1, DAY + timedelta(weeks=1), 40, final=True)
    import_rows(db, case, [("a", DAY, case["shared"], 60)])
    allocations = []
    produced_qty_map(db, include_mes=True, as_of=DAY, mes_allocations=allocations)
    assert [(a["finished_item_code"], a["quantity"]) for a in allocations] == [
        (case["items"][0].code, 10), (case["items"][1].code, 40), (case["items"][0].code, 10)]


def test_no_finish_plan_leaves_common_pool_unassigned(db, case):
    import_rows(db, case, [("a", DAY, case["shared"], 50)])
    credits, _ = produced_qty_map(db, include_mes=True, as_of=DAY)
    assert sum(credits.values()) == 0


def test_api_preview_does_not_write_and_confirmation_required(client, auth, db, case):
    content = workbook([["api1", DAY, case["shared"], 7, case["machine"].code]])
    files = {"file": ("mes.xlsx", content)}
    response = client.post("/api/mes/preview", files=files, headers=auth)
    assert response.status_code == 200, response.text
    assert db.query(MesDetail).count() == 0
    assert client.post("/api/mes/import", files=files, headers=auth).status_code == 422
    result = client.post("/api/mes/import", files=files, params={"token": response.json()["token"]}, headers=auth)
    assert result.status_code == 200, result.text
    assert db.query(MesDetail).count() == 1
    assert client.get("/api/mes/progress", params={"as_of": DAY}, headers=auth).status_code == 200


def test_finished_consumption_cannot_be_credited_to_second_product(db, case):
    plan(db, case, 1, DAY, 100, final=True)
    import_rows(db, case, [("w", DAY, case["shared"], 50), ("f", DAY, case["items"][0].code, 40)])
    allocations = []
    credits, _ = produced_qty_map(db, include_mes=True, as_of=DAY, mes_allocations=allocations)
    assert credits[(case["orders"][0].id, case["ops"][0][1].id)] == 40
    assert sum(a["quantity"] for a in allocations) == 10
    assert allocations[0]["finished_item_code"] == case["items"][1].code


def test_batch_credits_whole_batch_not_only_anchor_order(db, case):
    from app.models import ProductionBatch, ProductionBatchOrder
    first = case["orders"][0]
    other = Order(order_no="MES-extra-" + uuid4().hex[:6], item_id=first.item_id, due_date=DAY + timedelta(days=20), quantity=100)
    batch = ProductionBatch(batch_no="MES-batch-" + uuid4().hex[:6], item_id=first.item_id, due_date=DAY, quantity=200)
    db.add_all([other, batch]); db.flush()
    db.add_all([ProductionBatchOrder(batch_id=batch.id, order_id=first.id, quantity=100),
                ProductionBatchOrder(batch_id=batch.id, order_id=other.id, quantity=100)])
    p = plan(db, case, 0, DAY, 200, final=True)
    p.production_batch_id = batch.id
    db.commit()
    import_rows(db, case, [("a", DAY, case["shared"], 150)])
    allocations = []
    credits, _ = produced_qty_map(db, include_mes=True, as_of=DAY, mes_allocations=allocations)
    assert credits[(first.id, case["ops"][0][0].id)] == 150
    assert sum(a["quantity"] for a in allocations) == 150


def test_standard_hours_are_linear_for_multi_unit_machine_cycles(db, case):
    for first, _ in case["ops"]:
        first.time_basis = "machine_seconds_per_cycle"
        first.units_per_cycle = 4
        first.crew_size = 2
    db.commit()
    p, _ = import_rows(db, case, [("a", DAY, case["shared"], 3)])
    # 3 units × (3600 sec / 4 units × 2 people) = 1.5 standard hours.
    assert p["standard_hours"] == 1.5


def test_receipt_reduction_cannot_break_reservations(client, auth, db, case):
    from app.models import Reservation
    import_rows(db, case, [("f", DAY, case["items"][0].code, 10)])
    db.add(Reservation(item_id=case["items"][0].id, order_id=case["orders"][0].id, quantity=8))
    db.commit()
    content = workbook([["f", DAY, case["items"][0].code, 5, case["machine"].code]])
    p = mes.preview(db, content)
    response = client.post("/api/mes/import", files={"file": ("mes.xlsx", content)}, params={"token": p["token"]}, headers=auth)
    assert response.status_code == 409
    db.expire_all()
    assert db.get(MesDetail, "f").quantity == 10
    assert db.query(StockReceipt).filter_by(source="mes").one().quantity == 10


def test_export_is_filtered_and_identifiers_are_literal(client, auth, db, case):
    from openpyxl import load_workbook
    import_rows(db, case, [("=1+1", DAY, case["shared"], 7)])
    response = client.get("/api/mes/progress.xlsx", params={"as_of": DAY, "work_center_ids": case["wc"].id}, headers=auth)
    assert response.status_code == 200, response.text
    wb = load_workbook(BytesIO(response.content))
    assert wb["Plan dışı"]["A2"].value == "=1+1"
    assert wb["Plan dışı"]["A2"].data_type == "s"
    assert wb["Günlük"].max_row == 2
    assert wb["Günlük"]["B2"].value == 7


def test_sunday_included_and_missing_capacity_has_no_fake_percentage(db, case):
    sunday = DAY + timedelta(days=6)
    import_rows(db, case, [("a", sunday, case["shared"], 2)])
    r = mes_progress.progress(db, sunday)
    assert r["daily"][6]["hours"] == 2
    assert r["summary"]["output_vs_plan_pct"] is None


@pytest.mark.parametrize("unknown_machine", [False, True])
def test_mes_outside_route_credits_planned_center_once(db, case, unknown_machine, monkeypatch):
    from app.core.config import get_settings
    from app.services.mes_actuals import measure, weekly_kpis
    monkeypatch.setattr(get_settings(), "production_source", "mes")
    actual_wc = WorkCenter(code="ACTUAL-"+uuid4().hex[:7], name="Actual", is_planned=True)
    db.add(actual_wc); db.flush()
    if unknown_machine:
        case["machine"].code="UNKNOWN-"+uuid4().hex[:7]
        code=case["machine"].code
        db.delete(case["machine"])
    else:
        case["machine"].work_center_id=actual_wc.id
        code=case["machine"].code
    db.commit()
    slot=plan(db,case,0,DAY,100)
    plan(db,case,0,DAY,100,final=True)
    content=workbook([["outside",DAY,case["shared"],30,code]])
    preview=mes.preview(db,content)
    mapping=preview["rows"][0]["mapping"]
    assert preview["counts"]["unresolved"]==0
    assert mapping["resource_deviation"] is True
    assert mapping["work_center_id"]==case["wc"].id
    assert preview["standard_hours"]==30
    mes.apply_import(db,content,preview["token"],"test");db.commit()
    assert mes.preview(db,content)["counts"]["unchanged"]==1
    assert measure(db,DAY)["matches"][slot.id]["qty"]==30
    kpis=weekly_kpis(db,[case["wc"].id,actual_wc.id],DAY,DAY,DAY)
    assert kpis[(case["wc"].id,DAY)]["plan_adherence_remaining_hours"]==120
    assert (actual_wc.id,DAY) not in kpis
    assert produced_qty_map(db,as_of=DAY)[0][(case["orders"][0].id,case["ops"][0][0].id)]==30
    assert case["ops"][0][0].work_center_id==case["wc"].id


def test_outside_route_does_not_guess_conflicting_standard_times(db,case):
    case["ops"][1][0].cycle_time_sec=7200
    db.commit()
    content=workbook([["ambiguous-outside",DAY,case["shared"],20,"UNKNOWN"]])
    result=mes.preview(db,content)
    assert result["counts"]["free_stock"]==1
    assert "standart süre" in result["rows"][0]["mapping"]["reason"]
    assert result["standard_hours"]==0


def test_outside_route_finished_goods_still_create_single_receipt(db,case):
    content=workbook([["outside-fg",DAY,case["items"][0].code,12,"UNLISTED"]])
    result=mes.preview(db,content)
    assert result["counts"]["unresolved"]==0
    mes.apply_import(db,content,result["token"],"test");db.commit()
    again=mes.preview(db,content)
    mes.apply_import(db,content,again["token"],"test");db.commit()
    receipts=db.query(StockReceipt).filter_by(source="mes").all()
    assert len(receipts)==1 and receipts[0].quantity==12


@pytest.mark.parametrize("conflicting_parent", [False, True])
def test_wip_route_uses_parent_bom_without_orders(db, conflicting_parent):
    token=str(int(uuid4().hex[:8],16))
    wc=WorkCenter(code="PARENT-"+token,name="Weld",is_planned=True)
    wip=Item(code="5"+token+"-15",name="WIP")
    parents=[Item(code="6"+token+str(i),name="FG") for i in range(2)]
    db.add_all([wc,wip,*parents]);db.flush()
    machine=Machine(code="PARENT-M-"+token,work_center_id=wc.id)
    code="5"+token+"-23"
    previous="5"+token+"-22"
    op=RoutingOperation(item_id=wip.id,seq=20,operation_name="Weld",semi_finished_code=code,
                        work_center_id=wc.id,cycle_time_sec=900)
    db.add_all([machine,op])
    for i,parent in enumerate(parents):
        db.add_all([BomLine(item_id=parent.id,component_code=previous,source_wip=wip.code,recipe_seq=10,quantity=3 if conflicting_parent and i else 2),
                    BomLine(item_id=parent.id,component_code=code,source_wip=wip.code,recipe_seq=20,quantity=1)])
    db.commit()
    assert not wip.bom_lines
    content=workbook([["parent-"+token,DAY,code,4,machine.code]])
    result=mes.preview(db,content)
    mapping=result["rows"][0]["mapping"]
    if conflicting_parent:
        assert result["counts"]["pending_consumption"]==1
        assert result["rows"][0]["preview_category"] == "pending"
        assert mapping["status"] == "mapped" and mapping["inputs"] == {}
        assert mapping["input_candidates"] == [{previous: 2}, {previous: 3}]
        assert result["standard_hours"] == 1
        assert "tüketim" in mapping["reason"]
        return
    assert mapping["status"]=="mapped"
    assert mapping["inputs"]=={previous:2}
    assert mapping["candidates"]==sorted(p.code for p in parents)
    assert result["standard_hours"]==1
    mes.apply_import(db,content,result["token"],"test");db.commit()
    assert mes.preview(db,content)["counts"]["unchanged"]==1
    pool=mes.balances([mes.detail_dict(d) for d in db.query(MesDetail).all()])
    assert pool[code]==4 and pool[previous]==-8


def test_first_wip_operation_uses_parent_bom_without_fake_predecessor(db):
    token=str(int(uuid4().hex[:8],16))
    wc=WorkCenter(code="FIRST-"+token,name="Weld")
    wip=Item(code="5"+token+"-15",name="WIP")
    fg=Item(code="6"+token,name="FG")
    db.add_all([wc,wip,fg]);db.flush()
    code="5"+token+"-23"
    db.add(BomLine(item_id=fg.id,component_code=code,source_wip=wip.code,recipe_seq=10,quantity=1))
    db.add(RoutingOperation(item_id=wip.id,seq=10,operation_name="Weld",semi_finished_code=code,work_center_id=wc.id,cycle_time_sec=904.32))
    db.commit()
    result=mes.preview(db,workbook([["first-"+token,DAY,code,5,"UNLISTED"]]))
    assert result["counts"]["unresolved"]==0
    assert result["rows"][0]["mapping"]["inputs"]=={}
    assert result["standard_hours"]==pytest.approx(1.256)


@pytest.mark.parametrize("conflict", [False, True])
def test_successive_wip_route_preserves_pieces_despite_parent_component_ratios(db,conflict):
    token=str(int(uuid4().hex[:8],16))
    wc=WorkCenter(code="ASM-"+token,name="Press")
    wip=Item(code="5"+token+"-11",name="Assembly WIP")
    parents=[Item(code="6"+token+str(i),name="FG") for i in range(2)]
    db.add_all([wc,wip,*parents]);db.flush()
    previous="5"+token+"-03"
    db.add_all([RoutingOperation(item_id=wip.id,seq=10,operation_name="Forma",semi_finished_code=previous,work_center_id=wc.id,cycle_time_sec=76.8),
                RoutingOperation(item_id=wip.id,seq=20,operation_name="Etek Kesme",semi_finished_code=wip.code,work_center_id=wc.id,cycle_time_sec=48)])
    for i,parent in enumerate(parents):
        quantity=(i+1)*2
        db.add_all([BomLine(item_id=parent.id,component_code=wip.code,source_wip=wip.code,recipe_seq=0,quantity=quantity),
                    BomLine(item_id=parent.id,component_code=previous,source_wip=wip.code,recipe_seq=10,quantity=quantity*(3 if conflict and i else 2))])
    db.commit()
    content=workbook([["asm-"+token,DAY,wip.code,10,"UNLISTED"]])
    result=mes.preview(db,content)
    mapping=result["rows"][0]["mapping"]
    assert mapping["status"]=="mapped"
    assert mapping["inputs"]=={previous:1}
    assert result["standard_hours"]==pytest.approx(10*48/3600)
    assert mapping["candidates"]==sorted(p.code for p in parents)
    mes.apply_import(db,content,result["token"],"test");db.commit()
    pool=mes.balances([mes.detail_dict(d) for d in db.query(MesDetail).all()])
    assert pool[wip.code]==10 and pool[previous]==-10
    assert mes.preview(db,content)["counts"]["unchanged"]==1



def test_preview_maps_repeated_material_machine_once_but_refreshes_next_request(db,case,monkeypatch):
    original=mes.Mapper.map
    calls=[]
    def counted(self,row):
        calls.append((row["material_code"],row["machine_code"]))
        return original(self,row)
    monkeypatch.setattr(mes.Mapper,"map",counted)
    content=workbook([[f"repeat-{i}",DAY,case["shared"],1,case["machine"].code] for i in range(50)])
    first=mes.preview(db,content)
    assert len(calls)==1 and first["standard_hours"]==50
    for op,_ in case["ops"]:op.cycle_time_sec=1800
    db.commit()
    second=mes.preview(db,content)
    assert len(calls)==2 and second["standard_hours"]==25
    assert first["token"]!=second["token"]


def test_scoped_mapper_preserves_indirect_parents_and_ignores_unrelated_catalog(db):
    token=str(int(uuid4().hex[:8],16))
    wc=WorkCenter(code="SCOPE-"+token,name="Scope")
    wip=Item(code="5"+token+"-11",name="WIP")
    fg=Item(code="6"+token,name="FG")
    unrelated=Item(code="6"+token+"9",name="Unrelated")
    db.add_all([wc,wip,fg,unrelated]);db.flush()
    code="5"+token+"-03"
    db.add_all([BomLine(item_id=fg.id,component_code=wip.code,source_wip=wip.code,recipe_seq=0,quantity=1),
        BomLine(item_id=wip.id,component_code="100-test",source_wip=wip.code,recipe_seq=11,quantity=1),
        RoutingOperation(item_id=wip.id,seq=10,operation_name="Forma",semi_finished_code=code,work_center_id=wc.id,cycle_time_sec=60)])
    db.commit()
    row={"material_code":code,"machine_code":"UNLISTED"}
    full=mes.Mapper(db)
    scoped=mes.Mapper(db,{code},{"UNLISTED"})
    assert unrelated.code not in scoped.items
    assert scoped.map(row)==full.map(row)
    assert scoped.map(row)["candidates"]==[fg.code]



def test_unknown_material_is_free_stock_without_production_credit_and_idempotent(db,case,client,auth):
    code="5"+str(int(uuid4().hex[:8],16))+"-13"
    content=workbook([["free-a",DAY,code,6,case["machine"].code],["free-b",DAY,code,7,case["machine"].code]])
    p=mes.preview(db,content)
    assert p["counts"]["free_stock"]==2 and p["counts"]["unresolved"]==0
    assert p["standard_hours"]==0 and mes.free_stock_rows(db)==[]
    mes.apply_import(db,content,p["token"],"test");db.commit()
    again=mes.preview(db,content)
    mes.apply_import(db,content,again["token"],"test");db.commit()
    rows=client.get("/api/mes/free-stock",headers=auth).json()
    assert len(rows)==1 and rows[0]["quantity"]==13
    assert db.query(StockReceipt).filter_by(source="mes").count()==0
    assert produced_qty_map(db,include_mes=True)[0]=={}
    report=mes_progress.progress(db,DAY)
    assert report["summary"]["actual_hours"]==0 and report["unresolved"]==[]
    assert report["free_stock"][0]["quantity"]==13
    changed=workbook([["free-a",DAY,code,2,case["machine"].code]])
    mes.apply_import(db,changed,mes.preview(db,changed)["token"],"test");db.commit()
    assert mes.free_stock_rows(db)[0]["quantity"]==9
    assert mes.free_stock_rows(db,DAY-timedelta(days=1))==[]
    from app.services.mes_export import export_report
    from openpyxl import load_workbook
    wb=load_workbook(BytesIO(export_report(mes_progress.progress(db,DAY))))
    assert wb["Tanım bekleyen serbest stok"]["B2"].value==9


def test_free_stock_moves_to_mapped_pool_once_only_after_definition_and_reimport(db,case):
    code="5"+str(int(uuid4().hex[:8],16))+"-13"
    content=workbook([["free-remap",DAY,code,6,case["machine"].code]])
    mes.apply_import(db,content,mes.preview(db,content)["token"],"test");db.commit()
    item=Item(code=code,name="Now defined")
    db.add(item);db.commit()
    # Partially created master data must not lose already accepted free stock.
    mes.apply_import(db,content,mes.preview(db,content)["token"],"test");db.commit()
    assert mes.free_stock_rows(db)[0]["quantity"]==6
    db.add_all([BomLine(item_id=case["items"][0].id,component_code=code,source_wip=code,recipe_seq=10,quantity=1),
                RoutingOperation(item_id=item.id,seq=10,operation_name="New",semi_finished_code=code,work_center_id=case["wc"].id,cycle_time_sec=60)])
    db.commit()
    db.expire_all()  # Import preview is a new request after master-data edits.
    p=mes.preview(db,content)
    assert p["rows"][0]["mapping"]["status"]=="mapped"
    assert mes.free_stock_rows(db)[0]["quantity"]==6
    mes.apply_import(db,content,p["token"],"test");db.commit()
    assert mes.free_stock_rows(db)==[]
    assert mes.balances([mes.detail_dict(d) for d in db.query(MesDetail).all()])[code]==6


def test_known_wip_with_missing_route_preserves_output_in_free_stock(db,case):
    code="5"+str(int(uuid4().hex[:8],16))+"-13"
    db.add(Item(code=code,name="Known but incomplete"));db.commit()
    p=mes.preview(db,workbook([["known-no-route",DAY,code,6,"UNKNOWN"]]))
    assert p["counts"]["unresolved"]==0 and p["counts"]["free_stock"]==1



def test_wip_route_predecessor_wins_over_later_bom_component_and_nested_use(db):
    token=str(int(uuid4().hex[:8],16))
    wc=WorkCenter(code="PIECES-"+token,name="Heat")
    code="5"+token+"-02";previous="5"+token+"-01";component="5"+token+"-19"
    wip=Item(code=code,name="WIP");fg=Item(code="6"+token,name="FG")
    db.add_all([wc,wip,fg]);db.flush()
    db.add_all([RoutingOperation(item_id=wip.id,seq=10,operation_name="Form",semi_finished_code=previous,work_center_id=wc.id,cycle_time_sec=60),
        RoutingOperation(item_id=wip.id,seq=20,operation_name="Heat",semi_finished_code=code,work_center_id=wc.id,cycle_time_sec=186.24),
        BomLine(item_id=fg.id,component_code=code,source_wip=code,recipe_seq=0,quantity=2),
        BomLine(item_id=fg.id,component_code=previous,source_wip=code,recipe_seq=10,quantity=2),
        BomLine(item_id=fg.id,component_code=component,source_wip=code,recipe_seq=21,quantity=3),
        BomLine(item_id=fg.id,component_code=code,source_wip="5999999-23",recipe_seq=11,quantity=1)])
    db.commit()
    content=workbook([["pieces-in",DAY,previous,4,"UNLISTED"],["pieces-out",DAY,code,4,"UNLISTED"]])
    p=mes.preview(db,content)
    output=next(r for r in p["rows"] if r["material_code"]==code)
    assert output["mapping"]["status"]=="mapped"
    assert output["mapping"]["inputs"]=={previous:1}
    assert output["mapping"]["standard_unit_hours"]==pytest.approx(186.24/3600)
    pool=mes.balances(p["rows"])
    assert pool[previous]==0 and pool[code]==4
    assert component not in output["mapping"]["inputs"]


@pytest.mark.parametrize("qty", [0, 7])
def test_preview_free_stock_is_not_consumption_reconciliation(db, case, qty):
    code = "5" + str(int(uuid4().hex[:8], 16)) + "-13"
    content = workbook([["unknown-review", DAY, code, qty, "UNKNOWN"],
                        ["normal-review", DAY, case["shared"], 2, case["machine"].code]])
    result = mes.preview(db, content)
    assert result["counts"]["free_stock"] == 1
    assert result["counts"]["pending_consumption"] == 0
    assert [r["preview_category"] for r in result["rows"]] == ["free_stock", "mapped"]
    assert result["rows"][0]["mapping"]["consumption_status"] == "pending"
    mes.apply_import(db, content, result["token"], "test"); db.commit()
    assert mes.preview(db, content)["counts"]["unchanged"] == 2


def test_preview_categories_separate_conflicting_time_and_consumption():
    assert mes.preview_category({"status":"mapped", "consumption_status":"pending"}) == "pending"
    assert mes.preview_category({"status":"free_stock", "consumption_status":"pending",
                                 "input_candidates":[{"A":1},{"B":1}]}) == "free_stock"
    assert mes.preview_category({"status":"unresolved", "consumption_status":"pending"}) == "unresolved"
    assert mes.preview_category({"status":"mapped"}) == "mapped"


@pytest.mark.parametrize("referenced", [False, True])
def test_unlinked_erp_root_cannot_shadow_linked_output_route(db, case, referenced):
    token = str(int(uuid4().hex[:8], 16))
    root = Item(code="5"+token, name="Old root", product_group="ERP")
    output = root.code+"-02"
    real = Item(code=output, name="Actual output", product_group="ERP")
    db.add_all([root, real]); db.flush()
    for item, previous in [(root,root.code+"-01"),(real,root.code+"-19")]:
        db.add_all([RoutingOperation(item_id=item.id,seq=10,operation_name="Before",semi_finished_code=previous,work_center_id=case["wc"].id,cycle_time_sec=60),
                    RoutingOperation(item_id=item.id,seq=20,operation_name="Heat",semi_finished_code=output,work_center_id=case["wc"].id,cycle_time_sec=120),
                    BomLine(item_id=item.id,component_code="1000000",source_wip=item.code,quantity=1)])
    db.add(BomLine(item_id=case["items"][0].id,component_code=output,source_wip=output,quantity=1,recipe_seq=0))
    if referenced:
        db.add(BomLine(item_id=case["items"][1].id,component_code=root.code,source_wip=root.code,quantity=1,recipe_seq=0))
    db.commit()
    result=mes.preview(db,workbook([["root-check",DAY,output,9,case["machine"].code]]))
    mapping=result["rows"][0]["mapping"]
    assert mapping["status"]=="mapped"
    assert mapping["consumption_status"] == ("pending" if referenced else "known")
    if not referenced:
        assert mapping["inputs"]=={root.code+"-19":1}
        assert root.code in mapping["reason"]
    assert db.get(Item,root.id) is not None
    assert len(db.get(Item,root.id).operations)==2
