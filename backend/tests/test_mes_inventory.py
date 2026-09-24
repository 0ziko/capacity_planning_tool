from datetime import timedelta
from io import BytesIO
from openpyxl import load_workbook
from app.services.mes_inventory import inventory, planning_pool, report
from app.services import mes, mes_actuals
from app.models.mes import MesDetail
from tests.test_mes import DAY, case, workbook, import_rows, plan


def record(identity, day, code, qty, inputs=None, pending=None, free=False):
    mapping = {"status": "free_stock" if free else "mapped", "kind": "wip", "inputs": inputs or {},
               "standard_unit_hours": 0 if free else 0.5}
    if pending is not None:
        mapping.update(consumption_status="pending", input_candidates=pending, reason="Unknown predecessor")
    return dict(detail_id=identity, prod_date=day, material_code=code, quantity=qty, machine_code="M", mapping=mapping)


def by_code(result):
    return {r["material_code"]: r for r in result["rows"]}


def test_stock_accumulates_across_weeks_and_historical_cutoff():
    rows = [record(str(i), DAY + timedelta(days=i), "A", 100) for i in range(5)]
    rows.append(record("later", DAY + timedelta(days=8), "B", 180, {"A": 1}))
    assert by_code(inventory(rows, DAY + timedelta(days=4)))["A"]["balance"] == 500
    result = inventory(rows)
    assert result["opening_status"] == "zero_confirmed"
    assert by_code(result)["A"]["balance"] == 320
    assert by_code(result)["B"]["balance"] == 180


def test_five_operations_same_day_are_independent_of_file_order():
    rows = [record(str(i), DAY, str(i), 100-i*10, {str(i-1): 1} if i else {}) for i in range(5)]
    forward = inventory(rows)
    reverse = inventory(list(reversed(rows)))
    assert forward == reverse
    assert [r["balance"] for r in forward["rows"]] == [10, 10, 10, 10, 60]
    assert sum(m["standard_hours"] for m in forward["movements"]) == 200
    assert all(m["standard_hours"] == 0 for m in forward["movements"] if m["kind"] == "consumption")


def test_pending_inputs_hold_alternatives_but_output_can_be_consumed_today():
    rows = [record("a",DAY,"A",100),record("b",DAY,"B",100),
            record("c",DAY,"C",40,pending=[{"A":1},{"B":2}]),record("d",DAY,"D",30,{"C":1})]
    result = inventory(rows)
    balances = by_code(result)
    assert balances["A"]["balance"] == 100 and balances["B"]["balance"] == 100
    assert planning_pool(rows) == {"A":60,"B":20,"C":10,"D":30}
    assert len(result["pending"]) == 1
    # Replacing an amended MES detail removes tentative holds; no extra physical output.
    rows[2] = record("c",DAY,"C",40,{"A":1})
    assert planning_pool(rows) == {"A":60,"B":100,"C":10,"D":30}
    assert inventory(rows)["pending"] == []


def test_unknown_stock_consumed_in_same_import_and_amended_once(db, case, client, auth):
    code = "599999987-01"
    free = record("unknown-input", DAY, code, 100, free=True, pending=[])
    known = record("known-output", DAY, case["shared"], 70, {code:1})
    # Persist accepted mapping snapshots, just as import does.
    for r in (free, known):
        db.add(MesDetail(**r, updated_by="test"))
    db.commit()
    assert mes.free_stock_rows(db)[0]["quantity"] == 30
    response = client.get("/api/mes/inventory",params={"as_of":DAY},headers=auth)
    assert response.status_code == 200
    assert by_code(response.json())[code]["balance"] == 30
    exported = client.get("/api/mes/inventory.xlsx",params={"as_of":DAY},headers=auth)
    assert exported.status_code == 200
    book = load_workbook(BytesIO(exported.content))
    assert book.sheetnames == ["Bilgi","Bakiyeler","Hareketler","Belirsiz tüketimler"]
    detail = db.get(MesDetail,"known-output")
    detail.quantity = 90
    db.commit()
    assert mes.free_stock_rows(db)[0]["quantity"] == 10
    assert len(report(db)["movements"]) == 3


def test_import_consumption_does_not_add_capacity_hours_and_reimport_is_idempotent(db, case):
    plan(db,case,0,DAY,100)
    plan(db,case,0,DAY,100,final=True)
    rows=[("raw",DAY,case["shared"],100),("finished",DAY,case["items"][0].code,40)]
    result, content = import_rows(db,case,rows)
    assert result["standard_hours"] == 120
    kpi = mes_actuals.weekly_kpis(db,[case["wc"].id],DAY,DAY,DAY)[(case["wc"].id,DAY)]
    assert kpi["standard_hour_equivalent_output"] == 120
    assert kpi["plan_adherence_remaining_hours"] == 30
    assert by_code(report(db))[case["shared"]]["balance"] == 60
    mes.apply_import(db,content,mes.preview(db,content)["token"],"test");db.commit()
    assert mes_actuals.weekly_kpis(db,[case["wc"].id],DAY,DAY,DAY)[(case["wc"].id,DAY)] == kpi
    assert len(report(db)["movements"]) == 2
