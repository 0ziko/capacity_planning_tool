from datetime import date
from io import BytesIO
from openpyxl import load_workbook
from app.schemas import OrderScheduleOut
from app.services import orders, excel


def test_schedule_export_matches_filter_sort_and_columns(client, auth, monkeypatch):
    def row(oid, status, end, late):
        return OrderScheduleOut(order_id=oid, order_no=f"S-{oid}", position_no=str(oid),
            customer="Müşteri", item_code="600001", item_name="Ürün", quantity=10,
            due_date=date(2026, 9, 30), required_hours=10, planned_hours=10,
            coverage_pct=100, planned_end=end, lateness_days=late, plan_status=status,
            material_note="Koşullu plan: malzeme doğrulanmamış.")
    rows = [row(1, "late", date(2026, 10, 2), 2), row(2, "on_time", date(2026, 9, 29), -1),
            row(3, "late", date(2026, 10, 5), 5), row(4, "no_ops", None, None)]
    seen = []
    def schedule(db, ids):
        seen.append(ids)
        return list(rows)
    monkeypatch.setattr(orders, "order_schedule", schedule)
    response = client.get("/api/plan/orders.xlsx?work_center_ids=7&plan_status=late&sort=late", headers=auth)
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content))["Sipariş Bitiş Tarihleri"]
    values = list(sheet.values)
    assert seen == [[7]]
    assert [r[0] for r in values[1:]] == ["S-3", "S-1"]
    assert values[0][1:5] == ("Poz", "Müşteri", "Stok Kodu", "Stok Adı")
    assert values[1][1] == "3" and values[1][-2] == "Geç kalacak"
    assert values[1][-1] == rows[2].material_note
    assert len(values[0]) == len(values[1]) == 18
    empty = client.get("/api/plan/orders.xlsx?plan_status=partial", headers=auth)
    assert load_workbook(BytesIO(empty.content))["Sipariş Bitiş Tarihleri"].max_row == 1
    assert client.get("/api/plan/orders.xlsx?sort=invalid", headers=auth).status_code == 422
    assert client.get("/api/plan/orders.xlsx").status_code in (401, 403)
