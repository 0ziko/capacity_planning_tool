"""Tekil siparis CRUD, plan sonucu bitis tarihleri, siparis ilerlemesi ve birlestirme."""

import io
from datetime import date, timedelta

from tests.test_capacity_flow import _upload, _xlsx

WEEK = date(2026, 9, 7)  # Pazartesi


def _setup(client, auth):
    # test veritabani oturum boyunca paylasilir: onceki testlerin siparislerini temizle
    for st in ("open", "closed", "merged"):
        client.delete("/api/orders", headers=auth, params={"status": st})
    _upload(client, auth, "workcenters", ["İş Merkezi Kodu", "İş Merkezi Adı", "Planlanıyor (E/H)", "Birim Saat", "Kişi Başı Verimli Saat"], [["TZG-A", "A Tezgahı", "E", 10, 4], ["MON-1", "Montaj", "E", 10, 4]])
    _upload(client, auth, "shifts", ["İş Merkezi Kodu", "Vardiya", "Günler (Pzt=0..Paz=6)", "Başlangıç", "Bitiş", "Kişi Sayısı", "Kişi Başı Verimli Saat"], [["TZG-A", "Gündüz", "0,1,2,3,4", "08:00", "18:00", 10, 4], ["MON-1", "Gündüz", "0,1,2,3,4", "08:00", "18:00", 5, 4]])
    _upload(client, auth, "items", ["Stok Kodu", "Stok Adı", "Ürün Grubu"], [["MAM-1", "Ocak", "OCAK"]])
    # 36 sn kesim + 72 sn montaj => 1000 adet: 10 saat + 20 saat
    _upload(client, auth, "routing", ["Stok Kodu", "Sıra", "Operasyon", "İş Merkezi Kodu", "Çevrim Süresi (sn)"], [["MAM-1", 10, "Kesim", "TZG-A", 36], ["MAM-1", 20, "Montaj", "MON-1", 72]])


def test_single_order_crud(client, auth):
    _setup(client, auth)
    r = client.post("/api/orders", headers=auth, json={"order_no": "S-100", "customer": "ABC", "due_date": "2026-09-30", "item_code": "mam-1", "quantity": 1000})
    assert r.status_code == 201, r.text
    o = r.json()
    assert o["item_code"] == "MAM-1" and o["item_name"] == "Ocak" and o["status"] == "open"

    # ayni siparis no + stok tekrar eklenemez
    assert client.post("/api/orders", headers=auth, json={"order_no": "S-100", "due_date": "2026-09-30", "item_code": "MAM-1", "quantity": 5}).status_code == 400
    # bilinmeyen stok
    assert client.post("/api/orders", headers=auth, json={"order_no": "S-101", "due_date": "2026-09-30", "item_code": "YOK", "quantity": 5}).status_code == 400

    r = client.put(f"/api/orders/{o['id']}", headers=auth, json={"order_no": "S-100", "customer": "ABC Ltd", "due_date": "2026-10-05", "item_code": "MAM-1", "quantity": 1200, "note": "revize"})
    assert r.status_code == 200 and r.json()["quantity"] == 1200 and r.json()["customer"] == "ABC Ltd"

    assert client.delete(f"/api/orders/{o['id']}", headers=auth).status_code == 204
    assert client.get("/api/orders", headers=auth).json() == []


def test_schedule_progress_and_merge(client, auth):
    _setup(client, auth)
    for no, cust, due, qty in [("SP-1", "Musteri A", "2026-09-18", 1000), ("SP-2", "Musteri B", "2026-09-25", 500), ("SP-3", "Musteri C", "2026-10-02", 250)]:
        assert client.post("/api/orders", headers=auth, json={"order_no": no, "customer": cust, "due_date": due, "item_code": "MAM-1", "quantity": qty}).status_code == 201

    # plan yokken: unplanned
    sched = client.get("/api/plan/orders", headers=auth).json()
    assert len(sched) == 3 and all(s["plan_status"] == "unplanned" for s in sched)

    r = client.post("/api/plan/auto", headers=auth, json={"start_week": WEEK.isoformat(), "weeks": 6}).json()
    assert r["unplanned"] == []
    sched = client.get("/api/plan/orders", headers=auth).json()
    s1 = next(s for s in sched if s["order_no"] == "SP-1")
    assert s1["required_hours"] == 30 and s1["planned_hours"] == 30 and s1["coverage_pct"] == 100
    assert s1["planned_start"] == WEEK.isoformat()
    assert s1["planned_end"] is not None and s1["plan_status"] == "on_time"
    # plan durumu filtresi: planlanan = on_time + partial + late
    planned = client.get("/api/orders", headers=auth, params={"status": "open", "plan_status": "planned"}).json()
    assert any(o["order_no"] == "SP-1" for o in planned)
    unplanned = client.get("/api/orders", headers=auth, params={"status": "open", "plan_status": "unplanned"}).json()
    assert all(o["order_no"] != "SP-1" for o in unplanned)
    assert s1["lateness_days"] <= 0
    # tumu ilk haftaya sigar (52.5 saat < 200 / 100 kapasite); bitis gunu hafta icinde
    assert WEEK <= date.fromisoformat(s1["planned_end"]) <= WEEK + timedelta(days=4)

    # siparis ilerlemesi: S-1 icin 500 adet kesim (siparis no ile), 250 adet montaj (siparis no'suz -> FIFO S-1)
    _upload(client, auth, "production", ["Tarih", "İş Merkezi Kodu", "Stok Kodu", "Operasyon Sıra", "Sipariş No", "Miktar"], [[WEEK.isoformat(), "TZG-A", "MAM-1", 10, "SP-1", 500], [WEEK.isoformat(), "MON-1", "MAM-1", 20, "", 250]])
    prog = client.get("/api/progress/orders", headers=auth, params={"as_of": WEEK.isoformat()}).json()
    p1 = next(p for p in prog if p["order_no"] == "SP-1")
    assert p1["status"] == "in_progress"
    assert p1["ops"][0]["produced_qty"] == 500 and p1["ops"][0]["pct"] == 50
    assert p1["ops"][1]["produced_qty"] == 250 and p1["ops"][1]["pct"] == 25
    assert p1["produced_qty"] == 250  # son operasyondan cikan
    assert p1["earned_hours"] == 10  # 5 saat kesim + 5 saat montaj
    assert p1["last_prod_date"] == WEEK.isoformat()
    p2 = next(p for p in prog if p["order_no"] == "SP-2")
    assert p2["status"] == "not_started"
    r = client.get("/api/progress/orders.xlsx", headers=auth)
    assert r.status_code == 200 and r.content[:2] == b"PK"

    # birlestirme onerisi: 3 siparis ayni stok (terminler 7+ gun arayla -> opsiyonel grup)
    groups = client.get("/api/plan/merge-suggestions", headers=auth, params={"tolerance_days": 5}).json()
    assert len(groups) == 1 and groups[0]["recommended"] is False and groups[0]["order_count"] == 3
    rec = client.get("/api/plan/merge-suggestions", headers=auth, params={"tolerance_days": 30}).json()
    assert any(g["recommended"] and g["order_count"] == 3 for g in rec)
    g = groups[0]
    assert g["item_code"] == "MAM-1" and g["order_count"] == 3 and g["total_qty"] == 1750 and g["has_progress"] is True
    assert g["customers"] == ["Musteri A", "Musteri B", "Musteri C"]

    # SP-2 + SP-3 uretim partisi (siparisler acik kalir)
    ids = [o["id"] for o in g["orders"] if o["order_no"] in ("SP-2", "SP-3")]
    r = client.post("/api/plan/merge", headers=auth, json={"order_ids": ids})
    assert r.status_code == 201, r.text
    batch = r.json()
    assert batch["quantity"] == 750 and batch["due_date"] == "2026-09-25"
    assert batch["batch_no"].startswith("URT-MAM-1-")
    assert {o["order_no"] for o in batch["orders"]} == {"SP-2", "SP-3"}
    open_nos = sorted(o["order_no"] for o in client.get("/api/orders", headers=auth).json())
    assert open_nos == ["SP-1", "SP-2", "SP-3"]
    assert client.get("/api/orders", headers=auth, params={"status": "merged"}).json() == []
    batches = client.get("/api/plan/production-batches", headers=auth).json()
    assert len(batches) == 1 and batches[0]["batch_no"] == batch["batch_no"]
    # partideki siparis duzenlenebilir
    assert client.put(f"/api/orders/{ids[0]}", headers=auth, json={"order_no": "SP-2", "due_date": "2026-09-25", "item_code": "MAM-1", "quantity": 500}).status_code == 200
    # tek siparis partiye alinamaz
    assert client.post("/api/plan/merge", headers=auth, json={"order_ids": [ids[0]]}).status_code == 422

    # partiyi dagit
    r = client.delete(f"/api/plan/merge/{batch['id']}", headers=auth)
    assert r.status_code == 200 and r.json()["reopened"] == 2
    open_nos = sorted(o["order_no"] for o in client.get("/api/orders", headers=auth).json())
    assert open_nos == ["SP-1", "SP-2", "SP-3"]
    assert client.get("/api/plan/production-batches", headers=auth).json() == []

    # plan excel'i yeni sayfayla acilir
    r = client.get(f"/api/plan/export.xlsx?start={WEEK}", headers=auth)
    assert r.status_code == 200 and r.content[:2] == b"PK"


def test_user_role_cannot_write_orders(client, auth):
    _setup(client, auth)
    client.post("/api/users", headers=auth, json={"username": "izleyici2", "password": "123456", "role": "user"})
    tok = client.post("/api/auth/login", data={"username": "izleyici2", "password": "123456"}).json()["access_token"]
    h = {"Authorization": f"Bearer {tok}"}
    assert client.post("/api/orders", headers=h, json={"order_no": "X", "due_date": "2026-09-30", "item_code": "MAM-1", "quantity": 1}).status_code == 403
    assert client.post("/api/plan/merge", headers=h, json={"order_ids": [1, 2]}).status_code == 403
    assert client.get("/api/plan/merge-suggestions", headers=h).status_code == 200


def test_order_position_no_import(client, auth):
    """Ayni siparis no + farkli poz ve farkli stok; poz bos ise geriye uyum (siparis+stok)."""
    _setup(client, auth)
    _upload(client, auth, "items", ["Stok Kodu", "Stok Adı", "Ürün Grubu"], [["MAM-2", "Davlumbaz", "DAV"]])
    _upload(client, auth, "routing", ["Stok Kodu", "Sıra", "Operasyon", "İş Merkezi Kodu", "Çevrim Süresi (sn)"], [["MAM-2", 10, "Kesim", "TZG-A", 36]])

    r = _upload(
        client,
        auth,
        "orders",
        ["Sipariş No", "Poz No", "Müşteri", "Termin", "Stok Kodu", "Miktar", "Birim Fiyat"],
        [
            ["SP-MULTI", "10", "Musteri A", "2026-09-20", "MAM-1", 100, 10],
            ["SP-MULTI", "20", "Musteri A", "2026-09-20", "MAM-2", 50, 20],
            ["SP-LEG", "", "Musteri B", "2026-09-25", "MAM-1", 30, 5],
        ],
    )
    assert r["inserted"] == 3 and r["updated"] == 0 and r["errors"] == []

    rows = client.get("/api/orders", headers=auth).json()
    assert len(rows) == 3
    p10 = next(o for o in rows if o["order_no"] == "SP-MULTI" and o["position_no"] == "10")
    p20 = next(o for o in rows if o["order_no"] == "SP-MULTI" and o["position_no"] == "20")
    assert p10["item_code"] == "MAM-1" and p10["quantity"] == 100
    assert p20["item_code"] == "MAM-2" and p20["quantity"] == 50

    assert client.post("/api/orders", headers=auth, json={"order_no": "SP-MULTI", "position_no": "10", "due_date": "2026-09-20", "item_code": "MAM-2", "quantity": 1}).status_code == 400

    r2 = _upload(
        client,
        auth,
        "orders",
        ["Sipariş No", "Poz No", "Müşteri", "Termin", "Stok Kodu", "Miktar"],
        [["SP-MULTI", "10", "Musteri A", "2026-09-20", "MAM-1", 120]],
    )
    assert r2["updated"] == 1
    p10 = next(o for o in client.get("/api/orders", headers=auth).json() if o["position_no"] == "10")
    assert p10["quantity"] == 120

    stock = client.get("/api/stock/orders", headers=auth, params={"position": "20"}).json()
    assert len(stock) == 1 and stock[0]["position_no"] == "20" and stock[0]["item_code"] == "MAM-2"


def test_orders_import_preview_and_remove_missing(client, auth):
    """Onizleme farklari; listede olmayan acik siparisler istege bagli silinir."""
    _setup(client, auth)
    client.post("/api/orders", headers=auth, json={"order_no": "KEEP", "customer": "A", "due_date": "2026-09-30", "item_code": "MAM-1", "quantity": 10})
    client.post("/api/orders", headers=auth, json={"order_no": "DROP", "due_date": "2026-09-30", "item_code": "MAM-1", "quantity": 5})
    client.post("/api/orders", headers=auth, json={"order_no": "CLOSED", "due_date": "2026-09-30", "item_code": "MAM-1", "quantity": 3})
    closed = next(o for o in client.get("/api/orders", headers=auth).json() if o["order_no"] == "CLOSED")
    client.patch(f"/api/orders/{closed['id']}/status", headers=auth, params={"status": "closed"})

    content = _xlsx(
        ["Sipariş No", "Poz No", "Müşteri", "Termin", "Stok Kodu", "Miktar"],
        [["KEEP", "", "A", "2026-09-30", "MAM-1", 10], ["NEW-1", "1", "B", "2026-10-01", "MAM-1", 7]],
    )
    prev = client.post("/api/imports/orders/preview", headers=auth, files={"file": ("orders.xlsx", content, "application/octet-stream")}).json()
    assert prev["parse_errors"] == []
    assert prev["error_rows"] == []
    assert {r["order_no"] for r in prev["only_in_system"]} == {"DROP"}
    assert {r["order_no"] for r in prev["only_in_file"]} == {"NEW-1"}
    assert prev["unchanged_count"] == 1

    r = client.post(
        "/api/imports/orders",
        headers=auth,
        params={"remove_missing": True},
        files={"file": ("orders.xlsx", content, "application/octet-stream")},
    ).json()
    assert r["inserted"] == 1 and r["removed"] == 1 and r["errors"] == []

    open_nos = sorted(o["order_no"] for o in client.get("/api/orders", headers=auth, params={"status": "open"}).json())
    assert open_nos == ["KEEP", "NEW-1"]
    assert client.get("/api/orders", headers=auth, params={"status": "closed"}).json()[0]["order_no"] == "CLOSED"


def test_orders_import_preview_xlsx(client, auth):
    """Onizleme raporu Excel: hatali stok kodu, termin YYYY-MM-DD, duzenle sayfasi."""
    from datetime import datetime

    from openpyxl import load_workbook

    _setup(client, auth)
    content = _xlsx(
        ["Sipariş No", "Poz No", "Müşteri", "Termin", "Stok Kodu", "Miktar"],
        [
            ["KEEP", "", "A", datetime(2026, 5, 21, 16, 27, 8), "MAM-1", 10],
            ["BAD", "", "B", datetime(2026, 8, 27, 17, 0, 35), "YOK-STOK", 5],
        ],
    )
    prev = client.post("/api/imports/orders/preview", headers=auth, files={"file": ("orders.xlsx", content, "application/octet-stream")}).json()
    assert len(prev["error_rows"]) == 1
    assert prev["error_rows"][0]["item_code"] == "YOK-STOK"
    assert "YOK-STOK" in prev["missing_item_codes"]

    r = client.post("/api/imports/orders/preview.xlsx", headers=auth, files={"file": ("orders.xlsx", content, "application/octet-stream")})
    assert r.status_code == 200 and r.content[:2] == b"PK"
    wb = load_workbook(io.BytesIO(r.content), data_only=True)
    ws = wb["Düzenle ve yükle"]
    assert ws.cell(2, 4).value == "2026-05-21"
    assert ws.cell(3, 4).value == "2026-08-27"
    assert "T" not in str(ws.cell(2, 4).value)

    # onizleme raporunu tekrar yukle: duzenle sayfasi okunur
    r2 = client.post("/api/imports/orders/preview", headers=auth, files={"file": ("rapor.xlsx", r.content, "application/octet-stream")}).json()
    assert r2["file_row_count"] == 2
    assert r2["parse_errors"] == [] or len(r2["error_rows"]) == 1
