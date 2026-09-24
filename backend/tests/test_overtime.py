"""Fazla mesai (18:00-21:00, kişi başı ≤ 2,5 sa): kapasite, kısıtlar, Excel gidiş-dönüş, revizyon önerisi ve onayı."""
import io
import time
from datetime import date, timedelta

from openpyxl import load_workbook

from tests.test_capacity_flow import _plan_with_ack, _upload, _weekly_staffing, _xlsx

WEEK = date(2026, 9, 7)
HDR = ["İş Merkezi Kodu", "Hafta", "Kişi Sayısı", "Kişi Başı Verimli Saat", "Çalışma Günü", "Not", "Fazla Mesai Kişi", "Fazla Mesai Gün", "Fazla Mesai Saat (kişi/gün, ≤2,5)"]


def _master(client, auth, code="OT-WC"):
    # 10 kişi × 4 verimli saat × 5 gün = 200 sa; vardiya 08-18 (10 sa nominal) -> verim oranı 0,4
    _upload(client, auth, "workcenters", ["İş Merkezi Kodu", "İş Merkezi Adı", "Planlanıyor (E/H)", "Birim Saat", "Kişi Başı Verimli Saat"], [[code, "Fazla mesai WC", "E", 10, 4]])
    _upload(client, auth, "shifts", ["İş Merkezi Kodu", "Vardiya", "Günler (Pzt=0..Paz=6)", "Başlangıç", "Bitiş", "Kişi Sayısı", "Kişi Başı Verimli Saat"], [[code, "G", "0,1,2,3,4", "08:00", "18:00", 10, 4]])
    _weekly_staffing(client, auth, [[code, 10, 4, 5]])
    _upload(client, auth, "items", ["Stok Kodu", "Stok Adı", "Ürün Grubu"], [["OT-P", "Fazla mesai ürünü", "OT"]])
    _upload(client, auth, "routing", ["Stok Kodu", "Sıra", "Operasyon", "İş Merkezi Kodu", "Çevrim Süresi (sn)"], [["OT-P", 10, "Op", code, 3600]])
    return next(w for w in client.get("/api/workcenters", headers=auth).json() if w["code"] == code)["id"]


def _cap(client, auth, code, wk):
    return next(c for c in client.get("/api/capacity", headers=auth, params={"start": wk.isoformat()}).json() if c["work_center_code"] == code)["capacity_hours"]


def test_overtime_capacity_and_limits(client, auth):
    wc_id = _master(client, auth)
    assert _cap(client, auth, "OT-WC", WEEK) == 200
    # 4 kişi × 2,5 sa × 0,4 verim × 3 gün = 12 sa ek kapasite
    r = client.put(f"/api/workcenters/{wc_id}/weeks/{WEEK.isoformat()}", headers=auth,
                   json={"headcount": 10, "efficient_hours_per_person": 4, "working_days": 5, "overtime_headcount": 4, "overtime_days": 3, "overtime_hours_per_person": 2.5})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["overtime_headcount"] == 4 and body["overtime_days"] == 3 and body["overtime_efficiency_ratio"] == 0.4
    assert body["overtime_capacity_hours"] == 12 and body["base_capacity_hours"] == 200 and body["capacity_hours"] == 212
    assert _cap(client, auth, "OT-WC", WEEK) == 212
    load = client.get("/api/plan/load", headers=auth, params={"start": WEEK.isoformat(), "weeks": 1, "work_center_ids": [wc_id]}).json()
    assert load[0]["weeks"][0]["capacity_hours"] == 212 and load[0]["weeks"][0]["overtime_hours"] == 12
    # Kısıtlar: kişi > haftanın kişi sayısı, gün > çalışma günü, saat > 2,5
    for bad in ({"overtime_headcount": 11}, {"overtime_headcount": 2, "overtime_days": 6}, {"overtime_headcount": 2, "overtime_hours_per_person": 3}):
        rr = client.put(f"/api/workcenters/{wc_id}/weeks/{WEEK.isoformat()}", headers=auth, json={"headcount": 10, "efficient_hours_per_person": 4, "working_days": 5, **bad})
        assert rr.status_code in (400, 422), (bad, rr.text)
    # Gün verilmezse çalışma günleri (5): 2 kişi × 2,5 × 0,4 × 5 = 10
    r = client.put(f"/api/workcenters/{wc_id}/weeks/{WEEK.isoformat()}", headers=auth, json={"headcount": 10, "efficient_hours_per_person": 4, "working_days": 5, "overtime_headcount": 2})
    assert r.json()["overtime_capacity_hours"] == 10 and r.json()["overtime_days"] == 5
    # Fazla mesai kapatılınca 200'e döner
    r = client.put(f"/api/workcenters/{wc_id}/weeks/{WEEK.isoformat()}", headers=auth, json={"headcount": 10, "efficient_hours_per_person": 4, "working_days": 5})
    assert r.json()["capacity_hours"] == 200 and r.json()["overtime_headcount"] == 0


def test_overtime_excel_round_trip(client, auth):
    wc_id = _master(client, auth, "OT-XL")
    wk = WEEK + timedelta(weeks=1)
    res = _upload(client, auth, "wc_weeks", HDR, [["OT-XL", wk.isoformat(), 10, 4, 5, "", 3, 2, 2.5]])
    assert res["errors"] == []
    assert _cap(client, auth, "OT-XL", wk) == 200 + 3 * 2.5 * 0.4 * 2
    bad = client.post("/api/imports/wc_weeks", headers=auth, files={"file": ("w.xlsx", _xlsx(HDR, [["OT-XL", wk.isoformat(), 10, 4, 5, "", 12, 2, 2.5]]), "application/octet-stream")}).json()
    assert bad["errors"] and "asamaz" in bad["errors"][0]
    bad = client.post("/api/imports/wc_weeks", headers=auth, files={"file": ("w.xlsx", _xlsx(HDR, [["OT-XL", wk.isoformat(), 10, 4, 5, "", 2, 2, 3.0]]), "application/octet-stream")}).json()
    assert bad["errors"] and "2,5" in bad["errors"][0]
    exp = client.get("/api/exports/wc-weeks.xlsx", headers=auth, params={"start": wk.isoformat(), "weeks": 1, "work_center_ids": [wc_id]})
    ws = load_workbook(io.BytesIO(exp.content))["Haftalık İş Gücü"]
    row = [c.value for c in ws[2]]
    assert row[6:9] == [3, 2, 2.5]
    tmpl = client.get("/api/imports/template/wc_weeks", headers=auth)
    assert tmpl.status_code == 200 and "Fazla Mesai" in str(load_workbook(io.BytesIO(tmpl.content)).active[1][6].value)


def test_revision_overtime_suggestion_and_approval(client, auth):
    wc_id = _master(client, auth, "OT-RV")
    for st in ("open", "closed", "merged"):
        client.delete("/api/orders", headers=auth, params={"status": st})
    # 1 haftalık ufuk, 200 sa kapasite; 230 sa talep -> 30 sa plansız (kapasite yetersiz)
    a = client.post("/api/orders", headers=auth, json={"order_no": "OT-A", "due_date": (WEEK + timedelta(days=4)).isoformat(), "item_code": "OT-P", "quantity": 200}).json()
    b = client.post("/api/orders", headers=auth, json={"order_no": "OT-B", "due_date": (WEEK + timedelta(days=4)).isoformat(), "item_code": "OT-P", "quantity": 30}).json()
    assert a["id"] and b["id"]
    body = {"start_week": WEEK.isoformat(), "weeks": 1, "work_center_ids": [wc_id], "replace_existing": True, "use_overtime": False}
    assert _plan_with_ack(client, "/api/plan/auto", headers=auth, json=body).status_code == 200
    rev = client.post("/api/plan/revisions", headers=auth, json={"reason_codes": ["overtime_labor"], "note": "darboğaz", "start_week": WEEK.isoformat(), "weeks": 1, "work_center_ids": [wc_id], "mode": "due_date", "use_overtime": False}).json()  # elle öneri akışı: otomatik FM kapalı
    assert client.get(f"/api/plan/revisions/{rev['id']}/overtime-suggestions", headers=auth).status_code in (400, 409)  # önce hesap
    calc = client.post(f"/api/plan/revisions/{rev['id']}/calculate", headers=auth).json()
    assert calc["compare"]["proposed"]["unplanned"] + calc["compare"]["proposed"]["partial"] >= 1
    sug = client.get(f"/api/plan/revisions/{rev['id']}/overtime-suggestions", headers=auth)
    assert sug.status_code == 200, sug.text
    s = sug.json()
    assert s["total_shortfall_hours"] > 0 and s["max_hours_per_person"] == 2.5 and s["window"] == "18:00-21:00"
    assert len(s["suggestions"]) == 1
    row = s["suggestions"][0]
    # 30 sa açık: kişi başı hafta = 2,5 × 0,4 × 5 = 5 sa -> 6 kişi; sınır 10 kişi
    assert row["suggested_overtime_headcount"] == 6 and row["overtime_days"] == 5 and row["added_capacity_hours"] == 30 and row["shortfall_hours_after"] == 0
    assert row["work_center_code"] == "OT-RV" and "OT-B" in row["order_nos"] or "OT-A" in row["order_nos"]
    # Öneriyi taslağa ekle -> hesapla -> plansız kalmaz -> onay haftalık iş gücüne yazar
    changes = [
        {"entity_type": "wc_week", "extra_key": f"{wc_id}|{WEEK.isoformat()}", "field": "overtime_headcount", "new_value": str(row["suggested_overtime_headcount"])},
        {"entity_type": "wc_week", "extra_key": f"{wc_id}|{WEEK.isoformat()}", "field": "overtime_days", "new_value": str(row["overtime_days"])},
        {"entity_type": "wc_week", "extra_key": f"{wc_id}|{WEEK.isoformat()}", "field": "overtime_hours_per_person", "new_value": "2.5"},
    ]
    r = client.post(f"/api/plan/revisions/{rev['id']}/changes/bulk", headers=auth, json={"changes": changes})
    assert r.status_code == 200, r.text
    too_many = client.post(f"/api/plan/revisions/{rev['id']}/changes", headers=auth, json={"entity_type": "wc_week", "extra_key": f"{wc_id}|{WEEK.isoformat()}", "field": "overtime_hours_per_person", "new_value": "3"})
    assert too_many.status_code in (400, 409)
    calc2 = client.post(f"/api/plan/revisions/{rev['id']}/calculate", headers=auth).json()
    assert calc2["compare"]["proposed"]["unplanned"] == 0 and calc2["compare"]["proposed"]["planned_hours"] == 230
    assert _cap(client, auth, "OT-RV", WEEK) == 200  # canlı değişmedi
    ap = client.post(f"/api/plan/revisions/{rev['id']}/approve", headers=auth)
    assert ap.status_code == 200, ap.text
    assert _cap(client, auth, "OT-RV", WEEK) == 230
    wk = client.get(f"/api/workcenters/{wc_id}/weeks", headers=auth, params={"start": WEEK.isoformat(), "weeks": 1}).json()[0]
    assert wk["ov_overtime_headcount"] == 6 and wk["overtime_capacity_hours"] == 30
    # Fazla mesai kişi sayısı haftanın kişi sayısını aşan taslak hesapta reddedilir
    rev2 = client.post("/api/plan/revisions", headers=auth, json={"reason_codes": ["overtime_labor"], "note": "", "start_week": WEEK.isoformat(), "weeks": 1, "work_center_ids": [wc_id], "mode": "due_date"}).json()
    client.post(f"/api/plan/revisions/{rev2['id']}/changes", headers=auth, json={"entity_type": "wc_week", "extra_key": f"{wc_id}|{WEEK.isoformat()}", "field": "overtime_headcount", "new_value": "11"})
    bad = client.post(f"/api/plan/revisions/{rev2['id']}/calculate", headers=auth)
    assert bad.status_code in (400, 409) and "asamaz" in bad.text
    client.post(f"/api/plan/revisions/{rev2['id']}/cancel", headers=auth)


def test_daily_schedule_uses_overtime_window(db):
    """Günlük çizelge: fazla mesai olan günde makine aralığı 18:00-20:30'u da kapsar."""
    from datetime import datetime, time
    from app.models import Machine, WorkCenter, WorkCenterShift, WorkCenterWeek
    from app.services.calendar_capacity import machine_work_intervals
    from app.services.capacity import Overrides

    wc = WorkCenter(code="OT-DAILY", name="d", is_planned=True)
    db.add(wc)
    db.flush()
    m = Machine(work_center_id=wc.id, code="OT-DAILY-M1", is_active=True)
    db.add(WorkCenterShift(work_center_id=wc.id, name="G", weekdays="0,1,2,3,4", start_time=time(8, 0), end_time=time(18, 0), headcount=2))
    db.add(WorkCenterWeek(work_center_id=wc.id, week_start=WEEK, headcount=2, efficient_hours_per_person=4, working_days=5, overtime_headcount=2, overtime_days=2))
    db.add(m)
    db.commit()
    db.refresh(m)
    ovl = Overrides(db, wc.id)
    monday = machine_work_intervals(db, m, WEEK, ovl)
    ends = sorted(r.end for r in monday)
    assert ends[-1] == datetime.combine(WEEK, time(20, 30))
    wednesday = machine_work_intervals(db, m, WEEK + timedelta(days=2), ovl)  # overtime_days=2 -> Çarşamba yok
    assert max(r.end for r in wednesday) == datetime.combine(WEEK + timedelta(days=2), time(18, 0))


HDR_WE = HDR + ["Hafta Sonu FM Kişi", "Hafta Sonu FM Gün (0-2)", "Hafta Sonu FM Saat (kişi/gün, ≤8,5)"]


def test_weekend_overtime_capacity_limits_and_caps(client, auth, monkeypatch):
    """Cmt/Paz 08:00-18:00: kişi × 8,5 × verim × gün; kısıtlar; aylık yumuşak sınır; öneri/onay işareti."""
    wc_id = _master(client, auth, "OT-WE")
    url = f"/api/workcenters/{wc_id}/weeks/{WEEK.isoformat()}"
    base = {"headcount": 10, "efficient_hours_per_person": 4, "working_days": 5}
    # 4 kişi × 8,5 × 0,4 × 2 gün = 27,2 sa; hafta içi 2 kişi × 2,5 × 0,4 × 5 = 10 sa
    r = client.put(url, headers=auth, json={**base, "overtime_headcount": 2, "weekend_overtime_headcount": 4, "weekend_overtime_days": 2})
    assert r.status_code == 200, r.text
    b = r.json()
    assert b["weekend_overtime_days"] == 2 and b["weekend_overtime_hours_per_person"] == 8.5
    assert b["weekend_overtime_capacity_hours"] == 27.2 and b["overtime_capacity_hours"] == 37.2 and b["capacity_hours"] == 237.2
    assert b["overtime_person_hours_week"] == 5 * 2.5 + 2 * 8.5 and b["overtime_person_hours_ytd"] == 29.5 and b["overtime_proposed"] is False
    assert _cap(client, auth, "OT-WE", WEEK) == 237.2
    # Günlük kapasite takviminde Cumartesi/Pazar görünür
    days = next(c for c in client.get("/api/capacity", headers=auth, params={"start": WEEK.isoformat()}).json() if c["work_center_code"] == "OT-WE")["days"]
    assert {d["day"] for d in days} >= {(WEEK + timedelta(days=5)).isoformat(), (WEEK + timedelta(days=6)).isoformat()}
    # Yalnızca Cumartesi
    r = client.put(url, headers=auth, json={**base, "weekend_overtime_headcount": 4, "weekend_overtime_days": 1})
    assert r.json()["weekend_overtime_capacity_hours"] == 13.6 and r.json()["capacity_hours"] == 213.6
    # Kısıtlar: kişi > haftanın kişi sayısı, gün > 2, saat > 8,5, 6 günlük haftada 2 hafta sonu günü
    for bad in ({"weekend_overtime_headcount": 11}, {"weekend_overtime_headcount": 2, "weekend_overtime_days": 3}, {"weekend_overtime_headcount": 2, "weekend_overtime_hours_per_person": 9},
                {"working_days": 6, "weekend_overtime_headcount": 2, "weekend_overtime_days": 2}):
        rr = client.put(url, headers=auth, json={**base, **bad})
        assert rr.status_code in (400, 422), (bad, rr.text)
    # Yıl içi birikim: ikinci haftada 29,5 + 17 = 46,5
    wk2 = WEEK + timedelta(weeks=1)
    r = client.put(url, headers=auth, json={**base, "overtime_headcount": 2, "weekend_overtime_headcount": 2, "weekend_overtime_days": 2})
    r2 = client.put(f"/api/workcenters/{wc_id}/weeks/{wk2.isoformat()}", headers=auth, json={**base, "weekend_overtime_headcount": 2, "weekend_overtime_days": 2})
    assert r2.json()["overtime_person_hours_week"] == 17 and r2.json()["overtime_person_hours_ytd"] == 46.5
    # Aylık yumuşak sınır (ayar): 40 sa → ikinci hafta reddedilir; kapalıyken (0) serbest
    from app.core.config import get_settings
    monkeypatch.setattr(get_settings(), "overtime_monthly_cap_hours", 40.0)
    rr = client.put(f"/api/workcenters/{wc_id}/weeks/{wk2.isoformat()}", headers=auth, json={**base, "weekend_overtime_headcount": 2, "weekend_overtime_days": 2})
    assert rr.status_code == 400 and "aylik" in rr.text
    monkeypatch.setattr(get_settings(), "overtime_monthly_cap_hours", 0.0)
    # Öneri işareti: plan önerisi olarak yazılır, onaylanınca kalkar
    r = client.put(url, headers=auth, json={**base, "weekend_overtime_headcount": 2, "overtime_proposed": True})
    assert r.json()["overtime_proposed"] is True
    r = client.put(url, headers=auth, json={**base, "weekend_overtime_headcount": 2, "overtime_proposed": False})
    assert r.json()["overtime_proposed"] is False
    # Excel gidiş-dönüş (hafta sonu sütunları)
    wk3 = WEEK + timedelta(weeks=2)
    res = _upload(client, auth, "wc_weeks", HDR_WE, [["OT-WE", wk3.isoformat(), 10, 4, 5, "", 0, "", "", 3, 2, 8.5]])
    assert res["errors"] == [], res
    assert _cap(client, auth, "OT-WE", wk3) == 200 + 3 * 8.5 * 0.4 * 2
    bad = client.post("/api/imports/wc_weeks", headers=auth, files={"file": ("w.xlsx", _xlsx(HDR_WE, [["OT-WE", wk3.isoformat(), 10, 4, 5, "", 0, "", "", 3, 3, 8.5]]), "application/octet-stream")}).json()
    assert bad["errors"] and "0–2" in bad["errors"][0]
    # Revizyon taslağı hafta sonu alanını taşır ve sınırı denetler
    rev = client.post("/api/plan/revisions", headers=auth, json={"reason_codes": ["overtime_labor"], "note": "", "start_week": wk3.isoformat(), "weeks": 1, "work_center_ids": [wc_id], "mode": "due_date"}).json()
    ok = client.post(f"/api/plan/revisions/{rev['id']}/changes", headers=auth, json={"entity_type": "wc_week", "extra_key": f"{wc_id}|{wk3.isoformat()}", "field": "weekend_overtime_headcount", "new_value": "5"})
    assert ok.status_code == 200, ok.text
    too = client.post(f"/api/plan/revisions/{rev['id']}/changes", headers=auth, json={"entity_type": "wc_week", "extra_key": f"{wc_id}|{wk3.isoformat()}", "field": "weekend_overtime_hours_per_person", "new_value": "9"})
    assert too.status_code in (400, 409, 422)
