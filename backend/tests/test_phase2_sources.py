from datetime import timedelta
from io import BytesIO
from uuid import uuid4
from openpyxl import load_workbook
import pytest
from test_mes import case, plan, import_rows, DAY
from app.core.config import get_settings
from app.models import ProductionActual, Reservation, Order, ProductionBatch, ProductionBatchOrder, StockReceipt
from app.services.remaining_work import produced_qty_map, required_qty_for_batch
from app.services.mes_actuals import measure, weekly_kpis
from app.services.orders import order_progress
from app.services.stock import sync_progress_receipts


def legacy(db, c, qty):
    db.add(ProductionActual(prod_date=DAY, work_center_id=c["wc"].id, item_id=c["items"][0].id,
        operation_seq=10, order_no=c["orders"][0].order_no, quantity=qty, earned_hours=qty))
    db.commit()


def test_sources_never_added_and_import_does_not_activate(db, case, monkeypatch, client, auth):
    settings=get_settings()
    monkeypatch.setattr(settings, "production_source", "legacy")
    legacy(db,case,20)
    plan(db,case,0,DAY,100,final=True)
    import_rows(db,case,[("source",DAY,case["shared"],30)])
    key=(case["orders"][0].id,case["ops"][0][0].id)
    assert settings.production_source=="legacy"
    assert produced_qty_map(db,as_of=DAY)[0][key]==20
    assert produced_qty_map(db,as_of=DAY,include_mes=True)[0][key]==30
    monkeypatch.setattr(settings,"production_source","mes")
    assert produced_qty_map(db,as_of=DAY)[0][key]==30
    assert produced_qty_map(db,as_of=DAY-timedelta(days=1))[0]=={}
    state=client.get("/api/mes/source-status",headers=auth).json()
    assert state=={"production_source":"mes","has_mes_records":True,"automatic_cutover":False}


def test_matches_each_mes_unit_once_and_fulfillment_is_separate(db,case,monkeypatch,client,auth):
    monkeypatch.setattr(get_settings(),"production_source","mes")
    first=plan(db,case,0,DAY,30)
    second=plan(db,case,0,DAY+timedelta(weeks=1),30)
    import_rows(db,case,[("one",DAY,case["shared"],40),("two",DAY+timedelta(days=1),case["shared"],20)])
    measured=measure(db,DAY+timedelta(days=1))
    assert measured["matches"][first.id]["qty"]==30
    assert measured["matches"][second.id]["qty"]==30
    assert sum(m["qty"] for m in measured["matches"].values())==60
    kpis=weekly_kpis(db,[case["wc"].id],DAY,DAY+timedelta(weeks=1),DAY+timedelta(days=1))
    assert kpis[(case["wc"].id,DAY)]["standard_hour_equivalent_output"]==60
    assert kpis[(case["wc"].id,DAY+timedelta(weeks=1))]["standard_hour_equivalent_output"]==0
    row=next(r for r in order_progress(db,None,DAY+timedelta(days=1)) if r.order_id==case["orders"][0].id)
    assert row.earned_hours==60 and row.pct==0 and row.reserved_qty==0
    assert row.first_prod_date==DAY and row.last_prod_date==DAY+timedelta(days=1)
    from app.services.gantt import plan_gantt
    from app.services.progress import week_progress
    gantt=plan_gantt(db,case["wc"].id,DAY,DAY+timedelta(days=13),DAY+timedelta(days=1))
    assert sum(bar.produced_qty for bar in gantt.bars)==60
    assert week_progress(db,case["wc"],DAY,DAY+timedelta(days=1)).actual_hours_to_date==60

    export=client.get("/api/progress/orders.xlsx",params={"as_of":DAY},headers=auth)
    assert export.status_code==200,export.text
    wb=load_workbook(BytesIO(export.content))
    assert "Ölçüm açıklaması" in wb.sheetnames
    assert wb["Sipariş İlerleme"].max_column==19
    assert wb["Operasyon Detayı"].max_column==14


def test_batch_reservation_move_keeps_intended_shares_and_total_need(db,case,monkeypatch):
    monkeypatch.setattr(get_settings(),"production_source","mes")
    first=case["orders"][0]
    second=Order(order_no="other-"+uuid4().hex,item_id=first.item_id,due_date=DAY,quantity=100)
    batch=ProductionBatch(batch_no="batch-"+uuid4().hex,item_id=first.item_id,due_date=DAY,quantity=200)
    db.add_all([second,batch]);db.flush()
    db.add_all([ProductionBatchOrder(batch_id=batch.id,order_id=o.id,quantity=100) for o in [first,second]])
    reservation=Reservation(item_id=first.item_id,order_id=first.id,quantity=60,stock_provenance="external_finished_stock")
    db.add(reservation);db.commit()
    slot=plan(db,case,0,DAY,200,final=True)
    slot.production_batch_id=batch.id;db.commit()
    op=case["ops"][0][1]
    for recipient in [first,second]:
        reservation.order_id=recipient.id;db.commit()
        assert required_qty_for_batch(db,batch,first,produced={},netting_cache={})[op.id]==140
        rows={r.order_id:r for r in order_progress(db,None,DAY)}
        assert rows[recipient.id].reserved_qty==60 and rows[recipient.id].pct==60
        assert rows[first.id].planned_share_qty==rows[second.id].planned_share_qty==100
        assert sum(r.quantity for r in batch.orders)==200


def test_mes_mode_preserves_old_receipts_without_sync(db,case,monkeypatch):
    monkeypatch.setattr(get_settings(),"production_source","mes")
    receipt=StockReceipt(item_id=case["items"][0].id,quantity=17,source="progress",receipt_date=DAY)
    db.add(receipt);db.commit()
    legacy(db,case,100)
    result=sync_progress_receipts(db)
    assert result["added"]==result["adjusted"]==0
    assert db.get(StockReceipt,receipt.id).quantity==17


@pytest.mark.parametrize("provenance",["external_finished_stock","completed_production","legacy_unspecified"])
def test_reserved_mes_finished_goods_not_deducted_twice(db,case,monkeypatch,provenance):
    monkeypatch.setattr(get_settings(),"production_source","mes")
    plan(db,case,0,DAY,100,final=True)
    import_rows(db,case,[("done",DAY,case["items"][0].code,60)])
    db.add(Reservation(item_id=case["items"][0].id,order_id=case["orders"][0].id,quantity=60,stock_provenance=provenance));db.commit()
    row=next(r for r in order_progress(db,None,DAY) if r.order_id==case["orders"][0].id)
    assert row.reserved_qty==60 and row.unfulfilled_qty==40
    assert all(op.remaining_qty==40 for op in row.ops)
    assert row.remaining_hours==60
    assert row.produced_qty==60


def test_batch_setup_not_repeated_for_each_customer(db,case,monkeypatch):
    monkeypatch.setattr(get_settings(),"production_source","mes")
    first=case["orders"][0]
    other=Order(order_no="setup-"+uuid4().hex,item_id=first.item_id,due_date=DAY,quantity=100)
    batch=ProductionBatch(batch_no="setup-"+uuid4().hex,item_id=first.item_id,due_date=DAY,quantity=200)
    db.add_all([other,batch]);db.flush()
    db.add_all([ProductionBatchOrder(batch_id=batch.id,order_id=o.id,quantity=100) for o in [first,other]])
    for op in case["ops"][0]: op.setup_time_min=60
    db.commit()
    rows=[r for r in order_progress(db,None,DAY) if r.order_id in [first.id,other.id]]
    assert sum(r.required_hours for r in rows)==302
    assert sum(r.remaining_hours for r in rows)==302


def test_historical_mes_date_still_uses_current_fulfillment(db,case,monkeypatch):
    from app.models import Shipment
    monkeypatch.setattr(get_settings(),"production_source","mes")
    plan(db,case,0,DAY,100,final=True)
    import_rows(db,case,[("historical",DAY,case["items"][0].code,60)])
    db.add(Shipment(item_id=case["items"][0].id,order_id=case["orders"][0].id,quantity=60,ship_date=DAY+timedelta(days=1)));db.commit()
    row=next(r for r in order_progress(db,None,DAY) if r.order_id==case["orders"][0].id)
    assert row.shipped_qty==60 and row.unfulfilled_qty==40
    assert all(op.remaining_qty==40 for op in row.ops)
