"""Otomatik planlama on kontrol testleri."""

from datetime import date

from tests.test_capacity_flow import _upload
from tests.test_data_freshness import _clear_import_logs

WEEK = date(2026, 9, 8)


def _clean_orders(client, auth):
    for st in ("open", "closed", "merged"):
        client.delete("/api/orders", headers=auth, params={"status": st})


def _seed_daily_imports(client, auth, wc_code: str, item_code: str, order_no: str, wip_code: str = ""):
    """Bugunku gunluk veri importlarini tamamlar (orders haric — cagiran yukler)."""
    today = date.today().isoformat()
    prod_row = [today, wc_code, item_code, 10, order_no, 1]
    prod_cols = ["Tarih", "İş Merkezi Kodu", "Stok Kodu", "Operasyon Sıra", "Sipariş No", "Miktar"]
    if wip_code:
        prod_cols.append("Yarımamül Kodu")
        prod_row.append(wip_code)
    _upload(client, auth, "production", prod_cols, [prod_row])
    _upload(client, auth, "stock_receipts", ["Tarih", "Stok Kodu", "Miktar"], [[today, item_code, 1]])


def test_preflight_blocks_no_routing(client, auth, db):
    _clean_orders(client, auth)
    _upload(client, auth, "workcenters", ["İş Merkezi Kodu", "İş Merkezi Adı", "Planlanıyor (E/H)", "Birim Saat", "Kişi Başı Verimli Saat"], [["PF-1", "Preflight WC", "E", 10, 4]])
    wc = next(w for w in client.get("/api/workcenters", headers=auth).json() if w["code"] == "PF-1")
    _upload(client, auth, "items", ["Stok Kodu", "Stok Adı", "Ürün Grubu"], [["PF-NOROTA", "Rotasiz", "G"]])
    # Siparis giris kapisi rotasiz urunu artik reddeder; eski (kapidan once girilmis) siparis dogrudan yazilir.
    from tests.test_capacity_flow import _xlsx

    res = client.post("/api/imports/orders", headers=auth, files={"file": ("orders.xlsx", _xlsx(["Sipariş No", "Termin", "Stok Kodu", "Miktar"], [["PF-O1", "2026-10-01", "PF-NOROTA", 10]]), "application/octet-stream")}).json()
    assert res["errors"] and "Rota tanimi eksik" in res["errors"][0]
    from datetime import date as _date

    from app.models import Item, Order

    item = db.query(Item).filter(Item.code == "PF-NOROTA").one()
    db.add(Order(order_no="PF-O1", item_id=item.id, quantity=10, due_date=_date(2026, 10, 1), status="open"))
    db.commit()

    body = {"start_week": WEEK.isoformat(), "weeks": 4, "work_center_ids": [wc["id"]], "mode": "due_date"}
    pf = client.post("/api/plan/auto/preflight", headers=auth, json=body).json()
    assert pf["can_plan"] is False
    assert any(r["item_code"] == "PF-NOROTA" for r in pf["no_routing"])

    from io import BytesIO
    from openpyxl import load_workbook
    exported = client.post("/api/plan/auto/preflight/no-routing.xlsx", headers=auth, json=body)
    assert exported.status_code == 200, exported.text
    assert "attachment" in exported.headers["content-disposition"]
    wb = load_workbook(BytesIO(exported.content))
    ws = wb["Eksik Rotalar"]
    assert list(ws.values)[1:] == [(r["item_code"], r["item_name"], r["order_count"], ", ".join(r["order_nos"])) for r in pf["no_routing"]]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == ws.dimensions
    assert "Kontrol Kapsamı" in wb.sheetnames

    r = _plan_with_ack(client, "/api/plan/auto", headers=auth, json=body)
    assert r.status_code == 400


def test_preflight_warns_no_capacity(client, auth):
    _clean_orders(client, auth)
    _upload(client, auth, "workcenters", ["İş Merkezi Kodu", "İş Merkezi Adı", "Planlanıyor (E/H)", "Birim Saat", "Kişi Başı Verimli Saat"], [["PF-2", "Kapasitesiz", "E", 10, 4]])
    wc = next(w for w in client.get("/api/workcenters", headers=auth).json() if w["code"] == "PF-2")
    _upload(client, auth, "items", ["Stok Kodu", "Stok Adı", "Ürün Grubu"], [["PF-M", "Mamul", "G"]])
    _upload(client, auth, "routing", ["Stok Kodu", "Sıra", "Operasyon", "İş Merkezi Kodu", "Çevrim Süresi (sn)"], [["PF-M", 10, "Op", "PF-2", 3600]])
    _upload(client, auth, "orders", ["Sipariş No", "Termin", "Stok Kodu", "Miktar"], [["PF-O2", "2026-10-01", "PF-M", 1]])
    _seed_daily_imports(client, auth, "PF-2", "PF-M", "PF-O2")

    body = {"start_week": WEEK.isoformat(), "weeks": 2, "work_center_ids": [wc["id"]], "mode": "due_date"}
    pf = client.post("/api/plan/auto/preflight", headers=auth, json=body).json()
    assert pf["can_plan"] is True
    assert pf["needs_capacity_ack"] is True
    assert pf["needs_daily_data_ack"] is False
    assert any(r["work_center_code"] == "PF-2" for r in pf["no_capacity"])

    r = _plan_with_ack(client, "/api/plan/auto", headers=auth, json=body)
    assert r.status_code == 200


def test_preflight_warns_stale_daily_data(client, auth, db):
    _clear_import_logs(db)
    _clean_orders(client, auth)
    _upload(client, auth, "workcenters", ["İş Merkezi Kodu", "İş Merkezi Adı", "Planlanıyor (E/H)", "Birim Saat", "Kişi Başı Verimli Saat"], [["PF-4", "WC", "E", 10, 4]])
    wc = next(w for w in client.get("/api/workcenters", headers=auth).json() if w["code"] == "PF-4")
    _upload(client, auth, "employees", ["Sicil No", "Ad Soyad", "İş Merkezi Kodu"], [["PF4-1", "Ali", "PF-4"]])
    _upload(client, auth, "shifts", ["İş Merkezi Kodu", "Vardiya", "Günler (Pzt=0..Paz=6)", "Başlangıç", "Bitiş", "Kişi Sayısı", "Kişi Başı Verimli Saat"], [["PF-4", "G", "0,1,2,3,4", "08:00", "18:00", 1, 4]])
    _upload(client, auth, "items", ["Stok Kodu", "Stok Adı", "Ürün Grubu"], [["PF-D", "M", "G"]])
    _upload(client, auth, "routing", ["Stok Kodu", "Sıra", "Operasyon", "İş Merkezi Kodu", "Çevrim Süresi (sn)"], [["PF-D", 10, "Op", "PF-4", 360]])
    _upload(client, auth, "orders", ["Sipariş No", "Termin", "Stok Kodu", "Miktar"], [["PF-O4", "2026-10-01", "PF-D", 1]])

    body = {"start_week": WEEK.isoformat(), "weeks": 2, "work_center_ids": [wc["id"]], "mode": "due_date"}
    pf = client.post("/api/plan/auto/preflight", headers=auth, json=body).json()
    assert pf["needs_daily_data_ack"] is True
    assert len(pf["daily_data"]) == 4
    assert any(c["key"] == "open_orders" and c["status"] == "ok" for c in pf["daily_data"])
    assert any(c["status"] != "ok" for c in pf["daily_data"])


def test_preflight_ok_with_routing_and_capacity(client, auth):
    _clean_orders(client, auth)
    _upload(client, auth, "workcenters", ["İş Merkezi Kodu", "İş Merkezi Adı", "Planlanıyor (E/H)", "Birim Saat", "Kişi Başı Verimli Saat"], [["PF-3", "Tamam", "E", 10, 4]])
    wc = next(w for w in client.get("/api/workcenters", headers=auth).json() if w["code"] == "PF-3")
    _upload(client, auth, "employees", ["Sicil No", "Ad Soyad", "İş Merkezi Kodu"], [["PF3-1", "Ali", "PF-3"]])
    _upload(
        client,
        auth,
        "shifts",
        ["İş Merkezi Kodu", "Vardiya", "Günler (Pzt=0..Paz=6)", "Başlangıç", "Bitiş", "Kişi Sayısı", "Kişi Başı Verimli Saat"],
        [["PF-3", "Gündüz", "0,1,2,3,4", "08:00", "18:00", 1, 4]],
    )
    _upload(client, auth, "items", ["Stok Kodu", "Stok Adı", "Ürün Grubu"], [["PF-OK", "Mamul", "G"]])
    _upload(
        client,
        auth,
        "routing",
        ["Stok Kodu", "Sıra", "Operasyon", "İş Merkezi Kodu", "Çevrim Süresi (sn)", "Yarımamül Kodu"],
        [["PF-OK", 10, "Op", "PF-3", 360, "PF-OK-WIP"]],
    )
    _upload(client, auth, "orders", ["Sipariş No", "Termin", "Stok Kodu", "Miktar"], [["PF-O3", "2026-10-01", "PF-OK", 5]])
    _seed_daily_imports(client, auth, "PF-3", "PF-OK", "PF-O3", "PF-OK-WIP")

    body = {"start_week": WEEK.isoformat(), "weeks": 2, "work_center_ids": [wc["id"]], "mode": "due_date"}
    pf = client.post("/api/plan/auto/preflight", headers=auth, json=body).json()
    assert pf["can_plan"] is True
    # Positive shift defaults do not constitute an explicit weekly labor entry.
    assert len(pf["no_capacity"]) == 2
    assert all(r["capacity_hours"] == 0 for r in pf["no_capacity"])
    assert pf["needs_capacity_ack"] is True
    assert len(pf["missing_labor_weeks"]) == 2
    assert pf["needs_daily_data_ack"] is False
    assert all(c["status"] == "ok" for c in pf["daily_data"])
    assert pf["no_routing"] == []

from tests.test_capacity_flow import _plan_with_ack
