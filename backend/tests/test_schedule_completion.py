from datetime import date
from app.models import Order, PlanLine
from app.services.orders import order_schedule, plan_line_window
from app.services.gantt import _line_window_in_week
from app.services.excel import completion_schedule_sheet
from tests.test_orders_flow import _setup


def test_full_selected_scope_without_final_date_is_not_on_time(client, auth, db):
    from io import BytesIO
    from openpyxl import load_workbook
    _setup(client, auth)
    oid = client.post("/api/orders", headers=auth, json={"order_no": "NO-END", "item_code": "MAM-1", "quantity": 100, "due_date": "2026-09-30"}).json()["id"]
    order = db.get(Order, oid)
    first, last = sorted(order.item.operations, key=lambda op: op.seq)
    db.add(PlanLine(order_id=oid, operation_id=first.id, work_center_id=first.work_center_id,
                    week_start=date(2026, 9, 7), planned_qty=100, planned_hours=first.hours_for(100), mode="auto"))
    db.commit()
    selected = order_schedule(db, [first.work_center_id])[0]
    assert selected.coverage_pct == 100
    assert selected.planned_end is None and selected.lateness_days is None
    assert selected.plan_status == "finish_unknown"
    # In the full route scope the missing final operation remains a partial plan.
    assert order_schedule(db, None)[0].plan_status == "partial"
    response = client.get("/api/plan/orders.xlsx", headers=auth, params={
        "work_center_ids": first.work_center_id, "plan_status": "finish_unknown"})
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content))["Sipariş Bitiş Tarihleri"]
    assert sheet.max_row == 2
    assert list(sheet.values)[1][-2] == "Bitiş tarihi belirsiz"
    # Once the final operation is planned, a real date can establish on-time status.
    db.add(PlanLine(order_id=oid, operation_id=last.id, work_center_id=last.work_center_id,
                    week_start=date(2026, 9, 14), planned_qty=100, planned_hours=last.hours_for(100), mode="auto"))
    db.commit()
    complete = order_schedule(db, None)[0]
    assert complete.planned_end is not None and complete.plan_status == "on_time"


def test_final_operation_week_quantities_and_shared_gantt_date(client, auth, db):
    _setup(client, auth)
    oid = client.post("/api/orders", headers=auth, json={"order_no":"COMPLETE", "item_code":"MAM-1", "quantity":100, "due_date":"2026-09-30"}).json()["id"]
    order = db.get(Order,oid)
    first,last = sorted(order.item.operations, key=lambda o:o.seq)
    rows = []
    for op,wk,qty,mode in [(first,date(2026,9,21),100,"auto"),(last,date(2026,9,7),40,"auto"),(last,date(2026,9,14),60,"auto"),(last,date(2026,9,28),500,"forecast")]:
        row=PlanLine(order_id=oid,operation_id=op.id,work_center_id=op.work_center_id,
                     week_start=wk,planned_qty=qty,planned_hours=op.hours_for(qty),mode=mode)
        db.add(row); rows.append(row)
    db.commit()
    result=next(r for r in order_schedule(db,None) if r.order_id==oid)
    assert [w.quantity for w in result.completion_weeks]==[40,60]
    assert result.planned_end_week==date(2026,9,14)
    assert result.planned_end==_line_window_in_week(db,last.work_center,rows[2].week_start,rows[2].id,[rows[2]])[1]
    header,values=completion_schedule_sheet([result])
    assert [r[6] for r in values]==[40,60]
    assert len(header)==len(values[0])
    # Gantt includes other orders occupying the same week, even a closed order.
    closed = Order(order_no="CLOSED-CONTEXT", item_id=order.item_id, quantity=1,
                   due_date=date(2026,9,1), status="closed")
    db.add(closed); db.flush()
    other = PlanLine(order_id=closed.id,operation_id=last.id,work_center_id=last.work_center_id,
                     week_start=rows[2].week_start,planned_qty=1,planned_hours=90,mode="manual")
    db.add(other); db.commit()
    result=next(r for r in order_schedule(db,None) if r.order_id==oid)
    assert result.planned_end==_line_window_in_week(db,last.work_center,rows[2].week_start,rows[2].id,[other,rows[2]])[1]
    assert result.planned_end==date(2026,9,18)



def test_unrelated_later_order_does_not_shift_line_end(client, auth, db):
    from types import SimpleNamespace
    _setup(client,auth)
    from app.models import WorkCenter
    wc=db.query(WorkCenter).filter_by(code="MON-1").one()
    early=SimpleNamespace(id=1,order_id=1,order=SimpleNamespace(order_no="A",due_date=date(2026,9,10),revised_due_date=None),operation_seq=10,planned_hours=1)
    later=SimpleNamespace(id=2,order_id=2,order=SimpleNamespace(order_no="B",due_date=date(2026,9,20),revised_due_date=None),operation_seq=10,planned_hours=99)
    window=plan_line_window(db,wc,date(2026,9,7),early,[early,later])
    assert window==(date(2026,9,7),date(2026,9,7))
