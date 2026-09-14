from datetime import date, timedelta
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.models import Item, Order, ProductionActual, Reservation, RoutingOperation, Shipment, StockReceipt, WorkCenter


@pytest.fixture
def stock_case(db):
    item = Item(code=f"PRE-{uuid4().hex[:8]}", name="Önizleme ürünü")
    other = Item(code=f"PRE-{uuid4().hex[:8]}", name="Diğer ürün")
    db.add_all([item, other])
    db.flush()
    early = Order(order_no="PRE-A", position_no="001", customer="=Müşteri", item_id=item.id,
                  quantity=50, due_date=date.today())
    late = Order(order_no="PRE-B", position_no="002", customer="Geç müşteri", item_id=item.id,
                 quantity=40, due_date=date.today() + timedelta(days=10))
    unrelated = Order(order_no="PRE-C", position_no="003", customer="Diğer müşteri", item_id=other.id,
                      quantity=20, due_date=date.today())
    db.add_all([early, late, unrelated])
    db.flush()
    db.add_all([StockReceipt(item_id=item.id, quantity=75, receipt_date=date.today()),
                StockReceipt(item_id=other.id, quantity=10, receipt_date=date.today()),
                Reservation(item_id=item.id, order_id=late.id, quantity=10, source="manual"),
                Shipment(item_id=item.id, order_id=early.id, quantity=5, ship_date=date.today())])
    db.commit()
    return item.id, early.id, late.id, other.id


def preview(client, auth, item_id=None):
    r = client.post("/api/stock/reservations/auto/preview", headers=auth,
                    json={"item_ids": [item_id]} if item_id else {})
    assert r.status_code == 200, r.text
    return r.json()


def confirm(client, auth, p):
    return client.post("/api/stock/reservations/auto", headers=auth, json={"preview_token": p["preview_token"]})


def test_preview_cancel_and_exact_confirmation(client, auth, db, stock_case):
    iid, early, late, _ = stock_case
    receipts = db.query(StockReceipt).count()
    p = preview(client, auth, iid)
    assert db.query(Reservation).count() == 1
    assert db.query(StockReceipt).count() == receipts
    assert [(r["order_id"], r["allocate"]) for r in p["rows"]] == [(early, 45), (late, 15)]
    assert p["rows"][1]["remaining_after"] == 15
    assert p["rows"][0]["customer"] == "=Müşteri"
    assert p["reserved_qty"] == 60 and p["items"] == 1
    # Opening again after cancellation changes no stock and proposes identical rows.
    assert preview(client, auth, iid)["rows"] == p["rows"]
    r = confirm(client, auth, p)
    assert r.status_code == 200, r.text
    assert r.json()["created"] == 2 and r.json()["reserved_qty"] == 60
    db.expire_all()
    assert {(r.order_id, r.quantity) for r in db.query(Reservation).filter_by(source="auto")} == {(early, 45), (late, 15)}
    assert db.query(Reservation).filter_by(source="manual").one().quantity == 10
    assert confirm(client, auth, p).status_code == 409
    assert db.query(Reservation).count() == 3


@pytest.mark.parametrize("change", ["receipt", "quantity", "customer", "due_date", "reservation"])
def test_stale_preview_rejected_without_partial_writes(client, auth, db, stock_case, change):
    iid, early, _, _ = stock_case
    p = preview(client, auth, iid)
    if change == "receipt":
        db.query(StockReceipt).filter_by(item_id=iid).one().quantity += 1
    elif change == "reservation":
        db.query(Reservation).one().quantity += 1
    else:
        o = db.get(Order, early)
        setattr(o, change, {"quantity": 51, "customer": "Yeni müşteri", "due_date": date.today() + timedelta(days=20)}[change])
    db.commit()
    assert confirm(client, auth, p).status_code == 409
    assert db.query(Reservation).filter_by(source="auto").count() == 0


def test_all_scope_and_confirmation_requires_valid_preview(client, auth, stock_case):
    iid, _, _, other = stock_case
    p = preview(client, auth)
    assert {r["item_id"] for r in p["rows"]} == {iid, other}
    assert p["reserved_qty"] == 70
    assert client.post("/api/stock/reservations/auto", headers=auth, json={}).status_code == 422
    assert confirm(client, auth, {"preview_token": "invalid"}).status_code == 409
    assert client.post("/api/stock/reservations/auto/preview", json={}).status_code == 401
    assert client.get("/api/stock/summary", headers={"Authorization": f"Bearer {p['preview_token']}"}).status_code == 401


def test_export_values_filters_and_safe_text(client, auth, db, stock_case):
    _, _, late, _ = stock_case
    # Fully reserved but still open order remains in the report when the remaining filter is off.
    db.query(Reservation).filter_by(order_id=late).one().quantity = 40
    db.commit()
    r = client.get("/api/stock/orders/export.xlsx", headers=auth, params={"position": "001"})
    assert r.status_code == 200, r.text
    assert "attachment" in r.headers["content-disposition"]
    ws = load_workbook(BytesIO(r.content)).worksheets[0]
    assert ws.max_row == 2 and ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:N2"
    values = list(ws.values)[1]
    assert values[1] == "001" and values[2] == "=Müşteri" and ws["C2"].data_type == "s"
    assert values[6:11] == (50, 5, 45, 0, 45)
    assert values[3].date() == date.today()
    for only_remaining, count in [(False, 4), (True, 3)]:
        r = client.get("/api/stock/orders/export.xlsx", headers=auth, params={"only_remaining": only_remaining})
        assert load_workbook(BytesIO(r.content)).worksheets[0].max_row == count
    r = client.get("/api/stock/orders/export.xlsx", headers=auth, params={"position": "no-match"})
    assert load_workbook(BytesIO(r.content)).worksheets[0].max_row == 1


def test_production_receipts_are_persisted_only_after_confirmation(client, auth, db, stock_case):
    iid, _, _, _ = stock_case
    wc = WorkCenter(code=f"PRE-{uuid4().hex[:8]}", name="Önizleme merkezi")
    db.add(wc)
    db.flush()
    db.add(RoutingOperation(item_id=iid, work_center_id=wc.id, seq=10))
    db.add(ProductionActual(item_id=iid, work_center_id=wc.id, operation_seq=10,
                            quantity=5.125, prod_date=date.today()))
    db.commit()
    p = preview(client, auth, iid)
    assert p["reserved_qty"] == 65.125
    assert db.query(StockReceipt).filter_by(item_id=iid, source="progress").count() == 0
    assert confirm(client, auth, p).status_code == 200
    assert db.query(StockReceipt).filter_by(item_id=iid, source="progress").one().quantity == 5.125


def test_simultaneous_confirmations_do_not_duplicate(client, auth, db, stock_case):
    from concurrent.futures import ThreadPoolExecutor

    iid, _, _, _ = stock_case
    p = preview(client, auth, iid)
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: confirm(client, auth, p), range(2)))
    assert sorted(r.status_code for r in results) == [200, 409]
    assert sum(r.quantity for r in db.query(Reservation).filter_by(source="auto")) == 60
