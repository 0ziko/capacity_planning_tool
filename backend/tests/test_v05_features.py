"""v0.5: haftalik is gucu istisnalari, senaryo matrisi (operasyon gecis kurallari), stok & rezervasyon."""

from datetime import date, datetime, timedelta

from tests.test_capacity_flow import _upload


def _monday(offset_weeks: int = 0) -> str:
    d = date.today()
    return (d - timedelta(days=d.weekday()) + timedelta(weeks=offset_weeks)).isoformat()


def _wc(client, auth, code):
    return next(w for w in client.get("/api/workcenters", headers=auth).json() if w["code"] == code)


def _res_by_order(client, auth, item_id) -> dict:
    """order_no -> {quantity (toplam), source (kume), ids}"""
    out: dict = {}
    for x in client.get("/api/stock/reservations", headers=auth, params={"item_id": item_id}).json():
        r = out.setdefault(x["order_no"], {"quantity": 0.0, "sources": set(), "ids": []})
        r["quantity"] += x["quantity"]
        r["sources"].add(x["source"])
        r["ids"].append(x["id"])
    return out


def _week_cap(client, auth, wc_id, week: str) -> float:
    rows = client.get("/api/plan/load", headers=auth, params={"start": week, "weeks": 1, "work_center_ids": [wc_id]}).json()
    return next(r for r in rows if r["work_center_id"] == wc_id)["weeks"][0]["capacity_hours"]


def test_weekly_labor_override(client, auth):
    _upload(client, auth, "workcenters", ["İş Merkezi Kodu", "İş Merkezi Adı", "Planlanıyor (E/H)", "Kişi Başı Verimli Saat"], [["HFT-1", "Haftalık İM", "E", 4]])
    wc = _wc(client, auth, "HFT-1")
    _upload(client, auth, "employees", ["Sicil No", "Ad Soyad", "İş Merkezi Kodu"], [["H-1", "A", "HFT-1"], ["H-2", "B", "HFT-1"]])
    wk0, wk1 = _monday(), _monday(1)
    assert _week_cap(client, auth, wc["id"], wk0) == 2 * 4 * 5

    # Hafta profili: varsayilanlar
    rows = client.get(f"/api/workcenters/{wc['id']}/weeks", headers=auth, params={"start": wk0, "weeks": 2}).json()
    assert rows[0]["headcount"] == 2 and rows[0]["working_days"] == 5 and rows[0]["has_override"] is False

    # Ikinci haftaya istisna: 3 kisi, 4.5 saat, 4 gun
    r = client.put(f"/api/workcenters/{wc['id']}/weeks/{wk1}", headers=auth, json={"headcount": 3, "efficient_hours_per_person": 4.5, "working_days": 4, "note": "bayram"})
    assert r.status_code == 200, r.text
    assert r.json()["has_override"] and r.json()["capacity_hours"] == 3 * 4.5 * 4
    assert _week_cap(client, auth, wc["id"], wk1) == 54
    assert _week_cap(client, auth, wc["id"], wk0) == 40  # diger hafta degismedi

    # Excel ile hafta numarasiyla yukleme (yalnizca kisi sayisi); '2026-W37' formati
    iso = date.fromisoformat(wk0).isocalendar()
    r = _upload(client, auth, "wc_weeks", ["İş Merkezi Kodu", "Hafta", "Kişi Sayısı"], [["HFT-1", f"{iso[0]}-W{iso[1]}", 1]])
    assert not r["errors"], r["errors"]
    assert _week_cap(client, auth, wc["id"], wk0) == 1 * 4 * 5

    # Bos degerler => istisna silinir
    r = client.put(f"/api/workcenters/{wc['id']}/weeks/{wk0}", headers=auth, json={})
    assert r.json()["has_override"] is False
    assert _week_cap(client, auth, wc["id"], wk0) == 40


def test_scenario_matrix_rules_and_leadtime(client, auth):
    _upload(client, auth, "workcenters", ["İş Merkezi Kodu", "İş Merkezi Adı", "Planlanıyor (E/H)", "Kişi Başı Verimli Saat"], [["PRS-T", "Preshane Test", "E", 8]])
    _upload(client, auth, "employees", ["Sicil No", "Ad Soyad", "İş Merkezi Kodu"], [["P-1", "A", "PRS-T"]])
    _upload(client, auth, "items", ["Stok Kodu", "Stok Adı", "Ürün Grubu"], [["EVY-1", "Evye 1", "EVYE"], ["EVY-2", "Evye 2", "EVYE"]])
    # 3 operasyon; her biri 100 adet x 36 sn = 1 saat
    _upload(
        client, auth, "routing",
        ["Stok Kodu", "Sıra", "Operasyon", "İş Merkezi Kodu", "Çevrim Süresi (sn)"],
        [["EVY-1", 10, "Sıvama", "PRS-T", 36], ["EVY-1", 20, "Forma", "PRS-T", 36], ["EVY-1", 30, "Etek Kesme", "PRS-T", 36],
         ["EVY-2", 10, "Sıvama", "PRS-T", 36], ["EVY-2", 20, "Forma", "PRS-T", 36]],
    )
    groups = client.get("/api/scenarios/groups", headers=auth).json()
    g = next(x for x in groups if x["product_group"] == "EVYE")
    assert g["operations"] == ["Sıvama", "Forma", "Etek Kesme"] and g["item_count"] == 2

    flow = client.get("/api/scenarios/flow", headers=auth, params={"product_group": "EVYE"}).json()
    assert [n["name"] for n in flow["nodes"]] == ["Sıvama", "Forma", "Etek Kesme"]
    assert all(t["effective"]["source"] == "default" for t in flow["transitions"])
    assert next(i for i in flow["items"] if i["code"] == "EVY-2")["differs"] is True

    start = _monday(4)  # bos bir hafta
    def lt(code):
        r = client.post("/api/plan/leadtime", headers=auth, json={"item_code": code, "quantity": 100, "start": start})
        assert r.status_code == 200, r.text
        return r.json()

    base = lt("EVY-1")
    fmt = "%Y-%m-%d %H:%M"
    s0, e0 = datetime.strptime(base["steps"][0]["start"], fmt), datetime.strptime(base["steps"][0]["end"], fmt)
    s1 = datetime.strptime(base["steps"][1]["start"], fmt)
    assert s1 >= e0  # varsayilan: onceki bitince

    # Grup kurali: Sıvama basladiktan 5 cevrim sonra Forma baslar
    r = client.put("/api/scenarios/rules", headers=auth, json={"scope": "group", "product_group": "EVYE", "from_op": "sıvama", "to_op": "FORMA", "rule": "cycles", "lag_cycles": 5})
    assert r.status_code == 200, r.text
    assert "5 çevrim" in r.json()["description"]
    over = lt("EVY-1")
    s1b = datetime.strptime(over["steps"][1]["start"], fmt)
    e1b = datetime.strptime(over["steps"][1]["end"], fmt)
    assert s0 <= s1b < e0  # ic ice basladi
    assert e1b >= e0  # ama oncekinden once bitmez
    assert "5 çevrim" in over["steps"][1]["start_rule"] and "grup" in over["steps"][1]["start_rule"]
    assert datetime.strptime(over["end"], fmt) < datetime.strptime(base["end"], fmt)  # toplam termin kisaldi

    # Stok ozel kural grup kuralini ezer: EVY-2 icin Forma, Sıvama bitince + 60 dk bekleme
    r = client.put("/api/scenarios/rules", headers=auth, json={"scope": "item", "item_code": "EVY-2", "from_op": "Sıvama", "to_op": "Forma", "rule": "finish", "wait_minutes": 60})
    assert r.status_code == 200, r.text
    f2 = client.get("/api/scenarios/flow", headers=auth, params={"product_group": "EVYE", "item_code": "EVY-2"}).json()
    t = f2["transitions"][0]
    assert t["item_rule"]["rule"] == "finish" and t["group_rule"]["rule"] == "cycles" and t["effective"]["source"] == "item"
    l2 = lt("EVY-2")
    e0b = datetime.strptime(l2["steps"][0]["end"], fmt)
    s1c = datetime.strptime(l2["steps"][1]["start"], fmt)
    assert s1c >= e0b + timedelta(minutes=60)

    # Excel ile kural yukleme + kural listesi
    r = _upload(client, auth, "op_rules", ["Ürün Grubu", "Stok Kodu", "Önceki Operasyon", "Sonraki Operasyon", "Kural", "Çevrim Sayısı"], [["EVYE", "", "Forma", "Etek Kesme", "Çevrim", 5]])
    assert not r["errors"], r["errors"]
    rules = client.get("/api/scenarios/rules", headers=auth, params={"product_group": "EVYE"}).json()
    assert len(rules) == 3
    # sil
    assert client.delete(f"/api/scenarios/rules/{rules[0]['id']}", headers=auth).status_code == 204

    # Otomatik plan: cycles kuralinda ic ice hafta, kural yokken sonraki haftaya kayabilir — hata vermeden calisir
    _upload(client, auth, "orders", ["Sipariş No", "Müşteri", "Termin", "Stok Kodu", "Miktar"], [["SC-1", "M", (date.today() + timedelta(days=60)).isoformat(), "EVY-1", 100]])
    r = client.post("/api/plan/auto", headers=auth, json={"start_week": start, "weeks": 4, "work_center_ids": [_wc(client, auth, "PRS-T")["id"]], "replace_existing": True})
    assert r.status_code == 200, r.text


def test_stock_reservations_and_shipping(client, auth):
    _upload(client, auth, "items", ["Stok Kodu", "Stok Adı", "Ürün Grubu"], [["RZ-1", "Rezerve Ürün", "EVYE"]])
    d = date.today()
    _upload(
        client, auth, "orders",
        ["Sipariş No", "Müşteri", "Termin", "Stok Kodu", "Miktar"],
        [["RZ-A", "Erken", (d + timedelta(days=10)).isoformat(), "RZ-1", 50],
         ["RZ-B", "Orta", (d + timedelta(days=20)).isoformat(), "RZ-1", 30],
         ["RZ-C", "Geç", (d + timedelta(days=30)).isoformat(), "RZ-1", 40]],
    )
    # Depo girisi: 60 adet (manuel) + 20 adet (Excel)
    r = client.post("/api/stock/receipts", headers=auth, json={"item_code": "RZ-1", "receipt_date": d.isoformat(), "quantity": 60, "lot": "L1"})
    assert r.status_code == 201, r.text
    r = _upload(client, auth, "stock_receipts", ["Tarih", "Stok Kodu", "Miktar"], [[d.isoformat(), "RZ-1", 20]])
    assert not r["errors"], r["errors"]

    row = next(x for x in client.get("/api/stock/summary", headers=auth).json() if x["item_code"] == "RZ-1")
    assert row["on_hand"] == 80 and row["free"] == 80 and row["open_demand"] == 120 and row["open_orders"] == 3

    orders = {o["order_no"]: o for o in client.get("/api/stock/orders", headers=auth).json()}
    # Manuel: kritik musteri "Geç" (RZ-C) icin 30 adet ayir
    r = client.post("/api/stock/reservations", headers=auth, json={"order_id": orders["RZ-C"]["order_id"], "quantity": 30, "note": "kritik"})
    assert r.status_code == 201, r.text
    assert r.json()["source"] == "manual"
    # Siparisin kalanindan fazlasi rezerve edilemez
    assert client.post("/api/stock/reservations", headers=auth, json={"order_id": orders["RZ-C"]["order_id"], "quantity": 20}).status_code == 400

    # Otomatik: kalan 50 serbest -> termin sirasiyla RZ-A 50 (RZ-B'ye kalmaz)
    r = client.post("/api/stock/reservations/auto", headers=auth, json={})
    assert r.status_code == 200 and r.json()["created"] == 1 and r.json()["reserved_qty"] == 50
    by_order = _res_by_order(client, auth, row["item_id"])
    assert by_order["RZ-A"]["quantity"] == 50 and by_order["RZ-A"]["sources"] == {"auto"}
    assert by_order["RZ-C"]["quantity"] == 30 and by_order["RZ-C"]["sources"] == {"manual"}
    row = next(x for x in client.get("/api/stock/summary", headers=auth).json() if x["item_code"] == "RZ-1")
    assert row["free"] == 0 and row["reserved"] == 80

    # Manuel oncelik: serbest stok yokken RZ-B'ye 20 manuel rezerve => otomatik (RZ-A) 20 azalir
    r = client.post("/api/stock/reservations", headers=auth, json={"order_id": orders["RZ-B"]["order_id"], "quantity": 20})
    assert r.status_code == 201, r.text
    by_order = _res_by_order(client, auth, row["item_id"])
    assert by_order["RZ-A"]["quantity"] == 30 and by_order["RZ-B"]["quantity"] == 20
    # RZ-B'nin kalani 10: 10 daha manuel => otomatik RZ-A 20'ye duser
    assert client.post("/api/stock/reservations", headers=auth, json={"order_id": orders["RZ-B"]["order_id"], "quantity": 10}).status_code == 201
    # Siparis kalanindan fazlasi olamaz
    assert client.post("/api/stock/reservations", headers=auth, json={"order_id": orders["RZ-B"]["order_id"], "quantity": 1}).status_code == 400
    # Manuel rezervasyonlar cozulmez: RZ-A icin 30 manuel istenir, yalnizca 20 otomatik cozulebilir => 400 (ve geri alinir)
    assert client.post("/api/stock/reservations", headers=auth, json={"order_id": orders["RZ-A"]["order_id"], "quantity": 30}).status_code == 400
    by_order = _res_by_order(client, auth, row["item_id"])
    assert by_order["RZ-A"]["quantity"] == 20 and by_order["RZ-A"]["sources"] == {"auto"} and by_order["RZ-B"]["quantity"] == 30

    # Rezervasyonu baska siparise tasi: RZ-B'nin kalani 0 => 400
    a_id = by_order["RZ-A"]["ids"][0]
    assert client.patch(f"/api/stock/reservations/{a_id}/move", headers=auth, params={"order_id": orders["RZ-B"]["order_id"]}).status_code == 400
    # Kaldir ve yeniden otomatik dagit => RZ-A 20 (serbest 20, en erken termin)
    assert client.delete(f"/api/stock/reservations/{a_id}", headers=auth).status_code == 204
    r = client.post("/api/stock/reservations/auto", headers=auth, json={"item_ids": [row["item_id"]]})
    assert r.json()["reserved_qty"] == 20

    # Sevk: RZ-C'nin 30'unu sevk et (kismi: 10 + kalan)
    c_id = _res_by_order(client, auth, row["item_id"])["RZ-C"]["ids"][0]
    r = client.post(f"/api/stock/reservations/{c_id}/ship", headers=auth, json={"quantity": 10})
    assert r.status_code == 201, r.text
    r = client.post(f"/api/stock/reservations/{c_id}/ship", headers=auth, json={})
    assert r.status_code == 201
    row = next(x for x in client.get("/api/stock/summary", headers=auth).json() if x["item_code"] == "RZ-1")
    assert row["on_hand"] == 50 and row["shipped"] == 30 and row["reserved"] == 50 and row["free"] == 0
    o = {x["order_no"]: x for x in client.get("/api/stock/orders", headers=auth, params={"include_closed": True}).json()}
    assert o["RZ-C"]["shipped"] == 30 and o["RZ-C"]["remaining"] == 10 and o["RZ-C"]["status"] == "open"
    # RZ-A'nin tamamini (50) sevk etmek icin 30 daha depo girisi + rezervasyon
    client.post("/api/stock/receipts", headers=auth, json={"item_code": "RZ-1", "receipt_date": d.isoformat(), "quantity": 30})
    assert client.post("/api/stock/reservations", headers=auth, json={"order_id": orders["RZ-A"]["order_id"], "quantity": 30}).status_code == 201
    for rr in client.get("/api/stock/reservations", headers=auth, params={"order_id": orders["RZ-A"]["order_id"]}).json():
        assert client.post(f"/api/stock/reservations/{rr['id']}/ship", headers=auth, json={}).status_code == 201
    o = {x["order_no"]: x for x in client.get("/api/stock/orders", headers=auth, params={"include_closed": True}).json()}
    assert o["RZ-A"]["status"] == "closed" and o["RZ-A"]["remaining"] == 0
    # sevki geri al => siparis acilir, stok manuel rezervasyon olarak geri gelir
    ships = client.get("/api/stock/shipments", headers=auth, params={"order_id": orders["RZ-A"]["order_id"]}).json()
    assert client.delete(f"/api/stock/shipments/{ships[0]['id']}", headers=auth).status_code == 204
    o = {x["order_no"]: x for x in client.get("/api/stock/orders", headers=auth).json()}
    assert o["RZ-A"]["status"] == "open"

    # Yedekte yeni sayfalar
    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(client.get("/api/backup.xlsx", headers=auth).content))
    for name in ("Haftalık İş Gücü", "Senaryo Kuralları", "Depo Girişi", "Rezervasyonlar", "Sevkler"):
        assert name in wb.sheetnames, wb.sheetnames


def test_production_creates_finished_stock(client, auth):
    """Son operasyon uretim beyani otomatik depo girisi olusturur."""
    wk = _monday()
    _upload(client, auth, "workcenters", ["İş Merkezi Kodu", "İş Merkezi Adı", "Planlanıyor (E/H)", "Kişi Başı Verimli Saat"], [["STK-1", "Stok IM", "E", 8]])
    _upload(client, auth, "employees", ["Sicil No", "Ad Soyad", "İş Merkezi Kodu"], [["S1", "A", "STK-1"]])
    _upload(client, auth, "items", ["Stok Kodu", "Ürün Grubu"], [["STK-M", "GN"]])
    _upload(
        client, auth, "routing",
        ["Stok Kodu", "Sıra", "Operasyon", "İş Merkezi Kodu", "Çevrim Süresi (sn)", "Yarımamül Kodu"],
        [["STK-M", 10, "OP1", "STK-1", 36, "STK-M-10"], ["STK-M", 20, "OP2", "STK-1", 36, "STK-M-20"]],
    )
    _upload(client, auth, "production", ["Tarih", "Yarımamül Kodu", "Miktar"], [[wk, "STK-M-10", 50]])
    rows = client.get("/api/stock/summary", headers=auth).json()
    assert not any(x["item_code"] == "STK-M" for x in rows)
    _upload(client, auth, "production", ["Tarih", "Yarımamül Kodu", "Miktar"], [[wk, "STK-M-20", 40]])
    row = next(x for x in client.get("/api/stock/summary", headers=auth).json() if x["item_code"] == "STK-M")
    assert row["on_hand"] == 40 and row["free"] == 40
    rcpts = client.get("/api/stock/receipts", headers=auth, params={"item_id": row["item_id"]}).json()
    assert len(rcpts) == 1 and rcpts[0]["source"] == "progress" and rcpts[0]["quantity"] == 40
    assert client.delete(f"/api/stock/receipts/{rcpts[0]['id']}", headers=auth).status_code == 400
