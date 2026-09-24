from datetime import date, timedelta
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.models import WorkCenter, WorkCenterWeek


@pytest.fixture
def weekly_case(db):
    wc = WorkCenter(code=f"XL-{uuid4().hex[:8]}", name="Haftalık iş gücü", default_efficient_hours=4)
    db.add(wc)
    db.flush()
    monday = date(2026, 9, 14)
    db.add(WorkCenterWeek(work_center_id=wc.id, week_start=monday, headcount=8,
                          efficient_hours_per_person=4.5, working_days=5, note="Mevcut"))
    db.commit()
    return wc.id, monday


def export(client, auth, iid, monday):
    r = client.get("/api/exports/wc-weeks.xlsx", headers=auth,
                   params={"start": (monday + timedelta(days=2)).isoformat(), "weeks": 2, "work_center_ids": iid})
    assert r.status_code == 200, r.text
    return load_workbook(BytesIO(r.content))


def upload(client, auth, wb):
    buf = BytesIO()
    wb.save(buf)
    r = client.post("/api/imports/wc_weeks", headers=auth,
                    files={"file": ("haftalik.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    return r.json()


def test_export_round_trip_and_update(client, auth, db, weekly_case):
    iid, monday = weekly_case
    wb = export(client, auth, iid, monday)
    ws = wb.worksheets[0]
    assert ws.max_row == 3
    assert ws["B2"].value.date() == monday
    assert [ws.cell(2, c).value for c in range(3, 7)] == [8, 4.5, 5, "Mevcut"]
    assert ws["C3"].value is None  # defaults are not frozen into overrides
    assert ws["A2"].protection.locked and not ws["C2"].protection.locked
    assert ws.freeze_panes == "C2" and ws.auto_filter.ref == "A1:M3"  # + fazla mesai (hafta ici + hafta sonu) kisi/gun/saat sutunlari
    before = db.query(WorkCenterWeek).count()
    result = upload(client, auth, wb)
    assert not result["errors"] and result["inserted"] == result["updated"] == 0
    assert db.query(WorkCenterWeek).count() == before
    ws["C2"] = 6
    ws["C3"] = 0
    ws["D3"] = "4,25"
    ws["E3"] = 0
    wb.active = 1  # saving on the instruction sheet still imports the correct sheet
    result = upload(client, auth, wb)
    assert not result["errors"] and result["inserted"] == result["updated"] == 1
    db.expire_all()
    rows = db.query(WorkCenterWeek).filter_by(work_center_id=iid).order_by(WorkCenterWeek.week_start).all()
    assert rows[0].headcount == 6
    assert (rows[1].headcount, rows[1].efficient_hours_per_person, rows[1].working_days) == (0, 4.25, 0)
    assert db.get(WorkCenter, iid).default_efficient_hours == 4


@pytest.mark.parametrize("col,value", [("C", -1), ("C", 1.5), ("D", 25), ("D", "=2+2"), ("E", 8), ("E", 2.5), ("A", "UNKNOWN-WC")])
def test_invalid_file_is_atomic(client, auth, db, weekly_case, col, value):
    iid, monday = weekly_case
    wb = export(client, auth, iid, monday)
    ws = wb.worksheets[0]
    ws["C2"] = 99
    ws[f"{col}3"] = value
    result = upload(client, auth, wb)
    assert result["errors"] and result["inserted"] == result["updated"] == 0
    db.expire_all()
    assert db.query(WorkCenterWeek).filter_by(work_center_id=iid).one().headcount == 8


def test_blank_resets_only_file_week_and_duplicates_rejected(client, auth, db, weekly_case):
    iid, monday = weekly_case
    wb = export(client, auth, iid, monday)
    ws = wb.worksheets[0]
    ws["B3"] = monday
    assert upload(client, auth, wb)["errors"]
    ws.delete_rows(3)
    for col in "CDEF":
        ws[f"{col}2"] = None
    db.add(WorkCenterWeek(work_center_id=iid, week_start=monday + timedelta(weeks=10), headcount=7))
    db.commit()
    assert not upload(client, auth, wb)["errors"]
    remaining = db.query(WorkCenterWeek).filter_by(work_center_id=iid).one()
    assert remaining.week_start == monday + timedelta(weeks=10)
